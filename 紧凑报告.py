"""在原有公式之上组织五张阅读表，原始结果和计算依据集中留痕。"""
from collections import defaultdict
from copy import copy
from datetime import date, datetime
from decimal import Decimal
import math
import unicodedata

from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.views import Selection

from 报告列识别 import identify_columns


阅读表 = ('核对总览', '待处理事项', '月度核对', '每日统计', '对应明细')
阅读颜色 = 'BDD7EE'
留痕颜色 = 'C6E0B4'
人工方式 = ('人工全查', '人工抽样')


def _rows(sheet):
    headers = identify_columns(sheet)
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        yield number, {key: values[column - 1] for key, column in headers.items()}


def _lookup(book, name, label, key, key_label='事项编号', default='""'):
    headers = identify_columns(book[name])
    end = max(2, book[name].max_row)
    value_col, key_col = (get_column_letter(headers[x]) for x in (label, key_label))
    return f'IFERROR(INDEX(\'{name}\'!${value_col}$2:${value_col}${end},MATCH({key},\'{name}\'!${key_col}$2:${key_col}${end},0)),{default})'


def _short_result(expression):
    # 原始详细状态保留在绿色表中，这里只缩短显示文字。
    replacements = (
        ('所选对应重复使用原记录', '重复占用，请重选'),
        ('原对应已否定，记录恢复未对应', '已取消对应'),
        ('原对应被人工选择替代，记录未对应', '原对应已被替代'),
        ('已人工核对，保留未对应', '已查看，仍未对应'),
        ('已确认所选对应；原始差额及期间保留', '已确认，差额保留'),
        ('部分原记录已由其他事项建立对应，其余仍未对应', '部分记录已对应'),
        ('原记录已由其他事项建立对应', '已由其他事项对应'),
        ('程序已确认对应', '程序已对上'),
        ('待人工核对', '待核对'),
        ('留存备查，未作人工确认', '留存查看'),
        ('选择无效', '选择无效，请重选'),
    )
    for old, new in replacements:
        expression = f'SUBSTITUTE({expression},"{old}","{new}")'
    return '=' + expression


def _plain_reason(item, bank, journal, income_diff, expense_diff):
    parts = []
    if not bank:
        parts.append('账上有记录，尚未找到银行对应')
    elif not journal:
        parts.append('银行有记录，尚未找到账上对应')
    elif income_diff or expense_diff:
        parts.append('找到可能对应，但收支金额有差额')
    else:
        parts.append('金额相同，对应关系仍需确认')
    if item.get('跨期'):
        parts.append('双方记录跨月')
    if item.get('风险等级') == '范围未知':
        parts.append('资料或余额有缺口')
    reason = str(item.get('判断依据') or '')
    if '冲突' in reason:
        parts.append('关键业务信息有冲突')
    if '批次边界' in reason:
        parts.append('同一批业务的范围不清楚')
    if '多个' in reason or '歧义' in reason:
        parts.append('有多种可能对应')
    return '；'.join(parts)


def _description(records):
    if not records:
        return '尚未找到对应记录'
    summaries = list(dict.fromkeys(str(row.get('摘要') or '') for row in records if row.get('摘要')))
    parties = list(dict.fromkeys(str(row.get('对方及凭证') or '') for row in records if row.get('对方及凭证')))
    days = sorted({str(row['日期'])[:10] for row in records if row.get('日期')})
    day = days[0] if len(days) == 1 else days[0] + '至' + days[-1] if days else ''
    content = '；'.join(summaries[:2] + parties[:2])
    if len(summaries) > 2 or len(parties) > 2:
        content += '等，完整组成见明细'
    return f'{day} {len(records)}笔 ' + content


def _amounts(records):
    values = [Decimal(str(row.get('金额') or 0)) for row in records]
    return (float(sum((max(x, Decimal(0)) for x in values), Decimal(0))),
            float(sum((max(-x, Decimal(0)) for x in values), Decimal(0))))


def _text(cell, text, limit=48):
    """长说明留在批注中，主表先显示一句短说明，不截断原始证据。"""
    text = str(text or '').replace('\r', '').replace('\n', '；')
    cell.value = text if len(text) <= limit else text[:limit - 1] + '…'
    if len(text) > limit:
        cell.comment = Comment(text, '完整说明')
    if cell.value and str(cell.value)[0] in '=+-@':
        cell.value = "'" + cell.value


def _write_items(book):
    records = defaultdict(lambda: {'银行流水': [], '序时账': []})
    first_record = {}
    for number, row in _rows(book['核对明细']):
        records[row['事项编号']][row['来源']].append(row)
        first_record.setdefault(row['事项编号'], number)
    manual = {}
    for name in 人工方式:
        headers = identify_columns(book[name])
        for number, row in _rows(book[name]):
            if row.get('行别') == '核对事项':
                manual[row['事项编号']] = (name, number, headers, row)

    columns = ['核对编号', '处理要求', '哪里需要查看', '银行记录', '序时账记录',
               '收入差额', '支出差额', '人工核对结果', '当前结果', '备注', '完整明细', '事项编号']
    pending = book.create_sheet('待处理事项')
    pending.append(columns)
    matched = book.create_sheet('对应明细')
    matched.append(['核对编号', '对应方式', '银行记录', '银行收入', '银行支出',
                    '序时账记录', '序时账收入', '序时账支出', '收入差额', '支出差额',
                    '当前结果', '核对依据', '完整明细', '事项编号'])
    pending_rows = {}
    # 完整保留原分级名单，只把全查、抽样与留存查看集中到一个入口。
    items = list(_rows(book['复核事项索引']))
    priority = {'人工全查': 0, '人工抽样': 1, '留存备查': 2, '自动确认': 3}
    items.sort(key=lambda item: (priority.get(item[1]['核对方式'], 2), int(item[1]['核对编号'])))
    numbers = {item['事项编号']: item['核对编号'] for _, item in items}
    for _, item in items:
        identity = item['事项编号']
        bank = records[identity]['银行流水']
        journal = records[identity]['序时账']
        bi, be = _amounts(bank)
        ji, je = _amounts(journal)
        mode = item['核对方式']
        if mode != '自动确认':
            row = pending.max_row + 1
            pending.append([item['核对编号'], {'人工全查': '逐项核对', '人工抽样': '抽中核对'}.get(mode, '留存查看'),
                            _plain_reason(item, bank, journal, bi - ji, be - je),
                            _description(bank), _description(journal), bi - ji, be - je,
                            '', '', '', '查看', identity])
            pending_rows[identity] = row
            for column in (3, 4, 5):
                _text(pending.cell(row, column), pending.cell(row, column).value, 32)
            pending.cell(row, 3).comment = Comment('；'.join(str(item.get(key) or '') for key in ('入选原因', '判断依据')), '原判断依据')
            key = f'$L{row}'
            pending.cell(row, 9, _short_result(_lookup(book, '复核事项索引', '当前核对结论', key)))
            target = manual.get(identity)
            if target:
                name, source_row, headers, original = target
                pending.cell(row, 8, original.get('人工核对结果') or None)
                pending.cell(row, 10, original.get('备注') or None)
                old_cell = book[name].cell(source_row, headers['人工核对结果'])
                for validation in book[name].data_validations.dataValidation:
                    if old_cell.coordinate in validation.sqref:
                        new_validation = copy(validation)
                        new_validation.sqref = ''
                        pending.add_data_validation(new_validation)
                        new_validation.add(pending.cell(row, 8))
                # 稳定编号查找新的唯一填写入口，整行排序后仍能找到同一事项。
                owner_key = book[name].cell(source_row, headers['填报编号']).coordinate
                for label, column in (('人工核对结果', 'H'), ('备注', 'J')):
                    value = f'IFERROR(INDEX(\'待处理事项\'!${column}$2:${column}${max(2, len(items) + 1)},MATCH({owner_key},\'待处理事项\'!$L$2:$L${max(2, len(items) + 1)},0)),"")'
                    book[name].cell(source_row, headers[label], f'=IF({value}="","",{value})')
            else:
                pending.cell(row, 8, '本次无需填写')
            target_row = first_record.get(identity, 2)
            pending.cell(row, 11).hyperlink = f"#'{target[0]}'!A{target[1]}" if target else f"#'核对明细'!A{target_row}"
        if bank and journal:
            row = matched.max_row + 1
            matched.append([item['核对编号'], '逐笔对应' if len(bank) == len(journal) == 1 else '整组对应',
                            _description(bank), bi, be, _description(journal), ji, je,
                            bi - ji, be - je, '', item.get('判断依据') or '', '查看', identity])
            for column in (3, 6, 12):
                _text(matched.cell(row, column), matched.cell(row, column).value, 32)
            key = f'$N{row}'
            matched.cell(row, 11, _short_result(_lookup(book, '复核事项索引', '当前核对结论', key)))
            if identity in pending_rows:
                matched.cell(row, 1, f'=HYPERLINK("#\'待处理事项\'!H"&(MATCH({key},\'待处理事项\'!$L$2:$L${max(2,len(items)+1)},0)+1),"{item["核对编号"]}")')
            detail_name = '逐笔核对' if len(bank) == len(journal) == 1 else '整组核对'
            # 详细表保留原始组成、每条记录的当前去向及原文件出处。
            detail_headers = identify_columns(book[detail_name])
            detail_key = get_column_letter(detail_headers['事项编号'])
            matched.cell(row, 13, f'=HYPERLINK("#\'{detail_name}\'!A"&(MATCH({key},\'{detail_name}\'!${detail_key}$2:${detail_key}${max(2,book[detail_name].max_row)},0)+1),"查看")')
    # 绿色人工表只承接输入；黄色填写格仅出现在待处理事项中。
    for name in 人工方式:
        book[name].data_validations.dataValidation = []
    # 旧明细中的编号链接也回到新的填写处，绿色表不再出现第二个入口。
    for name in ('逐笔核对', '整组核对', '未对应记录'):
        headers = identify_columns(book[name])
        for row, item in _rows(book[name]):
            if item['事项编号'] not in pending_rows:
                continue
            key = book[name].cell(row, headers['事项编号']).coordinate
            number = str(item['核对编号'])
            if number.startswith('='):
                number = numbers[item['事项编号']]
            book[name].cell(row, headers['核对编号'], f'=HYPERLINK("#\'待处理事项\'!H"&(MATCH({key},\'待处理事项\'!$L$2:$L${max(2,len(items)+1)},0)+1),"{number}")')
    return pending_rows


def _period_flags(book, name):
    sheet = book[name]
    headers = identify_columns(sheet)
    daily = name == '每日统计'
    balance_labels = ('银行-期末余额', '日记账-期末余额') if daily else ('银行-月末余额', '序时账-月末余额')
    for label in (*balance_labels, '余额差额'):
        if label not in headers:
            column = sheet.max_column + 1
            sheet.cell(1, column, label)
            headers[label] = column
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column, '未提供')
    column = sheet.max_column + 1
    sheet.cell(1, column, '需查看')
    anomaly_days = set()
    if '余额连续性异常' in book:
        for _, item in _rows(book['余额连续性异常']):
            for label in ('日期', '异常日期', 'date'):
                if item.get(label):
                    anomaly_days.add(str(item[label])[:10])
    diff_labels = [label for label in ('收入金额差额', '支出金额差额', '余额差额', '余额衔接差') if label in headers]
    for row in range(2, sheet.max_row + 1):
        period = sheet.cell(row, headers['日期' if daily else '月份']).value
        if not daily and '合计' in str(period):
            continue
        tests = [f'AND(ISNUMBER({get_column_letter(headers[label])}{row}),{get_column_letter(headers[label])}{row}<>0)' for label in diff_labels]
        missing = [f'NOT(ISNUMBER({get_column_letter(headers[label])}{row}))' for label in balance_labels]
        if (daily and str(period)[:10] in anomaly_days) or (not daily and str(period)[:7] in {day[:7] for day in anomaly_days}):
            sheet.cell(row, column, '余额不连续')
        else:
            sheet.cell(row, column, f'=IF(OR({",".join(tests)}),"有差额",IF(OR({",".join(missing)}),"余额未提供",""))')
    return column


def _overview(book):
    sheet = book.create_sheet('核对总览')
    sheet.append(['项目', '结果', '怎么看'])
    rows = [
        ('核对范围', '核对范围', '资料有缺口时，结果要结合原始说明查看。'),
        ('资料情况', '范围说明', ''),
        ('银行记录数', '银行有效交易笔数', '进入本次核对的银行交易笔数。'),
        ('账上记录数', '日记账有效交易笔数', '进入本次核对的银行存款记账笔数。'),
        ('银行金额已对平比例', '银行当前已对平金额覆盖率', '收付金额分开累计；只计入已对应且差额为零的记录。'),
        ('账上金额已对平比例', '序时账当前已对平金额覆盖率', '人工查看过但仍有差额，不算已对平。'),
        ('还需人工核对', '待处理事项数', '在待处理事项的黄色格选择结果；留存查看不计入待办。'),
        ('已完成人工核对', '已处理事项数', '已选有效结果的事项数，有差额的仍保留差额。'),
        ('人工选择有冲突', '人工选择冲突数', '同一原记录被重复采用时，需要重新选择。'),
        ('还没找到对应的记录', '当前未对应原始记录数', '银行和账上合计笔数；未对应不等于漏记。'),
        ('期初余额差额', '期初余额差额', '银行余额减账上余额；未提供可用余额时不补零。'),
        ('余额核对情况', '余额核对', '每日和月度表保留收支、余额及差额。'),
    ]
    if '非工作日交易' in book:
        rows.extend([
            ('非工作日银行收付笔数', '非工作日银行收付笔数', '已对平的记录也保留提示；完整清单在绿色非工作日交易表。'),
            ('非工作日银行收款笔数', '非工作日银行收款笔数', ''),
            ('非工作日银行收款金额', '非工作日银行收款金额', ''),
            ('非工作日银行付款笔数', '非工作日银行付款笔数', ''),
            ('非工作日银行付款金额', '非工作日银行付款金额', ''),
            ('银行交易日历未覆盖笔数', '银行交易日历未覆盖笔数', '超出2021至2026年日历的记录，不能据此判断是否为工作日。'),
        ])
    for title, label, explanation in rows:
        value = _lookup(book, '核对结论', '数值', '"' + label + '"', '项目', '"未提供"')
        if label == '核对范围':
            for old, new in (('范围受限', '资料有缺口，部分结果需核实'), ('范围可用', '本次资料可以用于核对'), ('范围未验证', '尚未确认资料范围')):
                value = f'SUBSTITUTE({value},"{old}","{new}")'
        if label == '范围说明':
            original = next((item['数值'] for _, item in _rows(book['核对结论']) if item['项目'] == label), '')
            value = f'IF({value}="未发现范围疑点","未发现额外资料问题","有需要确认的资料，见原始说明")'
            explanation = str(original or '')
        if label == '余额核对':
            value = f'SUBSTITUTE({value},"余额核对未实施","未提供双方可用余额")'
        sheet.append([title, '=' + value, explanation])
        _text(sheet.cell(sheet.max_row, 3), explanation, 40)
        if '比例' in title:
            sheet.cell(sheet.max_row, 2).number_format = '0.00%'
        elif '差额' in title or '金额' in title:
            sheet.cell(sheet.max_row, 2).number_format = '#,##0.00;[Red]-#,##0.00'
    for name, title in (('月度核对', '有差额的月份'), ('每日统计', '有差额或余额不连续的日期')):
        headers = identify_columns(book[name])
        column = get_column_letter(headers['需查看'])
        rng = f"'{name}'!${column}$2:${column}${max(2,book[name].max_row)}"
        sheet.append([title, f'=COUNTIF({rng},"有差额")+COUNTIF({rng},"余额不连续")', '打开' + name + '，筛选“需查看”列。'])
        sheet.cell(sheet.max_row, 3).hyperlink = f"#'{name}'!A1"
    search = _lookup(book, '核对结论', '数值', '"组合搜索完整性"', '项目')
    sheet.append(['对应查找情况', f'=IF({search}="范围受限","部分查找未完成",IF({search}="存在多解","有多种可能对应","已在程序设定的范围内查找"))', '未找到不等于没有对应，详细检查范围见绿色运行参数表。'])
    sheet.append(['使用方法', '先看总览和每日、月度差额，再处理黄色格。', '浅蓝色表用于阅读；浅绿色表保留完整记录和计算依据。'])


def format_compact_sheet(sheet, *, reading=False):
    """统一字体和紧凑尺寸；留痕不改变值、公式或原行位置。"""
    headers = identify_columns(sheet)
    hidden = {'事项编号'} if sheet.title in ('待处理事项', '对应明细') else set()
    if reading and sheet.title in ('每日统计', '月度核对'):
        hidden.update(label for label in headers if '笔数' in label or '变动净额' in label or label.startswith(('已确认', '未确认')))
    widths = {'核对编号': 7, '处理要求': 9, '对应方式': 9, '哪里需要查看': 26,
              '银行记录': 28, '序时账记录': 28, '人工核对结果': 19, '当前结果': 21,
              '备注': 20, '完整明细': 7, '核对依据': 26, '日期': 11, '月份': 10,
              '需查看': 13, '余额说明': 21, '项目': 27, '结果': 43, '怎么看': 47}
    first = min(headers.values(), default=1)
    for column in range(1, first):
        sheet.column_dimensions[get_column_letter(column)].hidden = True
    for label, column in headers.items():
        dimension = sheet.column_dimensions[get_column_letter(column)]
        dimension.hidden = label in hidden
        amount = any(key in label for key in ('金额', '余额', '收入', '支出', '差额')) and not any(key in label for key in ('说明', '比例', '依据', '原因'))
        dimension.width = widths.get(label, 14 if amount else 28 if any(key in label for key in ('摘要', '原因', '说明', '依据', '记录')) else 16)
    last = max(headers.values(), default=1)
    for letter, dimension in sheet.column_dimensions.items():
        if column_index_from_string(letter) > last:
            dimension.hidden = True
    sheet.sheet_properties.tabColor = 阅读颜色 if reading else 留痕颜色
    sheet.sheet_state = 'visible'
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 100
    sheet.sheet_view.topLeftCell = f'{get_column_letter(first)}1'
    # 旧排版多次冻结会留下多个选区，重新设置前先清除旧分隔窗状态。
    sheet.sheet_view.pane = None
    sheet.sheet_view.selection = [Selection(activeCell=f'{get_column_letter(first)}2', sqref=f'{get_column_letter(first)}2')]
    sheet.freeze_panes = None if sheet.title == '核对总览' else 'A2'
    sheet.sheet_format.defaultRowHeight = 15
    sheet.row_dimensions[1].height = 28
    sheet.auto_filter.ref = None if sheet.title == '核对总览' else f'{get_column_letter(first)}1:{get_column_letter(last)}{sheet.max_row}'
    for cell in sheet[1]:
        if cell.value is not None:
            cell.font = Font(name='MiSans', size=9, bold=True, color='243746')
            cell.fill = PatternFill('solid', fgColor=阅读颜色 if reading else 留痕颜色)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    for row in sheet.iter_rows(min_row=2):
        height = 15
        for cell in row:
            if cell.value is None:
                continue
            font = copy(cell.font)
            font.name, font.sz = 'MiSans', 9
            font.scheme = None
            cell.font = font
            cell.fill = PatternFill(fill_type=None)
            label = str(sheet.cell(1, cell.column).value or '')
            numeric = isinstance(cell.value, (int, float, Decimal)) and not isinstance(cell.value, bool)
            cell.alignment = Alignment(horizontal='right' if numeric else 'left', vertical='top', wrap_text=reading)
            if numeric and any(key in label for key in ('金额', '余额', '收入', '支出', '差额')):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
                dimension = sheet.column_dimensions[cell.column_letter]
                dimension.width = max(dimension.width, len(f'{cell.value:,.2f}') + 1)
            if isinstance(cell.value, (datetime, date)):
                cell.number_format = 'yyyy-mm-dd'
            if reading and label not in hidden and cell.data_type != 'f':
                width = sheet.column_dimensions[cell.column_letter].width
                # 列宽以默认字号计量，9号正文要按较小的字宽估算，减少无用行高。
                text_width = sum(1.5 if unicodedata.east_asian_width(char) in ('W', 'F') else 0.75 for char in str(cell.value))
                height = max(height, math.ceil(text_width / max(1, width - 1)) * 12 + 3)
            if cell.hyperlink or (cell.data_type == 'f' and str(cell.value).startswith('=HYPERLINK(')):
                cell.font = Font(name='MiSans', size=9, color='0563C1', underline='single')
        # 留痕保持单行，完整文字可在编辑栏或原始明细查看。
        sheet.row_dimensions[row[0].row].height = height
        if not reading:
            sheet.row_dimensions[row[0].row].hidden = False
            sheet.row_dimensions[row[0].row].collapsed = False
    if sheet.title != '核对总览':
        sheet.print_title_rows = '1:1'
    sheet.print_area = f'{get_column_letter(first)}1:{get_column_letter(last)}{sheet.max_row}'
    if reading:
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True


def apply_compact_report(book):
    """最终输出调整，必须在原报告的公式和排版完成后调用。"""
    _write_items(book)
    _period_flags(book, '每日统计')
    _period_flags(book, '月度核对')
    _overview(book)
    for sheet in book:
        format_compact_sheet(sheet, reading=sheet.title in 阅读表)
    pending = book['待处理事项']
    for validation in pending.data_validations.dataValidation:
        for cell_range in validation.sqref:
            for row in range(cell_range.min_row, cell_range.max_row + 1):
                pending.cell(row, 8).fill = PatternFill('solid', fgColor='FFF2CC')
    for sheet in (book['待处理事项'], book['对应明细']):
        headers = identify_columns(sheet)
        for label in ('收入差额', '支出差额'):
            column = get_column_letter(headers[label])
            if sheet.max_row > 1:
                sheet.conditional_formatting.add(f'{column}2:{column}{sheet.max_row}', FormulaRule(
                    formula=[f'AND(ISNUMBER({column}2),{column}2<>0)'],
                    fill=PatternFill('solid', fgColor='FFF2CC'), font=Font(color='9C0006')))
    for name in ('每日统计', '月度核对'):
        sheet = book[name]
        column = get_column_letter(identify_columns(sheet)['需查看'])
        if sheet.max_row > 1:
            sheet.conditional_formatting.add(f'{column}2:{column}{sheet.max_row}', FormulaRule(
                formula=[f'OR({column}2="有差额",{column}2="余额不连续")'],
                fill=PatternFill('solid', fgColor='FFF2CC'), font=Font(color='9C0006')))
    order = [book[name] for name in 阅读表]
    book._sheets = order + [sheet for sheet in book if sheet not in order]
    book.active = 0
