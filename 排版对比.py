# -*- coding: utf-8 -*-
"""逐格对比两份 Excel 工作簿，供排版新旧路径的一致性验收。

用法:
    python 排版对比.py 旧文件.xlsx 新文件.xlsx
    from 排版对比 import 对比工作簿
    差异 = 对比工作簿(旧, 新)   # 空列表 = 完全一致
"""
import sys
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

容差 = 1e-9


def _颜色(color):
    if color is None:
        return None
    return (getattr(color, "rgb", None), getattr(color, "theme", None),
            getattr(color, "indexed", None), getattr(color, "tint", None),
            getattr(color, "auto", None))


def _值(value):
    if isinstance(value, float):
        return ("数", round(value, 9))
    if isinstance(value, (datetime, date)):
        return ("日", value.isoformat())
    if value is None:
        return None
    return ("文", str(value))


def _边(side):
    if side is None or side.style is None:
        return None
    return (side.style, _颜色(side.color))


def _格快照(cell):
    f = cell.font
    fill = cell.fill
    a = cell.alignment
    b = cell.border
    link = cell.hyperlink
    return {
        "值": _值(cell.value),
        "类型": cell.data_type if cell.value is not None else None,
        "字体": (f.name, f.size, bool(f.bold), _颜色(f.color), f.underline),
        "填充": (fill.fill_type, _颜色(fill.fgColor)),
        "对齐": (a.horizontal, a.vertical, bool(a.wrap_text)),
        "边框": (_边(b.top), _边(b.bottom), _边(b.left), _边(b.right)),
        "数字格式": cell.number_format if cell.number_format != "General" else None,
        "批注": (cell.comment.text, cell.comment.author) if cell.comment else None,
        "超链接": (getattr(link, "target", None), getattr(link, "location", None)) if link else None,
    }


_格属性中文 = {"值": "值", "类型": "类型", "字体": "字体", "填充": "填充", "对齐": "对齐",
             "边框": "边框", "数字格式": "数字格式", "批注": "批注", "超链接": "超链接"}


def _行快照(dimension):
    return {
        "行高": dimension.height,
        "隐藏": bool(dimension.hidden),
        "分级": dimension.outlineLevel or 0,
        "收起": bool(dimension.collapsed),
    }


def _列快照(dimension):
    return {"列宽": dimension.width, "隐藏": bool(dimension.hidden)}


def _验证快照(validation):
    return (validation.type, validation.formula1, validation.formula2,
            bool(validation.allowBlank), validation.error, validation.errorTitle,
            bool(validation.showErrorMessage), validation.errorStyle,
            tuple(sorted(str(r) for r in validation.sqref.ranges)))


def _条件格式快照(sheet):
    结果 = []
    for item in sheet.conditional_formatting:
        for rule in item.rules:
            填充 = None
            字体 = None
            if rule.dxf is not None:
                if rule.dxf.fill is not None:
                    填充 = _颜色(rule.dxf.fill.bgColor) or _颜色(rule.dxf.fill.fgColor)
                if rule.dxf.font is not None:
                    字体 = _颜色(rule.dxf.font.color)
            结果.append((str(item.sqref), rule.type, tuple(rule.formula or ()), 填充, 字体))
    return sorted(结果)


def _表级快照(sheet, book):
    页设置 = sheet.sheet_properties.pageSetUpPr
    return {
        "状态": sheet.sheet_state,
        "tab色": _颜色(sheet.sheet_properties.tabColor),
        "冻结": sheet.freeze_panes,
        "筛选": sheet.auto_filter.ref or None,
        "网格线": bool(sheet.sheet_view.showGridLines),
        "缩放": sheet.sheet_view.zoomScale or 100,
        "outline汇总在下": sheet.sheet_properties.outlinePr.summaryBelow if sheet.sheet_properties.outlinePr else None,
        "打印": (sheet.page_setup.orientation, sheet.page_setup.paperSize,
               sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight,
               bool(页设置.fitToPage) if 页设置 else None,
               sheet.print_title_rows, sheet.print_area or None,
               round(sheet.page_margins.left, 4), round(sheet.page_margins.right, 4)),
        "合并": tuple(sorted(str(r) for r in sheet.merged_cells.ranges)),
        "数据验证": sorted(_验证快照(v) for v in sheet.data_validations.dataValidation),
        "条件格式": _条件格式快照(sheet),
    }


def _快照(book):
    数据 = {}
    for sheet in book.worksheets:
        格子 = {}
        for (行, 列), cell in sheet._cells.items():
            if cell.value is not None or cell.has_style:
                格子[(行, 列)] = _格快照(cell)
        行 = {号: _行快照(dim) for 号, dim in sheet.row_dimensions.items()}
        列 = {}
        for 名, dim in sheet.column_dimensions.items():
            列[名] = _列快照(dim)
        数据[sheet.title] = {"格": 格子, "行": 行, "列": 列, "表级": _表级快照(sheet, book),
                         "默认行高": sheet.sheet_format.defaultRowHeight}
    return 数据


def _比字典(差异, 位置, 旧, 新, 键名):
    for 键 in sorted(set(旧) | set(新)):
        旧值, 新值 = 旧.get(键), 新.get(键)
        if 旧值 != 新值:
            差异.append(f"{位置}: {键名}{键} 旧={旧值} 新={新值}")


def 对比工作簿(旧路径, 新路径, 差异上限=200):
    """逐格对比两份工作簿，返回差异描述列表；空列表表示完全一致。"""
    旧簿 = load_workbook(旧路径)
    新簿 = load_workbook(新路径)
    差异 = []

    旧表名 = [s.title for s in 旧簿.worksheets]
    新表名 = [s.title for s in 新簿.worksheets]
    if 旧表名 != 新表名:
        差异.append(f"工作表顺序或名单: 旧={旧表名} 新={新表名}")

    旧计算 = 旧簿.calculation
    新计算 = 新簿.calculation
    for 属性 in ("calcMode", "fullCalcOnLoad", "forceFullCalc"):
        if getattr(旧计算, 属性) != getattr(新计算, 属性):
            差异.append(f"工作簿计算设置 {属性}: 旧={getattr(旧计算, 属性)} 新={getattr(新计算, 属性)}")
    if 旧簿.active.title != 新簿.active.title:
        差异.append(f"活动表: 旧={旧簿.active.title} 新={新簿.active.title}")

    旧数据 = _快照(旧簿)
    新数据 = _快照(新簿)
    for 表名 in 旧数据.keys() & 新数据.keys():
        旧表, 新表 = 旧数据[表名], 新数据[表名]
        _比字典(差异, f"{表名}", 旧表["表级"], 新表["表级"], "")
        _比字典(差异, f"{表名}", 旧表["列"], 新表["列"], "列")
        for 行号 in sorted(set(旧表["行"]) | set(新表["行"])):
            旧行 = 旧表["行"].get(行号, {"行高": None, "隐藏": False, "分级": 0, "收起": False})
            新行 = 新表["行"].get(行号, {"隐藏": False, "分级": 0, "收起": False, "行高": None})
            旧缺省 = 旧表["默认行高"] or "默认"
            新缺省 = 新表["默认行高"] or "默认"
            for 键 in ("隐藏", "分级", "收起"):
                if 旧行[键] != 新行[键]:
                    差异.append(f"{表名} 第{行号}行: {键} 旧={旧行[键]} 新={新行[键]}")
            旧高 = 旧行["行高"] if 旧行["行高"] is not None else 旧缺省
            新高 = 新行["行高"] if 新行["行高"] is not None else 新缺省
            if 旧高 != 新高:
                差异.append(f"{表名} 第{行号}行: 行高 旧={旧行['行高']} 新={新行['行高']}")
        for (行号, 列号) in sorted(set(旧表["格"]) | set(新表["格"])):
            旧格 = 旧表["格"].get((行号, 列号))
            新格 = 新表["格"].get((行号, 列号))
            坐标 = f"{get_column_letter(列号)}{行号}"
            if 旧格 is None or 新格 is None:
                缺失 = 旧格 if 旧格 is not None else 新格
                if 缺失 and (缺失["值"] is not None or any(v is not None for k, v in 缺失.items() if k != "值")):
                    差异.append(f"{表名}!{坐标}: 格缺失 旧={'有' if 旧格 else '无'} 新={'有' if 新格 else '无'}")
                continue
            for 键, 中文 in _格属性中文.items():
                if 旧格[键] != 新格[键]:
                    差异.append(f"{表名}!{坐标}: {中文} 旧={旧格[键]} 新={新格[键]}")
            if len(差异) >= 差异上限:
                差异.append(f"……差异超过 {差异上限} 条，截断")
                return 差异
    return 差异


def 对比并报告(旧路径, 新路径, 标签=""):
    差异 = 对比工作簿(旧路径, 新路径)
    前缀 = f"[{标签}] " if 标签 else ""
    if 差异:
        print(f"{前缀}不一致，共 {len(差异)} 条差异：")
        for 条目 in 差异[:30]:
            print("  " + 条目)
        if len(差异) > 30:
            print(f"  ……其余 {len(差异) - 30} 条省略")
        return False
    print(f"{前缀}逐格一致")
    return True


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("用法: python 排版对比.py 旧文件.xlsx 新文件.xlsx")
        sys.exit(2)
    sys.exit(0 if 对比并报告(sys.argv[1], sys.argv[2]) else 1)
