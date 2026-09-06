"""在原报告上组织月度、原始明细和可直接选择的人工核对结果。"""
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from 复核事项 import build_review_items
from 报告展示 import build_display_tables, format_display_sheets


主表 = ('核对结论', '月度核对', '核对明细', '人工全查', '人工抽样', '每日统计')
人工列 = ['事项编号', '行别', '核对原因', '银行记录', '银行金额', '序时账记录', '序时账金额', '差额', '人工核对结果', '当前核对结论', '备注', '原始出处', '填报编号']
原始列 = ['原始记录号', '事项编号', '来源', '日期', '对方及凭证', '摘要', '金额', '程序结论', '核对方式', '当前对应', '当前核对结果', '原始出处', '月份', '已对平金额', '已确认收入差', '已确认支出差', '收入差贡献', '支出差贡献']
索引列 = ['核对编号', '事项编号', '匹配ID', '程序结论', '风险等级', '核对方式', '影响金额', '差异金额', '首次日期', '入选原因', '判断依据', '跨期', '类型', '组金额', '最早日期', '人工核对结果', '当前核对结论']
选项列 = ['选项编号', '事项编号', '选择文字', '自动确认', '差异金额', '跨期', '关系说明', '人工选中', '有效对应', '选择冲突', '显示结果']
成员列 = ['选项编号', '事项编号', '原始记录号', '人工占用', '实际生效', '占用次数', '生效记录号', '显示结果', '已对平']


def _frame(rows, columns):
    return pd.DataFrame(rows, columns=columns)


def _record_key(side, row, index):
    return ('银行' if side == 'bank' else '账面') + str(int(row.get('original_file_row', row.get('original_idx', index))))


def _brief(row):
    fields = row.get('aux_text_fields', {})
    fields = fields if isinstance(fields, dict) else {}
    parts = [pd.Timestamp(row['date']).strftime('%Y-%m-%d'), str(row.get('summary', ''))]
    parts.extend(str(v) for k, v in fields.items() if v is not None and str(v).strip() and str(v) not in parts and any(s in k for s in ('对方', '单位', '凭证', '业务', '批次', '客户', '供应商')))
    voucher = str(row.get('voucher_no', '') or '')
    if voucher and voucher not in parts:
        parts.append('凭证' + voucher)
    return '；'.join(parts)


def build_review_tables(reporter, tables):
    matcher = reporter.matcher
    items = build_review_items(matcher)
    display_numbers = {item['事项编号']: f'{number:03d}' for number,item in enumerate(items,1)}
    source_info = {}
    for source in getattr(reporter.precheck_report, 'source_info', ()):
        side = 'bank' if source.get('来源') == '银行流水' else 'journal'
        source_info[side] = f"{Path(source.get('文件路径','')).name} · {source.get('工作表','')}"
    def origin(side,row,index):
        label = source_info.get(side, '银行流水' if side == 'bank' else '序时账')
        return f"{label} · 第{int(row.get('original_file_row',row.get('original_idx',index)))}行"
    candidates = {x.candidate_id: x for x in getattr(matcher, 'candidates', [])}
    selected = {x.final_match_id: x for x in getattr(matcher, 'selected_candidates', [])}
    index_rows, records, options, members, comparisons = [], [], [], [], []
    manual = {'人工全查': [], '人工抽样': []}
    owner = {}
    for item in items:
        item_id = item['事项编号']
        candidate = selected.get(item.get('匹配ID'))
        index = {k: item.get(k, '') for k in 索引列}
        index.update({'核对编号': display_numbers[item_id], '类型': candidate.match_type if candidate else '单边记录', '组金额': float(max(item['银行收入'] + item['银行支出'], item['日记账收入'] + item['日记账支出'])), '最早日期': item['首次日期'], '人工核对结果': '', '当前核对结论': ''})
        index_rows.append(index)
        for side, field in (('bank', '银行索引'), ('journal', '日记账索引')):
            for idx in item[field]:
                if (side, idx) in owner:
                    raise ValueError('同一原始记录进入多个复核事项')
                owner[(side, idx)] = item
        choices = []
        if candidate is not None:
            choices.append(('确认对应并保留差额' if item['差异金额'] else '确认对应', candidate))
            for alternative in candidate.evidence.get('alternative_candidate_ids', ()):
                other = candidates.get(alternative)
                if other and (other.bank_idxs, other.journal_idxs) != (candidate.bank_idxs, candidate.journal_idxs):
                    if all((other.bank_idxs, other.journal_idxs) != (x.bank_idxs, x.journal_idxs) for _, x in choices):
                        choices.append((f'采用其他对应{len(choices)}', other))
        for number, (label, choice) in enumerate(choices):
            option_id = item_id + f'|{number}'
            difference = float(choice.metrics.total_diff_li) / 10000
            cross = len({pd.Timestamp(x).strftime('%Y-%m') for x in choice.bank_dates + choice.journal_dates}) > 1
            result = ('已确认对应，差额保留' if difference else '已确认对应') + ('；跨期保留' if cross else '')
            if number == 0 and item['核对方式'] == '自动确认':
                result = '程序已确认对应' + ('；跨期保留' if cross else '')
            options.append({'选项编号': option_id, '事项编号': item_id, '选择文字': label, '自动确认': int(number == 0 and item['核对方式'] == '自动确认'), '差异金额': difference, '跨期': int(cross), '关系说明': choice.processing_reason or item['判断依据'], '显示结果': result})
            for side, indexes in (('bank', choice.bank_idxs), ('journal', choice.journal_idxs)):
                frame = getattr(matcher, side)
                for idx in indexes:
                    source_row = frame.loc[idx]
                    key = _record_key(side, source_row, idx)
                    members.append({'选项编号': option_id, '事项编号': item_id, '原始记录号': key})
                    if number:
                        comparisons.append({'事项编号': item_id, '可选结果': label, '来源': '银行流水' if side == 'bank' else '序时账', '日期': source_row['date'], '摘要及对方': _brief(source_row), '金额': float(source_row['amount']), '关系说明': choice.processing_reason or choice.evidence.get('business_basis', ''), '原始出处': origin(side,source_row,idx)})
        mode = item['核对方式']
        if mode in manual:
            bank_rows = [matcher.bank.loc[x] for x in item['银行索引']]
            journal_rows = [matcher.journal.loc[x] for x in item['日记账索引']]
            manual[mode].append({'事项编号': item_id, '填报编号': item_id, '行别': '核对事项', '核对原因': item['入选原因'] + ('；跨期' if item['跨期'] else '') + ('；' + item['判断依据'] if candidate else ''), '银行记录': _brief(bank_rows[0]) if len(bank_rows) == 1 else f'{len(bank_rows)}笔，完整组成如下' if bank_rows else '未找到银行对应', '银行金额': float(sum((x['amount'] for x in bank_rows), Decimal(0))) if bank_rows else None, '序时账记录': _brief(journal_rows[0]) if len(journal_rows) == 1 else f'{len(journal_rows)}笔，完整组成如下' if journal_rows else '未找到序时账对应', '序时账金额': float(sum((x['amount'] for x in journal_rows), Decimal(0))) if journal_rows else None, '差额': float(item['差异金额']), '人工核对结果': '', '当前核对结论': '', '备注': '', '原始出处': '；'.join(origin(side,getattr(matcher,side).loc[idx],idx) for side,field in (('bank','银行索引'),('journal','日记账索引')) for idx in item[field])})
            if len(bank_rows) > 1 or len(journal_rows) > 1:
                for label, side_rows in (('银行', bank_rows), ('序时账', journal_rows)):
                    for row in side_rows:
                        manual[mode].append({'事项编号': item_id, '行别': label+'组成', label+'记录': _brief(row), label+'金额': float(row['amount']), '原始出处': str(int(row.get('original_file_row', row.get('original_idx', 0))))})
    for side, source in (('bank', '银行流水'), ('journal', '序时账')):
        frame = getattr(matcher, side)
        for idx, row in frame.sort_values(['date', 'original_idx'], kind='stable').iterrows():
            item = owner[(side, idx)]
            amount = Decimal(str(row['amount']))
            sign = 1 if side == 'bank' else -1
            fields = row.get('aux_text_fields', {})
            fields = fields if isinstance(fields, dict) else {}
            party = '；'.join(str(v) for k, v in fields.items() if v is not None and any(s in k for s in ('对方', '单位', '凭证', '业务', '批次', '客户', '供应商')))
            voucher = str(row.get('voucher_no', '') or '')
            if voucher and voucher not in party:
                party += ('；' if party else '') + '凭证' + voucher
            records.append({'原始记录号': _record_key(side,row,idx), '事项编号': item['事项编号'], '来源': source, '日期': row['date'], '对方及凭证': party, '摘要': row.get('summary',''), '金额': float(amount), '程序结论': item['程序结论'], '核对方式': item['核对方式'], '当前对应': '', '当前核对结果': '', '原始出处': origin(side,row,idx), '月份': row['date'].strftime('%Y-%m'), '收入差贡献': float(max(amount,Decimal(0))*sign), '支出差贡献': float(max(-amount,Decimal(0))*sign)})
    records.sort(key=lambda row: (display_numbers[row['事项编号']], row['来源'] != '银行流水', row['日期'], row['原始记录号']))
    totals = defaultdict(lambda: {'收入差额':Decimal(0), '支出差额':Decimal(0)})
    for record in records:
        key = (record['月份'],record['事项编号'])
        totals[key]['收入差额'] += Decimal(str(record['收入差贡献']))
        totals[key]['支出差额'] += Decimal(str(record['支出差贡献']))
    contributions = [{'月份': month, '事项编号': item_id, **values, '净发生额差': values['收入差额']-values['支出差额'], '当前核对结论': ''} for (month,item_id),values in sorted(totals.items())]
    monthly = tables['月度统计'].copy()
    monthly = monthly.rename(columns={c: c.replace('日记账','序时账') for c in monthly.columns})
    # 月度原值不变；确认与未确认贡献使用实际原记录计算。
    for column in ['已确认收入差', '未确认收入差', '已确认支出差', '未确认支出差', '余额衔接差', '余额说明']:
        monthly[column] = ''
    result = {
        '月度核对': monthly,
        '核对明细': _frame(records, 原始列),
        '人工全查': _frame(manual['人工全查'],人工列),
        '人工抽样': _frame(manual['人工抽样'],人工列),
        '月度差异组成': _frame(contributions,['月份','事项编号','收入差额','支出差额','净发生额差','已确认收入差','未确认收入差','已确认支出差','未确认支出差','当前核对结论']),
        '其他对应供选择': _frame(comparisons,['事项编号','可选结果','来源','日期','摘要及对方','金额','关系说明','原始出处']),
        '复核事项索引': _frame(index_rows,索引列),
        '复核候选选项': _frame(options,选项列),
        '复核候选组成': _frame(members,成员列),
    }
    for name in ('核对明细','人工全查','人工抽样','月度差异组成','其他对应供选择'):
        result[name].insert(0,'核对编号',result[name]['事项编号'].map(display_numbers))
    result['核对明细'].insert(result['核对明细'].columns.get_loc('当前对应'),'当前核对编号','')
    return build_display_tables(result)


def apply_review_presentation(book):
    """普通Excel公式承接人工选择；按稳定事项和原记录号查找，不依赖行序。"""
    if '复核事项索引' not in book.sheetnames:
        return
    sheets = {name: book[name] for name in ('复核事项索引','复核候选选项','复核候选组成','核对明细','月度核对','人工全查','人工抽样')}
    headers = {name:{cell.value:cell.column for cell in sheet[1] if cell.value is not None} for name,sheet in sheets.items()}
    def cell(name, label, row):
        return f'{get_column_letter(headers[name][label])}{row}'
    def rng(name,label):
        letter=get_column_letter(headers[name][label])
        return f"'{name}'!${letter}$2:${letter}${max(2,sheets[name].max_row)}"
    def put(name,label,row,formula):
        sheets[name].cell(row,headers[name][label],formula)
    def find(name,label,key,key_label='事项编号',default='""'):
        return f'IFERROR(INDEX({rng(name,label)},MATCH({key},{rng(name,key_label)},0)),{default})'
    ix,op,mem,detail='复核事项索引','复核候选选项','复核候选组成','核对明细'
    index_rows={str(sheets[ix].cell(r,headers[ix]['事项编号']).value):r for r in range(2,sheets[ix].max_row+1)}
    option_rows=defaultdict(list)
    for row in range(2,sheets[op].max_row+1):
        item_id=str(sheets[op].cell(row,headers[op]['事项编号']).value)
        option_rows[item_id].append(row)
    # 稳定编号查找人工输入，用户排序或复制完整行后仍能回到正确事项。
    for item_id,row in index_rows.items():
        key=cell(ix,'事项编号',row)
        mode=sheets[ix].cell(row,headers[ix]['核对方式']).value
        raw_choice=find(mode,'人工核对结果',key,'填报编号') if mode in ('人工全查','人工抽样') else '""'
        choice=f'IF({raw_choice}="","",{raw_choice})'
        put(ix,'人工核对结果',row,'='+choice)
    for mode in ('人工全查','人工抽样'):
        sheet=sheets[mode]
        for row in range(2,sheet.max_row+1):
            item_id=sheet.cell(row,headers[mode]['事项编号']).value
            if not item_id:
                continue
            if sheet.cell(row,headers[mode]['行别']).value != '核对事项':
                put(mode,'人工核对结果',row,'='+find(ix,'人工核对结果',cell(mode,'事项编号',row)))
                put(mode,'当前核对结论',row,'='+find(ix,'当前核对结论',cell(mode,'事项编号',row)))
                continue
            choices=[sheets[op].cell(r,headers[op]['选择文字']).value for r in option_rows.get(str(item_id),[])]
            choices.append('否定对应' if choices else '保留未对应')
            validation=DataValidation(type='list',formula1='"'+','.join(choices)+'"',allow_blank=True)
            validation.errorTitle='请选择本事项的结果'
            validation.error='请选择下拉列表中的对应结果。'
            validation.showErrorMessage=True
            validation.errorStyle='stop'
            sheet.add_data_validation(validation)
            validation.add(cell(mode,'人工核对结果',row))
            put(mode,'当前核对结论',row,'='+find(ix,'当前核对结论',cell(mode,'事项编号',row)))
    for row in range(2,sheets[op].max_row+1):
        key=cell(op,'事项编号',row)
        choice=find(ix,'人工核对结果',key)
        put(op,'人工选中',row,f'=--({choice}={cell(op,"选择文字",row)})')
    for row in range(2,sheets[mem].max_row+1):
        opt=cell(mem,'选项编号',row)
        key=cell(mem,'原始记录号',row)
        put(mem,'人工占用',row,'='+find(op,'人工选中',opt,'选项编号','0'))
        put(mem,'占用次数',row,f'=SUMIF({rng(mem,"原始记录号")},{key},{rng(mem,"人工占用")})')
    for row in range(2,sheets[op].max_row+1):
        opt=cell(op,'选项编号',row)
        manual=cell(op,'人工选中',row)
        choice=find(ix,'人工核对结果',cell(op,'事项编号',row))
        put(op,'选择冲突',row,f'=COUNTIFS({rng(mem,"选项编号")},{opt},{rng(mem,"占用次数")},">1")')
        # 明确人工选择优先；被替代的自动组整体退出，释放全部成员。
        put(op,'有效对应',row,f'=IF({manual}=1,--({cell(op,"选择冲突",row)}=0),IF(AND({cell(op,"自动确认",row)}=1,{choice}=""),--(SUMIF({rng(mem,"选项编号")},{opt},{rng(mem,"占用次数")})=0),0))')
    for row in range(2,sheets[mem].max_row+1):
        opt=cell(mem,'选项编号',row)
        put(mem,'实际生效',row,'='+find(op,'有效对应',opt,'选项编号','0'))
        put(mem,'生效记录号',row,f'=IF({cell(mem,"实际生效",row)}=1,{cell(mem,"原始记录号",row)},"")')
        put(mem,'显示结果',row,'='+find(op,'显示结果',opt,'选项编号'))
        put(mem,'已对平',row,f'=IF({cell(mem,"实际生效",row)}=1,--({find(op,"差异金额",opt,"选项编号","1")}=0),0)')
    for item_id,row in index_rows.items():
        key=cell(ix,'事项编号',row)
        choice=cell(ix,'人工核对结果',row)
        selected=f'SUMIFS({rng(op,"人工选中")},{rng(op,"事项编号")},{key})'
        conflict=f'SUMIFS({rng(op,"选择冲突")},{rng(op,"事项编号")},{key},{rng(op,"人工选中")},1)'
        effective=f'SUMIFS({rng(op,"有效对应")},{rng(op,"事项编号")},{key})'
        mode = sheets[ix].cell(row,headers[ix]['核对方式']).value
        default = '待人工核对' if mode in ('人工全查','人工抽样') else '原对应被人工选择替代，记录未对应' if mode == '自动确认' else '留存备查，未作人工确认'
        formula=f'=IF({choice}="否定对应","原对应已否定，记录恢复未对应",IF({choice}="保留未对应","已人工核对，保留未对应",IF({selected}>0,IF({conflict}>0,"所选对应重复使用原记录","已确认所选对应；原始差额及期间保留"),IF({effective}>0,"程序已确认对应",IF({choice}<>"","选择无效","{default}")))))'
        # 单边原记录可能已经被其他事项的人工候选采用，显示实际去向。
        if not option_rows.get(item_id):
            all_records = f'COUNTIF({rng(detail,"事项编号")},{key})'
            effective_records = f'COUNTIFS({rng(detail,"事项编号")},{key},{rng(detail,"当前对应")},"?*")'
            formula = f'=IF({effective_records}>0,IF({effective_records}={all_records},"原记录已由其他事项建立对应","部分原记录已由其他事项建立对应，其余仍未对应"),{formula[1:]})'
        put(ix,'当前核对结论',row,formula)
    for row in range(2,sheets[detail].max_row+1):
        key=cell(detail,'原始记录号',row)
        lookup=f'MATCH({key},{rng(mem,"生效记录号")},0)'
        put(detail,'当前对应',row,f'=IFERROR(INDEX({rng(mem,"事项编号")},{lookup}),"")')
        put(detail,'当前核对编号',row,'='+find(ix,'核对编号',cell(detail,'当前对应',row)))
        initial_id = str(sheets[detail].cell(row,headers[detail]['事项编号']).value)
        initial=find(ix,'当前核对结论',cell(detail,'事项编号',row))
        if option_rows.get(initial_id):
            choice = find(ix,'人工核对结果',cell(detail,'事项编号',row))
            initial = f'IF({choice}<>"","该记录当前未对应；请按当前对应列查看",IF({find(ix,"核对方式",cell(detail,"事项编号",row))}="自动确认","原对应已被替代，该记录当前未对应",{initial}))'
        put(detail,'当前核对结果',row,f'=IFERROR(INDEX({rng(mem,"显示结果")},{lookup}),{initial})')
        put(detail,'已对平金额',row,f'=IFERROR(INDEX({rng(mem,"已对平")},{lookup})*ABS({cell(detail,"金额",row)}),0)')
        for category in ('收入','支出'):
            put(detail,'已确认'+category+'差',row,f'=IF({cell(detail,"当前对应",row)}<>"",{cell(detail,category+"差贡献",row)},0)')
    monthly=sheets['月度核对']
    for row in range(2,monthly.max_row+1):
        month=cell('月度核对','月份',row)
        for category in ('收入','支出'):
            put('月度核对','已确认'+category+'差',row,f'=SUMIF({rng(detail,"月份")},{month},{rng(detail,"已确认"+category+"差")})')
            put('月度核对','未确认'+category+'差',row,f'={cell("月度核对",category+"金额差额",row)}-{cell("月度核对","已确认"+category+"差",row)}')
        bal_headers=headers['月度核对']
        if '余额差额' not in bal_headers:
            put('月度核对','余额说明',row,'本次未提供双方可用余额')
        else:
            b=cell('月度核对','银行-月末余额',row)
            j=cell('月度核对','序时账-月末余额',row)
            put('月度核对','余额说明',row,f'=IF(AND(ISNUMBER({b}),ISNUMBER({j})),"按双方原始余额核对","本月余额无法完整确认")')
            if row>2:
                previous=cell('月度核对','余额差额',row-1)
                current=cell('月度核对','余额差额',row)
                inc=cell('月度核对','收入金额差额',row)
                exp=cell('月度核对','支出金额差额',row)
                put('月度核对','余额衔接差',row,f'=IF(AND(ISNUMBER({b}),ISNUMBER({j}),ISNUMBER({previous})),{current}-{previous}-{inc}+{exp},"")')
    if '月度差异组成' in book:
        sheet=book['月度差异组成']
        hs={c.value:c.column for c in sheet[1] if c.value}
        for row in range(2,sheet.max_row+1):
            key=f'{get_column_letter(hs["事项编号"])}{row}'
            month=f'{get_column_letter(hs["月份"])}{row}'
            for category in ('收入','支出'):
                confirmed = f'{get_column_letter(hs["已确认"+category+"差"])}{row}'
                difference = f'{get_column_letter(hs[category+"差额"])}{row}'
                sheet.cell(row,hs['已确认'+category+'差'],f'=SUMIFS({rng(detail,"已确认"+category+"差")},{rng(detail,"事项编号")},{key},{rng(detail,"月份")},{month})')
                sheet.cell(row,hs['未确认'+category+'差'],f'={difference}-{confirmed}')
            remaining=f'COUNTIFS({rng(detail,"事项编号")},{key},{rng(detail,"月份")},{month})-COUNTIFS({rng(detail,"事项编号")},{key},{rng(detail,"月份")},{month},{rng(detail,"当前对应")},"?*")'
            sheet.cell(row,hs['当前核对结论'],f'=IF({remaining}>0,"本月原记录仍有未对应；详见核对明细","本月原记录均已建立对应，原始差额保留")')
    _summary_formulas(book,sheets,headers,rng)
    _format_sheets(book)
    format_display_sheets(book)


def _summary_formulas(book,sheets,headers,rng):
    sheet=book['核对结论']
    hs={c.value:c.column for c in sheet[1] if c.value}
    if not {'项目','数值'} <= hs.keys():
        return
    rows={sheet.cell(r,hs['项目']).value:r for r in range(2,sheet.max_row+1)}
    def set_value(label,value):
        row=rows.get(label)
        if row is None:
            row=sheet.max_row+1
            sheet.cell(row,hs['项目'],label)
            rows[label]=row
        sheet.cell(row,hs['数值'],value)
    ix='复核事项索引'
    for risk in ('低风险','中风险','高风险','范围未知'):
        set_value(risk+'事项数',f'=COUNTIF({rng(ix,"风险等级")},"{risk}")')
        set_value(risk+'事项金额',f'=SUMIF({rng(ix,"风险等级")},"{risk}",{rng(ix,"影响金额")})')
    set_value('中风险抽样总数',f'=COUNTIF({rng(ix,"风险等级")},"中风险")')
    set_value('中风险抽中待核查数',f'=COUNTIF({rng(ix,"核对方式")},"人工抽样")')
    set_value('当前未对应原始记录数',f'=COUNTA({rng("核对明细","原始记录号")})-COUNTIF({rng("核对明细","当前对应")},"?*")')
    set_value('使用顺序','先看月度发生额和余额，再看核对明细；人工全查及人工抽样只在黄色格选择结果。多笔业务的组成直接展开。')
    set_value('金额差额口径','收入与支出分别比较，差额不相互抵销。人工表金额为各侧净额，差额为收入差和支出差绝对值之和；原始各笔金额见组成。')
    set_value('当前覆盖率口径','只统计有效对应且差额为零的原记录绝对金额；已人工查看但仍未对应或仍有差额，不计作已对平。')
    set_value('月度说明','月度始终显示双方原始发生额和余额；已确认差表示对应关系已确认，金额与记账月份均未改写；月度差异组成可查看仍未确认部分。')
    set_value('专题说明','退款冲销重付、手续费及净额、截止性差异、重复线索及余额检查等，有事项时直接显示；月度笔数与解释分项列、技术留存表可在Excel取消隐藏查看。')
    pending=[]
    for mode in ('人工全查','人工抽样'):
        mode_range=rng(ix,'核对方式')
        choice=rng(ix,'人工核对结果')
        set_value(mode+'事项数',f'=COUNTIF({mode_range},"{mode}")')
        expression=f'COUNTIFS({mode_range},"{mode}",{choice},"")+COUNTIFS({mode_range},"{mode}",{rng(ix,"当前核对结论")},"选择无效")+COUNTIFS({mode_range},"{mode}",{rng(ix,"当前核对结论")},"所选对应重复使用原记录")'
        pending.append(expression)
        set_value(mode+'未核对数','='+expression)
    set_value('待处理事项数','='+'+'.join(pending))
    set_value('已处理事项数',f'=COUNTIF({rng(ix,"人工核对结果")},"<>")')
    set_value('人工选择冲突数',f'=COUNTIF({rng(ix,"当前核对结论")},"所选对应重复使用原记录")')
    set_value('人工结果说明','在人工全查或人工抽样中选择结果；确认不改变原始金额和期间，否定取消对应，候选冲突不计入有效覆盖。')
    for source,label in (('银行流水','银行'),('序时账','序时账')):
        numerator=f'SUMIF({rng("核对明细","来源")},"{source}",{rng("核对明细","已对平金额")})'
        total=sum(abs(float(sheets['核对明细'].cell(r,headers['核对明细']['金额']).value or 0)) for r in range(2,sheets['核对明细'].max_row+1) if sheets['核对明细'].cell(r,headers['核对明细']['来源']).value==source)
        set_value(label+'当前已对平金额覆盖率',f'={numerator}/{total}' if total else 0)
        sheet.cell(rows[label+'当前已对平金额覆盖率'],hs['数值']).number_format='0.00%'
    visible = ['使用顺序','核对范围','范围说明','银行有效交易笔数','日记账有效交易笔数','实际执行重要性水平','明显微小错报临界值','人工全查事项数','人工全查未核对数','人工抽样事项数','人工抽样未核对数','待处理事项数','已处理事项数','人工选择冲突数','当前未对应原始记录数','银行当前已对平金额覆盖率','序时账当前已对平金额覆盖率','余额核对','期初余额差额','期初余额状态','人工结果说明','金额差额口径','当前覆盖率口径','月度说明','专题说明']
    for label,row in rows.items():
        sheet.row_dimensions[row].hidden = label not in visible
    set_value('已处理事项数',f'=COUNTIF({rng(ix,"人工核对结果")},"?*")-COUNTIF({rng(ix,"当前核对结论")},"选择无效")-COUNTIF({rng(ix,"当前核对结论")},"所选对应重复使用原记录")')
    contents = [(sheet.cell(row,hs['项目']).value,sheet.cell(row,hs['数值']).value,sheet.cell(row,hs['数值']).number_format) for row in range(2,sheet.max_row+1)]
    priority={label:number for number,label in enumerate(visible)}
    contents.sort(key=lambda item:priority.get(item[0],len(priority)))
    for row,(label,value,number_format) in enumerate(contents,2):
        sheet.cell(row,hs['项目'],label)
        sheet.cell(row,hs['数值'],value).number_format=number_format
        sheet.row_dimensions[row].hidden=label not in visible




def _format_sheets(book):
    duplicates={'疑点事项','自动归集事项','银行侧待查','日记账侧待查','逐笔匹配','整组勾稽','匹配组成','月度统计','其他可能对应明细'}
    technical={'运行参数','大模型辅助明细','解析异常明细','运行资料与映射','数据入口处置','复核事项索引','复核候选选项','复核候选组成'}
    order=[book[name] for name in 主表 if name in book]
    book._sheets=order+[sheet for sheet in book if sheet not in order]
    for sheet in book:
        has_rows=sheet.max_row>1 and any(c.value is not None for row in sheet.iter_rows(min_row=2) for c in row)
        sheet.sheet_state='visible' if sheet.title in 主表 or (sheet.title not in duplicates|technical and has_rows) else 'hidden'
        if sheet.title not in 主表 and sheet.title not in ('月度差异组成','其他对应供选择'):
            continue
        hs={c.value:c.column for c in sheet[1] if c.value is not None}
        if not hs:
            continue
        first=min(hs.values())
        for column in range(1,first):
            sheet.column_dimensions[get_column_letter(column)].hidden=True
        widths={'核对编号':10,'当前核对编号':12,'事项编号':24,'行别':12,'核对原因':28,'银行记录':36,'序时账记录':36,'银行金额':16,'序时账金额':16,'差额':15,'人工核对结果':24,'当前核对结论':36,'备注':25,'原始出处':24,'来源':12,'日期':13,'摘要':30,'对方及凭证':24,'金额':16,'程序结论':18,'核对方式':14,'当前对应':24,'当前核对结果':34,'月份':12,'余额说明':30,'项目':29,'数值':70,'摘要及对方':38,'关系说明':40}
        hidden={'原始记录号','收入差贡献','支出差贡献','已对平金额','已确认收入差','已确认支出差'} if sheet.title=='核对明细' else set()
        if '核对编号' in hs:
            hidden.add('事项编号')
        if sheet.title=='核对明细':
            hidden.update(('月份','当前对应'))
        if sheet.title in ('人工全查','人工抽样'):
            hidden.add('填报编号')
        if sheet.title == '月度核对':
            hidden.update(label for label in hs if '笔数' in label or '变动净额' in label or label.startswith('已确认') or label.startswith('未确认'))
        for label,column in hs.items():
            letter=get_column_letter(column)
            sheet.column_dimensions[letter].hidden=label in hidden
            sheet.column_dimensions[letter].width=widths.get(label,17)
        sheet.sheet_view.showGridLines=False
        sheet.sheet_view.zoomScale=90
        sheet.freeze_panes=f'{get_column_letter(first+1)}2'
        sheet.auto_filter.ref=f'{get_column_letter(first)}1:{get_column_letter(sheet.max_column)}{max(1,sheet.max_row)}'
        for c in sheet[1]:
            c.font=Font(name='微软雅黑',size=11,bold=True,color='FFFFFF')
            c.fill=PatternFill('solid',fgColor='24445B')
            c.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[1].height=36
        for row in range(2,sheet.max_row+1):
            height=32
            for label,column in hs.items():
                c=sheet.cell(row,column)
                c.font=Font(name='微软雅黑',size=10,color='243746')
                c.fill=PatternFill('solid',fgColor='F1F5F8' if row%2 else 'FFFFFF')
                c.alignment=Alignment(wrap_text=True,vertical='top')
                if any(k in label for k in ('金额','收入','支出','余额','差额','发生额差')) and label not in ('余额说明',):
                    c.number_format='#,##0.00;[Red]-#,##0.00'
                if label=='日期':
                    c.number_format='yyyy-mm-dd'
                if c.data_type!='f' and label not in hidden:
                    width=widths.get(label,17)
                    lines=sum(max(1,(sum(2 if ord(x)>127 else 1 for x in part)+width-1)//width) for part in str(c.value or '').split('\n'))
                    height=max(height,lines*15+5)
                if label=='人工核对结果' and sheet.title in ('人工全查','人工抽样') and sheet.cell(row,hs['行别']).value == '核对事项':
                    c.fill=PatternFill('solid',fgColor='FFF2CC')
            sheet.row_dimensions[row].height=min(height,300)
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape'
        sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1
        sheet.page_setup.fitToHeight=0
        sheet.print_title_rows='1:1'
        sheet.print_area=f'{get_column_letter(first)}1:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        sheet.page_margins.left=sheet.page_margins.right=0.25
    book.active=0
    book.calculation.calcMode='auto'
    book.calculation.fullCalcOnLoad=True
    book.calculation.forceFullCalc=True
