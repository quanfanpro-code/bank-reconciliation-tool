"""将现有核对证据组织为三张供人阅读的表，不改变匹配或风险判断。"""

from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MAIN_SHEETS = ("核对概览", "待复核事项", "全部核对明细")
DETAIL_COLUMNS = ["事项编号", "行别", "核对结果", "银行日期", "银行摘要", "银行金额",
                  "日记账日期", "日记账摘要", "日记账金额", "差额", "依据及说明", "原始出处"]
ISSUE_COLUMNS = ["事项编号", "需要关注的问题", "银行记录", "日记账记录", "核查建议", "后续状态", "处理说明"]
SOURCE_COLUMNS = ["来源", "原文件行号", "日期", "金额", "摘要", "辅助文字", "凭证号", "账户", "币种", "原始出处", "事项编号"]
INDEX_COLUMNS = ["事项编号", "匹配ID", "原匹配ID", "类型", "最终状态", "风险等级", "判断依据", "检索文字", "最早日期",
                 "组金额", "银行笔数", "日记账笔数", "需要复核", "跨期", "全局事项", "核对结果", "需要关注的问题",
                 "银行记录", "日记账记录", "核查建议"]


def _text(value):
    return "" if value is None or pd.isna(value) else str(value)


def _records(tables, name):
    return tables.get(name, pd.DataFrame()).fillna("").to_dict("records")


def _sum(rows):
    return float(sum((Decimal(str(row["金额"])) for row in rows), Decimal(0)))


def _date(value):
    return pd.Timestamp(value).strftime("%Y-%m-%d") if _text(value) else ""


def _period(rows):
    dates = sorted({_date(row["日期"]) for row in rows if _text(row["日期"])})
    return (dates[0] if len(dates) == 1 else f"{dates[0]} 至 {dates[-1]}") if dates else "无记录"


def _summary(rows):
    if not rows:
        return "无对应记录"
    if len(rows) == 1:
        row = rows[0]
        voucher = f"；凭证 {row['凭证号']}" if row.get("凭证号") else ""
        auxiliary = row.get("辅助文字", "")
        extra = f"；{auxiliary}" if auxiliary else ""
        return f"{row['摘要']}{extra}{voucher}"
    income = _sum([r for r in rows if float(r["金额"]) > 0])
    expense = -_sum([r for r in rows if float(r["金额"]) < 0])
    return f"{len(rows)} 笔；收入 {income:,.2f}；支出 {expense:,.2f}"


def _fact(rows):
    if not rows:
        return "无对应记录"
    return f"{_period(rows)}\n{_summary(rows)}\n净额 {_sum(rows):,.2f}"


def build_source_records(bank, journal, source_info):
    """保留双方每条标准化记录，用于独立核验展示的完整性。"""
    info = source_info.fillna("").to_dict("records") if source_info is not None else []
    rows = []
    for side, frame in (("银行流水", bank), ("日记账", journal)):
        meta = next((r for r in info if ("流水" in str(r.get("来源", ""))) == (side == "银行流水")), {})
        path = str(meta.get("文件路径", ""))
        origin = " / ".join(x for x in (Path(path).name if path else side, str(meta.get("工作表", ""))) if x)
        for index, row in frame.iterrows():
            original = int(row.get("original_file_row", row.get("original_idx", index)))
            auxiliary = row.get("aux_text_fields", {})
            rows.append({"来源": side, "原文件行号": original, "日期": _date(row.get("date")),
                         "金额": float(row["amount"]), "摘要": _text(row.get("summary")),
                         "辅助文字": "；".join(f"{k}：{v}" for k, v in auxiliary.items() if _text(v) and _text(v) != _text(row.get("summary"))) if isinstance(auxiliary, dict) else _text(auxiliary),
                         "凭证号": _text(row.get("voucher_no")), "账户": _text(row.get("account")),
                         "币种": _text(row.get("currency")), "原始出处": f"{origin} 第 {original} 行", "事项编号": ""})
    return pd.DataFrame(rows, columns=SOURCE_COLUMNS)


def _side(value):
    return "银行流水" if "流水" in str(value) else "日记账"


def _key(row):
    return _side(row["来源"]), str(int(float(row["原文件行号"])))


def _explanation(group, issue, cross):
    """陈述业务事实及实际限制；评分细节仍保留在技术附表。"""
    reason = str(issue.get("判断依据", group.get("判断依据", "")))
    difference = float(group.get("总差额", 0) or 0)
    parts = [f"双方收支分别比较，差异 {difference:,.2f}。" if difference else "双方金额一致。"]
    small = "明显微小" in reason and "自动确认" in reason
    if small:
        parts.append("按明显微小金额规则保留，对应关系未独立验证。")
    elif group.get("最终状态") == "整组勾稽一致" or int(group.get("银行笔数", 0) or 0) > 1 or int(group.get("日记账笔数", 0) or 0) > 1:
        parts.append("按完整业务组核对，组内各条记录不表示逐笔对应。")
    else:
        parts.append("结合金额、日期和业务文字建立对应。")
    constraints = [(("歧义", "多个", "多解"), "存在其他可能对应，不能证明逐笔唯一。"),
                   (("依据不足", "可信度", "证据不足"), "现有业务依据不足，需要结合回单或凭证判断。"),
                   (("冲突", "不一致"), "业务文字或所属期间存在不一致，需要核实业务归属。"),
                   (("批次", "边界", "范围闭合"), "业务批次或组成范围仍需核实。"),
                   (("断档", "余额连续"), "相关期间余额衔接异常，需要核查是否缺少记录。"),
                   (("范围未知", "范围受限", "总体资料"), "核对资料范围尚未确认完整。")]
    for words, sentence in constraints:
        if any(word in reason for word in words) and not (small and words in (("歧义", "多个", "多解"), ("依据不足", "可信度", "证据不足"))):
            parts.append(sentence)
    if "未抽中" in reason:
        parts.append("本次中风险抽样未抽中，留存备查。")
    elif "中风险等距抽样：抽中" in reason:
        parts.append("本次中风险抽样已抽中。")
    if cross:
        parts.append("双方记录涉及跨期，金额对应不代表入账期间正确。")
    if group.get("类型") == "手续费净额":
        parts.append(f"按手续费净额关系核对，费用金额 {float(group.get('费用金额', 0) or 0):,.2f}。")
    return "".join(parts)


def build_readable_tables(tables):
    """根据既有事实构造阅读主表及筛选索引。"""
    sources = _records(tables, "阅读来源记录")
    for source in sources:
        source["事项编号"] = ""
    by_key = {_key(r): r for r in sources}
    components = defaultdict(list)
    for row in _records(tables, "匹配组成"):
        if _key(row) in by_key:
            components[str(row["匹配ID"])].append(by_key[_key(row)])
    issues = {str(r["匹配ID"]): r for r in _records(tables, "疑点事项") if r.get("匹配ID")}
    cross_ids = {str(r["匹配ID"]) for r in _records(tables, "截止性差异") if r.get("匹配ID")}
    events = defaultdict(list)
    clues = defaultdict(list)
    for row in _records(tables, "退款冲销重付"):
        events[_key(row)].append(row)
    for row in _records(tables, "重复线索"):
        clues[_key(row)].append(row)
    index_rows, details = [], []

    def add(group, members, issue=None, global_problem="", global_action=""):
        issue = issue or {}
        number = f"事项{len(index_rows) + 1:04d}"
        bank = [r for r in members if r["来源"] == "银行流水"]
        journal = [r for r in members if r["来源"] == "日记账"]
        cross = str(group.get("匹配ID", "")) in cross_ids or any(e.get("是否跨期") == "是" for r in members for e in events[_key(r)])
        status = group.get("最终状态", "尚未找到对应")
        risk = issue.get("风险等级", group.get("风险等级", "范围未知"))
        reason = str(issue.get("判断依据", group.get("判断依据", "")))
        retained = risk == "低风险" or (risk == "中风险" and "未抽中" in str(issue.get("建议动作", "")))
        needs = bool(global_problem or cross or (not bank or not journal) or (risk in {"高风险", "范围未知"}) or (status in {"疑点事项", "自动归集事项"} and not retained))
        result = "资料或余额需核查" if global_problem else "尚未找到对应" if not bank or not journal else "留存备查" if retained else str(status)
        explanation = global_problem or (_explanation(group, issue, cross) if bank and journal else "本侧有记录，对侧尚未找到足够依据建立对应关系。")
        if cross and "跨期" not in explanation:
            explanation += "记录涉及跨期，需要核查入账期间。"
        linked_events = {e["业务链ID"]: e for r in members for e in events[_key(r)]}
        if linked_events:
            explanation += "涉及" + "、".join(sorted({str(e["业务类型"]) for e in linked_events.values()})) + "，需连同相关退回、冲销或重付记录整体查看。"
        if any(clues[_key(r)] for r in members):
            explanation += "存在相似重复记录线索，已保留原记录，不能据此认定重复入账。"
        if result == "自动确认":
            result = "金额一致，期间待核查" if cross else "金额一致，已自动核对"
        elif result in {"自动归集事项", "疑点事项"}:
            result = "金额差异，需核查" if float(group.get("总差额", 0) or 0) else "对应关系需核查"
        action = global_action or ("核对回单、入账凭证及相关期间，确认业务归属与差异原因。" if needs else "留存本次核对结果，按需要查阅组成记录。")
        if cross:
            action += "核查两侧入账期间及截止性。"
        if risk == "范围未知" and not global_problem:
            action = "先核实相关期间资料是否完整，再" + action
        problem_parts = []
        if not bank or not journal:
            problem_parts.append("尚未找到对应记录")
        difference = float(group.get("总差额", 0) or 0)
        if difference:
            problem_parts.append(f"收支分别比较，金额差异 {difference:,.2f}")
        if cross:
            problem_parts.append("跨期，需要核查入账期间")
            if not difference and bank and journal:
                action = "核查两侧入账期间及截止性，结合回单和凭证确认。"
        if not problem_parts:
            problem_parts.append("对应依据或业务归属仍需核实" if needs else "按既有规则保留核对结果")
        problem = global_problem or "；".join(problem_parts) + "。"
        row = {"事项编号": number, "匹配ID": number, "原匹配ID": group.get("匹配ID", ""),
               "类型": group.get("类型", "资料或余额" if global_problem else "单边记录"), "最终状态": status,
               "风险等级": risk, "判断依据": reason or explanation,
               "检索文字": "；".join(str(r.get(k, "")) for r in members for k in ("摘要", "辅助文字", "凭证号")),
               "最早日期": min((r["日期"] for r in members), default=group.get("最早日期", "")),
               "组金额": max(sum(abs(float(r["金额"])) for r in bank), sum(abs(float(r["金额"])) for r in journal)),
               "银行笔数": len(bank), "日记账笔数": len(journal), "需要复核": "是" if needs else "否",
               "跨期": "是" if cross else "否", "全局事项": "是" if global_problem else "否", "核对结果": result,
               "需要关注的问题": (f"{risk}：" if needs and risk in {"高风险", "中风险", "范围未知"} else "") + problem,
               "银行记录": _fact(bank), "日记账记录": _fact(journal), "核查建议": action}
        index_rows.append(row)
        for member in members:
            member["事项编号"] = number
        base = {c: "" for c in DETAIL_COLUMNS}
        base.update({"事项编号": number, "行别": "业务合计" if len(members) > 2 or len(bank) > 1 or len(journal) > 1 else "逐笔记录",
                     "核对结果": result, "银行日期": _period(bank) if bank else "", "银行摘要": _summary(bank),
                     "银行金额": _sum(bank) if bank else "", "日记账日期": _period(journal) if journal else "",
                     "日记账摘要": _summary(journal), "日记账金额": _sum(journal) if journal else "",
                     "差额": float(group.get("总差额", 0) or 0) if bank and journal else "",
                     "依据及说明": explanation, "原始出处": "\n".join(r["原始出处"] for r in members) if len(members) <= 2 else "展开下方组成记录查看原始出处"})
        if global_problem:
            base.update({"行别": "范围事项", "银行摘要": "", "日记账摘要": "", "原始出处": group.get("出处", "输入检查及余额证据")})
        details.append(base)
        if base["行别"] == "业务合计":
            for member in members:
                line = {c: "" for c in DETAIL_COLUMNS}
                prefix = "银行" if member["来源"] == "银行流水" else "日记账"
                line.update({"事项编号": number, "行别": "组成记录", f"{prefix}日期": member["日期"],
                             f"{prefix}摘要": _summary([member]), f"{prefix}金额": member["金额"],
                             "原始出处": member["原始出处"]})
                details.append(line)

    groups = _records(tables, "逐笔匹配") + _records(tables, "整组勾稽")
    groups.sort(key=lambda r: (_text(r.get("最早日期")), str(r.get("匹配ID", ""))))
    for group in groups:
        key = str(group["匹配ID"])
        add(group, components[key], issues.get(key))
    # 尚未被双方关系占用的同侧业务链归为一项；其他单边记录逐条保留。
    for record in sources:
        if record["事项编号"]:
            continue
        chain_ids = {e["业务链ID"] for e in events[_key(record)]}
        members = [record]
        if chain_ids:
            members = [r for r in sources if not r["事项编号"] and any(e["业务链ID"] in chain_ids for e in events[_key(r)])]
        add({}, members)
    for issue in _records(tables, "疑点事项"):
        if not issue.get("匹配ID"):
            amount = float(issue.get("差异金额", 0) or 0)
            add({"出处": "疑点事项（月度累计）"}, [], issue,
                f"{issue.get('月份', '')} 同类差异累计 {amount:,.2f}，达到既有风险处置条件。", "结合相关业务逐项核实，不以不同方向差异抵销。")
    summary = {r["项目"]: r["数值"] for r in _records(tables, "核对结论")}
    if summary.get("核对范围") == "范围受限":
        add({}, [], global_problem="核对范围受限：" + str(summary.get("范围说明", "资料尚未完整")), global_action="先核实账户、币种、期间及资料完整性，再使用受影响的核对结论。")
    for name in ("余额连续性异常", "余额差异明细"):
        for position, problem in enumerate(_records(tables, name), 2):
            description = "；".join(f"{k}：{_text(v)}" for k, v in problem.items() if _text(v) and not str(k).startswith("Unnamed:"))
            add({"出处": f"{name} 第 {position} 行"}, [], global_problem=f"{name}。{description}", global_action="核对相关日期的余额及区间收支，检查是否存在缺失或重复入账。")
    index = pd.DataFrame(index_rows, columns=INDEX_COLUMNS)
    source_table = pd.DataFrame(sources, columns=SOURCE_COLUMNS)
    issue_table = pd.DataFrame([{**{k: r.get(k, "") for k in ISSUE_COLUMNS}, "后续状态": "", "处理说明": ""}
                                for r in index_rows if r["需要复核"] == "是"], columns=ISSUE_COLUMNS)
    result = {"核对概览": build_overview(index, source_table, tables), "待复核事项": issue_table,
              "全部核对明细": pd.DataFrame(details, columns=DETAIL_COLUMNS), "阅读事项索引": index, "阅读来源记录": source_table}
    return result


def build_overview(index, sources, tables, filter_description=""):
    items = index.fillna("").to_dict("records")
    records = sources.fillna("").to_dict("records")
    summary = {r["项目"]: r["数值"] for r in _records(tables, "核对结论")}
    needs = sum(r["需要复核"] == "是" for r in items)
    rows = [("处理结果", f"{'筛选结果已生成' if filter_description else '程序分析已完成'}；有 {needs} 项需要查看" if needs else "筛选结果无待复核事项" if filter_description else "程序分析已完成，本次没有待复核事项", "待复核事项"),
            ("阅读顺序", "先看概览，再查看待复核事项，按事项编号进入完整明细。", "全部核对明细")]
    if filter_description:
        rows += [("筛选条件", filter_description, ""), ("筛选范围说明", "仅展示本次选中的完整业务；资料和余额限制沿用全量报告，单列保留。", "")]
    for side, label in (("银行流水", "银行"), ("日记账", "日记账")):
        selected = [r for r in records if r["来源"] == side]
        accounts = "、".join(sorted({_text(r.get("账户")) for r in selected if _text(r.get("账户"))})) or "未提供账户标识"
        currencies = "、".join(sorted({_text(r.get("币种")) for r in selected if _text(r.get("币种"))})) or "未提供币种"
        rows += [(f"{label}账户及币种", f"{accounts}；{currencies}", ""),
                 (f"{label}记录期间", _period(selected), "全部核对明细"), (f"{label}记录笔数", len(selected), "全部核对明细")]
    rows += [("已核对业务数", sum(r["最终状态"] in {"自动确认", "整组勾稽一致"} and r["全局事项"] != "是" for r in items), "全部核对明细"),
             ("留存备查业务数", sum(r["核对结果"] == "留存备查" for r in items), "全部核对明细"),
             ("待复核事项数", needs, "待复核事项")]
    crosses = sum(r["跨期"] == "是" for r in items)
    if crosses:
        rows.append(("涉及跨期事项数", crosses, "待复核事项"))
    rows += [("核对范围", summary.get("核对范围", "范围未验证"), "待复核事项"),
             ("余额核对", summary.get("余额核对", "余额核对未实施"), "待复核事项")]
    if str(summary.get("组合搜索完整性", "")) not in {"", "完整", "已完成"} and ("未穷尽" in str(summary.get("组合搜索说明", ""))):
        rows.append(("对应关系限制", "部分记录存在多个可能对应或未能检查全部组合；未找到对应不等于不存在对应。", "全部核对明细"))
    rows += [("后续处理（可选）", "无需填写即可使用报告；自愿填写的状态在下方汇总。", "待复核事项"),
             ("已处理事项数", 0, "待复核事项"),
             ("金额及计数说明", "金额正数为收入、负数为支出。业务合计与组成记录不重复加总；已核对业务中仍可能有期间等问题。", ""),
             ("查看技术明细", "Excel 中右击工作表标签，选择取消隐藏，可查评分、运行参数及原始解析证据。", "")]
    return pd.DataFrame(rows, columns=["项目", "结果", "查看位置"])


def apply_readable_presentation(workbook):
    """全量报告及筛选报告共用阅读排版。"""
    if not all(name in workbook.sheetnames for name in MAIN_SHEETS):
        return
    for sheet in workbook:
        sheet.sheet_state = "visible" if sheet.title in MAIN_SHEETS else "hidden"
    for position, name in enumerate(MAIN_SHEETS):
        sheet = workbook[name]
        workbook.move_sheet(sheet, offset=position - workbook.index(sheet))
    workbook.active = 0
    targets = {}
    detail = workbook["全部核对明细"]
    detail_headers = {c.value: c.column for c in detail[1] if c.value}
    for row in range(2, detail.max_row + 1):
        number = detail.cell(row, detail_headers["事项编号"]).value
        kind = detail.cell(row, detail_headers["行别"]).value
        if number and kind != "组成记录":
            targets[str(number)] = row
    for name in MAIN_SHEETS:
        sheet = workbook[name]
        headers = {c.value: c.column for c in sheet[1] if c.value}
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 80 if name == "全部核对明细" else 90
        first = min(headers.values())
        for column in range(1, first):
            sheet.column_dimensions[get_column_letter(column)].hidden = True
        widths = ({"项目": 23, "结果": 80, "查看位置": 20} if name == "核对概览" else
                  {"事项编号": 12, "需要关注的问题": 45, "银行记录": 30, "日记账记录": 30, "核查建议": 34, "后续状态": 13, "处理说明": 30} if name == "待复核事项" else
                  {"事项编号": 12, "行别": 12, "核对结果": 22, "银行日期": 14, "银行摘要": 25, "银行金额": 16,
                   "日记账日期": 14, "日记账摘要": 25, "日记账金额": 16, "差额": 15, "依据及说明": 48, "原始出处": 36})
        for header, column in headers.items():
            sheet.column_dimensions[get_column_letter(column)].width = widths.get(header, 20)
            sheet.column_dimensions[get_column_letter(column)].hidden = False
        sheet.freeze_panes = f"{get_column_letter(first + (3 if name == '全部核对明细' else 1))}2"
        sheet.auto_filter.ref = f"{get_column_letter(first)}1:{get_column_letter(sheet.max_column)}{max(1, sheet.max_row)}" if name != "核对概览" else None
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_properties.outlinePr.summaryBelow = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if name == "全部核对明细" else sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1 if name == "核对概览" else 0
        sheet.page_margins.left = sheet.page_margins.right = 0.3
        sheet.page_margins.top = sheet.page_margins.bottom = 0.4
        sheet.print_title_rows = "1:1"
        sheet.print_options.horizontalCentered = True
        sheet.print_area = f"{get_column_letter(first)}1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        sheet.oddFooter.center.text = "第 &P 页 / 共 &N 页"
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="203D56")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for row in range(2, sheet.max_row + 1):
            height = 22 if name == "核对概览" else 28
            for header, column in headers.items():
                cell = sheet.cell(row, column)
                cell.font = Font(name="微软雅黑", size=11, color="243746")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if header in {"银行金额", "日记账金额", "差额"}:
                    cell.number_format = '#,##0.00;[Red]-#,##0.00;"—"'
                width = widths.get(header, 20)
                lines = sum(max(1, (sum(2 if ord(c) > 127 else 1 for c in line) + width - 1) // width) for line in str(cell.value or "").split("\n"))
                height = max(height, lines * 15 + 8)
                cell.fill = PatternFill("solid", fgColor="FFFFFF" if row % 2 == 0 else "F1F5F8")
            sheet.row_dimensions[row].height = min(409, height)
            if name == "全部核对明细":
                kind = sheet.cell(row, headers["行别"]).value
                if kind == "组成记录":
                    sheet.row_dimensions[row].outlineLevel = 1
                    sheet.row_dimensions[row].hidden = True
                elif kind == "业务合计":
                    sheet.row_dimensions[row].collapsed = True
                    for cell in sheet[row]:
                        cell.fill = PatternFill("solid", fgColor="E0EBF2")
            if name == "待复核事项":
                cell = sheet.cell(row, headers["事项编号"])
                target = targets.get(str(cell.value))
                if target:
                    cell.hyperlink = f"#'全部核对明细'!{get_column_letter(detail_headers['事项编号'])}{target}"
                    cell.font = Font(name="微软雅黑", size=11, color="1565A0", underline="single")
        if sheet.max_row == 1 and name != "核对概览":
            sheet.cell(2, first, "本次没有待复核事项" if name == "待复核事项" else "本次没有选中的业务记录")
            sheet.column_dimensions[get_column_letter(first)].width = 24
        if name == "待复核事项":
            validation = DataValidation(type="list", formula1='"已关注,已处理,无需处理"', allow_blank=True)
            sheet.add_data_validation(validation)
            if sheet.max_row > 1:
                letter = get_column_letter(headers["后续状态"])
                validation.add(f"{letter}2:{letter}{sheet.max_row}")
    overview = workbook["核对概览"]
    headers = {c.value: c.column for c in overview[1] if c.value}
    issue = workbook["待复核事项"]
    issue_headers = {c.value: c.column for c in issue[1] if c.value}
    for row in range(2, overview.max_row + 1):
        link = overview.cell(row, headers["查看位置"])
        if link.value in MAIN_SHEETS:
            link.hyperlink = f"#'{link.value}'!B1"
            link.font = Font(name="微软雅黑", size=11, color="1565A0", underline="single")
        if overview.cell(row, headers["项目"]).value == "已处理事项数":
            letter = get_column_letter(issue_headers["后续状态"])
            overview.cell(row, headers["结果"], f'=COUNTIF(\'待复核事项\'!{letter}2:{letter}{max(2, issue.max_row)},"已处理")')
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
