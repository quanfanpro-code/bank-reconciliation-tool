"""用户在 Excel 保存人工结果后，原最近报告入口仍能打开报告。"""

from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

import gui
from 项目记录 import LocalProjectStore


def test_人工修改报告后最近报告可打开且原指纹仍保留(tmp_path, monkeypatch):
    银行 = tmp_path / '银行.csv'
    账 = tmp_path / '序时账.csv'
    银行.write_text('日期,金额\n2026-01-01,100\n', encoding='utf-8-sig')
    账.write_text('日期,金额\n2026-01-01,100\n', encoding='utf-8-sig')
    报告 = tmp_path / '核对报告.xlsx'
    簿 = Workbook()
    簿.active.title = '人工全查'
    簿.active.append(['事项编号', '人工核对结果'])
    簿.active.append(['R001', ''])
    簿.save(报告)
    存储 = LocalProjectStore(tmp_path / '历史')
    记录 = 存储.save_project(bank_path=银行, journal_path=账, report_path=报告,
                          bank_mapping={}, journal_mapping={}, parameters={}, result_counts={})
    历史路径 = 存储.projects_dir / (记录['project_id'] + '.json')
    原记录 = 历史路径.read_bytes()
    簿 = load_workbook(报告)
    簿['人工全查']['B2'] = '确认对应'
    簿.save(报告)
    with pytest.raises(ValueError, match='内容已经改变'):
        存储.resolve_report(记录['project_id'])
    打开, 报错 = [], []
    monkeypatch.setattr(gui.os, 'startfile', 打开.append)
    monkeypatch.setattr(gui.messagebox, 'showerror', lambda *args, **kwargs: 报错.append(args))

    gui.ReconciliationApp.open_latest_report(SimpleNamespace(project_store=存储))

    assert 打开 == [str(报告)]
    assert 报错 == []
    assert 历史路径.read_bytes() == 原记录
    assert load_workbook(报告)['人工全查']['B2'].value == '确认对应'
    with pytest.raises(ValueError, match='内容已经改变'):
        存储.resolve_report(记录['project_id'])
