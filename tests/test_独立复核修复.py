# -*- coding: utf-8 -*-
"""独立复核发现问题的回归测试：空待查表空行、未达差异误降级、业务链交叉引用、表名错别字。"""
import pandas as pd
import pytest
from openpyxl import load_workbook

from application import run_reconciliation
from data_structures import MatcherConfig, ProcessingStatus


映射 = {
    "date": "日期", "amount": "金额", "summary": "摘要", "mode": "signed_amount",
    "account": "本方账号", "currency": "币种", "amount_basis": "本位币",
    "auxiliary_text_columns": ["摘要", "业务编号"],
}


def _写表(tmp_path, 银行行, 账行):
    银行路径 = tmp_path / "银行.csv"
    账路径 = tmp_path / "账.csv"
    pd.DataFrame(银行行).to_csv(银行路径, index=False, encoding="utf-8-sig")
    pd.DataFrame(账行).to_csv(账路径, index=False, encoding="utf-8-sig")
    输出 = tmp_path / "核对结果.xlsx"
    捕获 = []
    run_reconciliation(
        str(银行路径), str(账路径), 映射, 映射, MatcherConfig(),
        bank_skiprows=0, journal_skiprows=0, bank_header_rows=1, journal_header_rows=1,
        output_path=str(输出), matcher_ready=捕获.append,
    )
    return 输出, 捕获[0]


def _行(日期, 金额, 摘要, 编号):
    return {"日期": 日期, "金额": 金额, "摘要": 摘要, "业务编号": 编号, "本方账号": "A001", "币种": "CNY"}


def test_两侧全部对应时待查表只有表头不多出空行(tmp_path):
    行 = [_行("2026-08-10", 100, "甲公司回款", "P01"), _行("2026-08-11", -200, "付乙公司款", "P02")]
    输出, _ = _写表(tmp_path, 行, [dict(r) for r in 行])

    簿 = load_workbook(输出)
    for 表名 in ("银行侧待查", "日记账侧待查"):
        表 = 簿[表名]
        assert 表.max_row == 1, f"{表名}应只有表头，实际max_row={表.max_row}"


def test_起止日期和合计不一致属正常未达不再整体降级(tmp_path):
    银行 = [
        _行("2026-08-10", 100, "甲公司回款", "P01"),
        _行("2026-08-11", -200, "付乙公司款", "P02"),
        _行("2026-08-25", 50, "丙公司回款", "P03"),
    ]
    账 = [
        _行("2026-08-10", 100, "甲公司回款", "P01"),
        _行("2026-08-11", -200, "付乙公司款", "P02"),
    ]
    _, 匹配器 = _写表(tmp_path, 银行, 账)

    assert 匹配器.overall_control.period_status == "疑点"
    assert 匹配器.overall_scope_limited is False
    已确认 = [c for c in 匹配器.selected_candidates if c.processing_status is ProcessingStatus.AUTO_CONFIRMED]
    assert len(已确认) == 2, "P01、P02应保持自动确认，仅P03留待查"


def test_结构性问题仍然限制范围并降级(tmp_path):
    银行 = [_行("2026-08-10", 100, "甲公司回款", "P01"), _行("2026-08-11", 0, "空摘要", "")]
    账 = [_行("2026-08-10", 100, "甲公司回款", "P01"), _行("2026-08-11", 0, "空摘要", "")]
    银行[1]["金额"] = "无法解析"
    _, 匹配器 = _写表(tmp_path, 银行, 账)

    assert 匹配器.overall_scope_limited is True


def test_待查行交叉引用所属业务链(tmp_path):
    银行 = [
        _行("2026-08-12", -5000, "支付丙公司设备款", "P03"),
        _行("2026-08-13", 5000, "丙公司设备款退回", "P03"),
        _行("2026-08-15", -5000, "重新支付丙公司设备款", "P03"),
    ]
    账 = [_行("2026-08-15", -5000, "付丙公司设备款", "P03")]
    输出, _ = _写表(tmp_path, 银行, 账)

    簿 = load_workbook(输出)
    待查 = 簿["银行侧待查"]
    表头 = [c.value for c in 待查[1]]
    依据列 = 表头.index("判断依据")
    金额列 = 表头.index("金额")
    退回行 = [r for r in 待查.iter_rows(min_row=2, values_only=True) if r[金额列] == 5000]
    assert 退回行, "退回行应保留在银行侧待查"
    assert "业务链" in str(退回行[0][依据列]) and "退款冲销重付" in str(退回行[0][依据列])


def test_报告使用数据入口处置表名(tmp_path):
    行 = [_行("2026-08-10", 100, "甲公司回款", "P01")]
    输出, _ = _写表(tmp_path, 行, [dict(r) for r in 行])

    簿 = load_workbook(输出)
    assert "数据入口处置" in 簿.sheetnames
    assert "数据人口处置" not in 簿.sheetnames


def test_同日行内乱序不再误报余额连续性():
    from decimal import Decimal
    from balance import check_balance_continuity

    frame = pd.DataFrame({
        "date": pd.to_datetime(["2026-01-05", "2026-01-06", "2026-01-06", "2026-01-06", "2026-01-07"]),
        "amount": [Decimal("5000"), Decimal("100"), Decimal("200"), Decimal("300"), Decimal("50")],
        # 01-06 真实过账 +100→5100、+300→5400、+200→5600，文件把后两行写反
        "balance": [Decimal("5000"), Decimal("5100"), Decimal("5600"), Decimal("5400"), Decimal("5650")],
        "original_file_row": [2, 3, 4, 5, 6],
    })
    assert check_balance_continuity(frame, source="银行流水") == []


def test_日级断点仍被检出并定位到具体日期():
    from decimal import Decimal
    from balance import check_balance_continuity

    frame = pd.DataFrame({
        "date": pd.to_datetime(["2026-01-05", "2026-01-06", "2026-01-06", "2026-01-07"]),
        "amount": [Decimal("5000"), Decimal("100"), Decimal("300"), Decimal("50")],
        # 01-06 开盘倒推 4000，与上日末 5000 差 1000（真实缺行）
        "balance": [Decimal("5000"), Decimal("4100"), Decimal("4400"), Decimal("4450")],
        "original_file_row": [2, 3, 4, 5],
    })
    异常 = check_balance_continuity(frame, source="银行流水")
    assert len(异常) == 1
    assert "2026-01-06" in str(异常[0]["日期"])
    assert 异常[0]["差额"] == Decimal("1000")
