"""独立复核：失败、取消和另存不能破坏已有底稿。"""

from hashlib import sha256
from types import SimpleNamespace

import pytest

from gui import ReconciliationApp
from reporter import Reporter
from tests.test_第五阶段交付边界 import _输入, _运行
from tests.test_复核保存筛选 import _构造报告
from tests.test_第三阶段特殊业务 import _df
from matcher import Matcher
from data_structures import MatcherConfig
from 底稿筛选 import FilterCriteria, export_filtered_workpaper
from 项目记录 import LocalProjectStore


def test_正式核对不能覆盖已有人工底稿(tmp_path):
    输入 = _输入(tmp_path)
    输出 = tmp_path / "已有底稿.xlsx"
    输出.write_bytes(b"existing reviewed workbook")
    原文 = 输出.read_bytes()
    with pytest.raises(FileExistsError):
        _运行(*输入, 输出, LocalProjectStore(tmp_path / "历史"))
    assert 输出.read_bytes() == 原文


@pytest.mark.parametrize("已有输出", [False, True])
def test_报告末段失败不能留下半成品或毁掉旧文件(tmp_path, monkeypatch, 已有输出):
    frame = _df([("2026-01-02", 100, {"摘要": "服务款", "业务编号": "P01"})])
    匹配器 = Matcher(frame, frame, MatcherConfig())
    匹配器.run()
    输出 = tmp_path / "报告.xlsx"
    if 已有输出:
        输出.write_bytes(b"existing reviewed workbook")
    def 末段故障(*args):
        raise OSError("模拟末段写入故障")
    monkeypatch.setattr(Reporter, "_apply_report_presentation", 末段故障)
    with pytest.raises(OSError, match="末段"):
        Reporter(匹配器).generate_report(str(输出))
    if 已有输出:
        assert 输出.read_bytes() == b"existing reviewed workbook"
    else:
        assert not 输出.exists()


def test_匹配后报告前取消不交付报告(tmp_path):
    输入 = _输入(tmp_path)
    输出 = tmp_path / "报告.xlsx"
    状态 = {}
    def 进度(value):
        if value == 0.8:
            状态["匹配器"].set_stopping(True)
    with pytest.raises(InterruptedError):
        _运行(*输入, 输出, LocalProjectStore(tmp_path / "历史"),
              matcher_ready=lambda value: 状态.update(匹配器=value), progress_callback=进度)
    assert not 输出.exists()


def test_读取阶段点击停止不应继续生成报告(tmp_path, monkeypatch):
    import gui
    from tests.test_第五阶段交付边界 import 映射
    输入 = _输入(tmp_path)
    日志 = []
    # 使用真实 run_process 和 stop_process，仅替换窗口显示以免阻塞测试。
    app = SimpleNamespace(matcher=None, project_store=LocalProjectStore(tmp_path / "历史"),
                          last_report_path=None, _set_progress=lambda value: None,
                          _set_stop_enabled=lambda value: None, _set_start_enabled=lambda value: None,
                          _confirm_precheck_warnings=lambda value: True)
    def 记录(value):
        日志.append(value)
        if value == "开始读取银行流水和银行存款序时账":
            ReconciliationApp.stop_process(app)
    app.log = 记录
    原哈希 = [sha256(path.read_bytes()).hexdigest() for path in 输入]
    ReconciliationApp.run_process(app, {
        "bank_path": str(输入[0]), "journal_path": str(输入[1]),
        "bank_mapping": 映射, "journal_mapping": 映射, "config": MatcherConfig(),
        "llm_config": None, "bank_skip": 0, "journal_skip": 0,
        "bank_header_rows": 1, "journal_header_rows": 1, "date_format": "auto",
    })
    assert app.last_report_path is None
    assert not list(tmp_path.glob("*核对报告*.xlsx"))
    assert any("取消" in value for value in 日志)
    assert [sha256(path.read_bytes()).hexdigest() for path in 输入] == 原哈希


def test_筛选另存不能覆盖已有底稿(tmp_path):
    来源 = _构造报告(tmp_path)
    输出 = tmp_path / "已有筛选.xlsx"
    输出.write_bytes(b"reviewed filter")
    with pytest.raises(FileExistsError):
        export_filtered_workpaper(来源, 输出, FilterCriteria())
    assert 输出.read_bytes() == b"reviewed filter"


@pytest.mark.parametrize("比例", [float("nan"), -0.1, 1.1])
def test_筛选比例越界不能悄悄变成全选或空选(tmp_path, 比例):
    来源 = _构造报告(tmp_path)
    输出 = tmp_path / "筛选.xlsx"
    with pytest.raises(ValueError):
        export_filtered_workpaper(来源, 输出, FilterCriteria(coverage_ratio=比例))
    assert not 输出.exists()


def test_报告写入期间取消不交付正式文件(tmp_path):
    输入 = _输入(tmp_path)
    输出 = tmp_path / "报告.xlsx"
    状态 = {"取消": False}
    def 日志(value):
        if value == "开始应用审计报告排版和复核标记":
            状态["取消"] = True
    with pytest.raises(InterruptedError):
        _运行(*输入, 输出, LocalProjectStore(tmp_path / "历史"), logger=日志,
              cancel_requested=lambda: 状态["取消"])
    assert not 输出.exists()
    assert not (tmp_path / "历史").exists()


def test_筛选保存中途失败不留下正式半成品(tmp_path, monkeypatch):
    from pathlib import Path
    from openpyxl import Workbook
    来源 = _构造报告(tmp_path)
    原哈希 = sha256(来源.read_bytes()).digest()
    输出 = tmp_path / "筛选.xlsx"
    def 保存故障(self, path):
        Path(path).write_bytes(b"partial workbook")
        raise OSError("模拟磁盘写入失败")
    monkeypatch.setattr(Workbook, "save", 保存故障)
    with pytest.raises(OSError, match="磁盘"):
        export_filtered_workpaper(来源, 输出, FilterCriteria())
    assert not 输出.exists()
    assert sha256(来源.read_bytes()).digest() == 原哈希


def test_任务运行期间出现同名底稿也不能覆盖(tmp_path):
    输入 = _输入(tmp_path)
    输出 = tmp_path / "报告.xlsx"
    def 日志(value):
        if value == "开始应用审计报告排版和复核标记":
            输出.write_bytes(b"another completed workpaper")
    with pytest.raises(FileExistsError):
        _运行(*输入, 输出, LocalProjectStore(tmp_path / "历史"), logger=日志)
    assert 输出.read_bytes() == b"another completed workpaper"


@pytest.mark.parametrize("列名,原值", [("原日期", 20260105), ("年份", 2026), ("月份", 1)])
def test_日期列的原始数字不能被排版改成错误日期(tmp_path, 列名, 原值):
    from contextlib import closing
    import pandas as pd
    from openpyxl import load_workbook
    from make_excel import make_excel, beautify
    输出 = tmp_path / "原值.xlsx"
    make_excel(pd.DataFrame({列名: [原值]}), str(输出))
    with closing(load_workbook(输出)) as book:
        assert book.active["B2"].value == 原值
        assert book.active["B2"].data_type == "n"
    美化后 = tmp_path / "美化.xlsx"
    beautify(str(输出), str(美化后))
    with closing(load_workbook(美化后)) as book:
        assert book.active["B2"].value == 原值


def test_真实日期对象仍按日期展示(tmp_path):
    from contextlib import closing
    from datetime import datetime
    import pandas as pd
    from openpyxl import load_workbook
    from make_excel import make_excel
    输出 = tmp_path / "真实日期.xlsx"
    原日期 = datetime(2026, 1, 5)
    make_excel(pd.DataFrame({"日期": [原日期]}), str(输出))
    with closing(load_workbook(输出)) as book:
        assert book.active["B2"].value == 原日期
        assert book.active["B2"].is_date
