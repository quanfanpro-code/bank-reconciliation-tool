"""筛选金额含义、完整组成、列偏好与真实界面验收。"""
from contextlib import closing
from hashlib import sha256
from decimal import Decimal

import pandas as pd
import pytest
from openpyxl import load_workbook

from 底稿筛选 import FilterCriteria, _filter_groups, export_filtered_workpaper
from tests.test_人工结果计算 import _报告, excel, _打开, _选择, _表


def _工资():
    groups = pd.DataFrame([{'事项编号':'工资','匹配ID':'工资','组金额':800000}, {'事项编号':'银行独有','匹配ID':'银行独有','组金额':20000}])
    details = pd.DataFrame([{'事项编号':'工资','来源':'银行流水','金额':-8000,'原始记录号':str(i)} for i in range(100)] + [{'事项编号':'工资','来源':'序时账','金额':-800000,'原始记录号':'账1'}, {'事项编号':'银行独有','来源':'银行流水','金额':-20000,'原始记录号':'单边'}])
    return groups, details


@pytest.mark.parametrize('basis,expected', [('银行单笔',['银行独有']), ('序时账单笔',['工资']), ('整组金额',['工资','银行独有'])])
def test_工资银行单笔与账面单笔和整组不是同一口径(basis, expected):
    groups, details = _工资()
    chosen = _filter_groups(groups, FilterCriteria(amount_basis=basis,min_amount=10000), details)
    assert chosen['事项编号'].tolist() == expected


def test_金额区间含边界并支持统一带符号值():
    groups, details = _工资()
    assert _filter_groups(groups,FilterCriteria(amount_basis='银行单笔',min_amount=-20000,max_amount=-20000,amount_absolute=False),details)['事项编号'].tolist() == ['银行独有']
    assert _filter_groups(groups,FilterCriteria(amount_basis='银行单笔',min_amount=20000,max_amount=20000),details)['事项编号'].tolist() == ['银行独有']


@pytest.mark.parametrize('values',[{'min_amount':'NaN'}, {'min_amount':'Infinity'}, {'min_amount':20,'max_amount':10}, {'coverage_ratio':1.1}, {'start_date':'2026-02-30'}, {'start_date':'2026-02-02','end_date':'2026-01-01'}])
def test_非法条件不能被静默接受(values):
    with pytest.raises(ValueError):
        _filter_groups(*[pd.DataFrame(),FilterCriteria(amount_basis='整组金额',**values)])


def test_单笔筛选缺少组成时明确说明():
    groups,_ = _工资()
    with pytest.raises(ValueError,match='明细|组成'):
        _filter_groups(groups,FilterCriteria(amount_basis='银行单笔',min_amount=10))


def test_单笔命中后整组完整且原报告不变(tmp_path):
    source=_报告(tmp_path,[8000,8000],[16000],[((0,1),(0,))])
    before=sha256(source.read_bytes()).hexdigest()
    output=tmp_path/'账面大额.xlsx'
    export_filtered_workpaper(source,output,FilterCriteria(amount_basis='序时账单笔',min_amount=16000))
    assert sha256(source.read_bytes()).hexdigest()==before
    with closing(load_workbook(output)) as book:
        sheet=book['整组核对']
        rows=[tuple(c.value for c in sheet[r]) for r in range(2,sheet.max_row+1) if not sheet.row_dimensions[r].hidden]
        assert len(rows)==4
        text=' '.join(str(c.value) for row in book['筛选说明'] for c in row)
        assert '序时账单笔' in text and '16000' in text


def test_列偏好重启读取且保留所有身份列(tmp_path):
    from 报告偏好 import save_column_preferences, load_column_preferences, apply_column_preferences
    path=tmp_path/'偏好.json'
    save_column_preferences({'逐笔核对':['银行金额','银行记录','银行金额','已废弃']},path)
    tables={'逐笔核对':pd.DataFrame([{'核对编号':1,'银行记录':'甲','银行金额':100,'事项编号':'M1','新增列':'保留'}])}
    apply_column_preferences(tables,load_column_preferences(path))
    assert list(tables['逐笔核对'])==['核对编号','银行金额','银行记录','事项编号','新增列']
    assert tables['逐笔核对'].iloc[0].to_dict()=={'核对编号':1,'银行记录':'甲','银行金额':100,'事项编号':'M1','新增列':'保留'}


def test_列重排后Excel人工选择仍生效(excel,tmp_path,monkeypatch):
    from 报告偏好 import save_column_preferences
    monkeypatch.setenv('LOCALAPPDATA',str(tmp_path))
    save_column_preferences({'人工全查':['人工核对结果','备注','序时账记录','银行记录'], '逐笔核对':['当前核对结论','银行金额','序时账金额']})
    source=_报告(tmp_path,[100],[100],[((0,),(0,))])
    with closing(load_workbook(source)) as book:
        assert [c.value for c in book['人工全查'][1] if c.value is not None][:3]==['核对编号','人工核对结果','备注']
    with _打开(excel,source) as book:
        _选择(excel,book,'M1','否定对应')
        assert '否定' in _表(book,'逐笔核对')[0]['当前核对结论']


def test_集中窗口收集条件及取消不返回结果(tmp_path):
    import tkinter as tk
    from 筛选窗口 import FilterDialog
    root=tk.Tk();root.withdraw()
    try:
        dialog=FilterDialog(root, preferences_path=tmp_path/'偏好.json')
        dialog.variables['amount_basis'].set('银行单笔')
        dialog.variables['min_amount'].set('10000')
        dialog.variables['include_text'].set('工资;奖金；报销')
        value=dialog.build_criteria()
        assert value.min_amount==Decimal('10000')
        assert value.include_text==('工资','奖金','报销')
        dialog.cancel()
        assert dialog.result is None
    finally:
        root.destroy()
