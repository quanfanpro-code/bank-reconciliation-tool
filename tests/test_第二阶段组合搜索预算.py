"""组合搜索预算及验收报告披露的回归测试。"""

from decimal import Decimal

import pandas as pd
from openpyxl import load_workbook

from data_structures import MatcherConfig
from matcher import Matcher
import matcher as 匹配模块
from precision_engine import PrecisionEngine
from reporter import Reporter


def _数据(金额列表, 来源, 摘要="", 对方=""):
    return pd.DataFrame(
        [
            {
                "date": pd.Timestamp("2026-08-31"),
                "amount": Decimal(str(金额)),
                "amount_decimal": PrecisionEngine.to_integer_li(金额),
                "summary": 摘要,
                "aux_text_fields": {"摘要": 摘要, "对方户名": 对方},
                "balance": None,
                "voucher_no": "",
                "source": 来源,
                "original_idx": 序号,
                "original_file_row": 序号 + 2,
            }
            for 序号, 金额 in enumerate(金额列表)
        ]
    )


def _运行(银行金额, 账簿金额, **配置项):
    配置 = MatcherConfig(
        allow_greedy_fallback=False,
        clearly_trivial_threshold=Decimal("0"),
        **配置项,
    )
    实例 = Matcher(_数据(银行金额, "bank"), _数据(账簿金额, "journal"), 配置)
    实例.match_dfs_combinations()
    return 实例


def _工作表键值(工作表, 键列名, 值列名):
    表头 = {单元格.value: 单元格.column for 单元格 in 工作表[1]}
    return {
        工作表.cell(行, 表头[键列名]).value: 工作表.cell(行, 表头[值列名]).value
        for 行 in range(2, 工作表.max_row + 1)
    }


def test_异号过滤后不足两笔不计候选截断或预算耗尽():
    实例 = _运行([100], [50, *([-1] * 30)])

    搜索 = 实例.run_parameters["combination_search"]
    assert 搜索["generic_source_rows"] == 0
    assert 搜索["truncated_source_rows"] == 0
    assert 搜索["budget_exhausted_source_rows"] == 0


def test_动态深度漏搜时明确记录深度受限():
    实例 = _运行(
        [100],
        [1000, 20, 20, 20, 20, 20],
        combination_node_limit_per_source=100_000,
    )

    搜索 = 实例.run_parameters["combination_search"]
    assert 搜索["generic_source_rows"] == 1
    assert 搜索["depth_limited_source_rows"] == 1
    assert 搜索["fully_searched_source_rows"] == 0
    assert 搜索["budget_exhausted_source_rows"] == 1


def test_同一来源同时触发深度和节点限制时耗尽来源只计一次():
    实例 = _运行(
        [100],
        [1000, 20, 20, 20, 20, 20],
        combination_node_limit_per_source=1,
    )

    搜索 = 实例.run_parameters["combination_search"]
    assert 搜索["depth_limited_source_rows"] == 1
    assert 搜索["node_budget_exhausted_source_rows"] == 1
    assert 搜索["budget_exhausted_source_rows"] == 1


def test_单线程组合搜索受单任务时间限制():
    实例 = _运行(
        [101],
        [50, 40, 30, 20],
        combination_task_timeout_seconds=0,
        combination_global_time_limit_seconds=10,
    )

    搜索 = 实例.run_parameters["combination_search"]
    assert 搜索["task_timeout_source_rows"] == 1
    assert 搜索["global_timeout_unprocessed_source_rows"] == 0
    assert 搜索["budget_exhausted_source_rows"] == 1


def test_单线程组合搜索受全局时间限制且未处理来源留痕():
    实例 = _运行(
        [101],
        [50, 40, 30, 20],
        combination_task_timeout_seconds=10,
        combination_global_time_limit_seconds=0,
    )

    搜索 = 实例.run_parameters["combination_search"]
    # 零全局预算覆盖双向来源：银行1笔、日记账4笔，去重后均计5笔。
    assert 搜索["global_timeout_unprocessed_source_rows"] == 5
    assert 搜索["fully_searched_source_rows"] == 0
    assert 搜索["budget_exhausted_source_rows"] == 5


def test_验收Excel在可见结论和隐藏运行参数逐项披露搜索预算(tmp_path):
    配置 = MatcherConfig(
        combination_node_limit_per_source=1234,
        combination_task_timeout_seconds=12,
        combination_global_time_limit_seconds=34,
    )
    实例 = Matcher(_数据([100], "bank"), _数据([100], "journal"), 配置)
    实例.run_parameters["combination_search"] = {
        "business_group_bank_rows": 0,
        "business_group_journal_rows": 0,
        "generic_source_rows": 5,
        "fully_searched_source_rows": 1,
        "truncated_source_rows": 1,
        "depth_limited_source_rows": 1,
        "node_budget_exhausted_source_rows": 1,
        "task_timeout_source_rows": 1,
        "global_timeout_unprocessed_source_rows": 1,
        "budget_exhausted_source_rows": 4,
        "candidate_limit": 30,
        "search_budget": {
            "generic_candidate_limit": 30,
            "max_depth": 30,
            "exact_split_max_depth": 8,
            "node_limit_per_source": 1234,
            "task_timeout_seconds": 12,
            "global_time_limit_seconds": 34,
        },
    }
    实例.run_parameters["candidate_search"] = {
        "whitelist": {"examined": 20, "retained": 5, "truncated_source_rows": 2},
        "exact": {"examined": 30, "retained": 6, "truncated_source_rows": 3},
        "tolerance": {"examined": 40, "retained": 7, "truncated_source_rows": 4},
        "atomic_voucher": {"examined": 8, "retained": 5, "truncated_source_rows": 1},
        "generic_combination": {"examined": 50, "retained": 10, "truncated_source_rows": 1},
    }
    输出 = tmp_path / "组合搜索预算披露.xlsx"

    Reporter(实例).generate_report(str(输出), config=配置)

    工作簿 = load_workbook(输出, data_only=False)
    参数表 = 工作簿["运行参数"]
    结论表 = 工作簿["核对结论"]
    assert 参数表.sheet_state == "hidden"
    assert 结论表.sheet_state == "hidden"
    assert 工作簿["核对概览"].sheet_state == "visible"
    assert any("未能检查全部组合" in str(cell.value) for row in 工作簿["核对概览"] for cell in row)

    参数 = _工作表键值(参数表, "参数名称", "参数值")
    assert 参数["组合搜索每来源节点上限"] == 1234
    assert 参数["组合搜索单任务时间上限（秒）"] == 12
    assert 参数["组合搜索全局时间上限（秒）"] == 34
    assert 参数["组合搜索深度受限来源笔数"] == 1
    assert 参数["组合搜索节点预算耗尽来源笔数"] == 1
    assert 参数["组合搜索单任务超时来源笔数"] == 1
    assert 参数["组合搜索全局超时未处理来源笔数"] == 1
    assert 参数["组合搜索预算耗尽来源笔数"] == 4
    assert 参数["精确候选检查数"] == 30
    assert 参数["精确候选保留数"] == 6
    assert 参数["精确候选截断来源笔数"] == 3
    assert 参数["容差候选截断来源笔数"] == 4
    assert 参数["白名单候选截断来源笔数"] == 2

    结论 = _工作表键值(结论表, "项目", "数值")
    assert 结论["组合搜索完整性"] == "范围受限"
    assert "预算内完整搜索1笔" in 结论["组合搜索覆盖"]
    assert "候选截断1笔" in 结论["组合搜索受限原因"]
    assert "节点预算1笔" in 结论["组合搜索受限原因"]
    assert "精确截断3笔" in 结论["组合搜索受限原因"]
    assert "容差截断4笔" in 结论["组合搜索受限原因"]
    assert "未穷尽" in 结论["组合搜索说明"]
    assert "未找到对应不代表不存在组合" in 结论["组合搜索说明"]


def test_无耗尽时Excel明确显示通用来源均在预算内完成(tmp_path):
    实例 = _运行(
        [100],
        [60, 40],
        combination_task_timeout_seconds=5,
        combination_global_time_limit_seconds=10,
    )
    搜索 = 实例.run_parameters["combination_search"]
    assert 搜索["generic_source_rows"] == 1
    assert 搜索["fully_searched_source_rows"] == 1
    assert 搜索["budget_exhausted_source_rows"] == 0

    输出 = tmp_path / "组合搜索预算内完成.xlsx"
    Reporter(实例).generate_report(str(输出), config=实例.config)
    结论 = _工作表键值(load_workbook(输出)["核对结论"], "项目", "数值")
    assert 结论["组合搜索完整性"] == "预算内完成"
    assert 结论["组合搜索说明"] == "通用来源均在所列预算内完成。"


def test_员工报销无共同业务边界不能由金额多集回退成整组一致():
    配置 = MatcherConfig()
    实例 = Matcher(
        _数据([400, 600, 400, 600], "bank", "员工报销"),
        _数据([400, 600, 400, 600], "journal", "员工报销"),
        配置,
    )

    实例.run()

    assert not any(
        len(候选.bank_idxs) == len(候选.journal_idxs) == 4
        and 候选.processing_status.value in {"自动确认", "整组勾稽一致"}
        for 候选 in 实例.selected_candidates
    )
    assert any(
        候选.processing_status.value == "疑点事项" or 候选.is_ambiguous
        for 候选 in 实例.selected_candidates
    )


def test_仅相同摘要无对方或编号时综合可信度不能满分():
    实例 = Matcher(
        _数据([60000], "bank", "设备采购"),
        _数据([60000], "journal", "设备采购"),
        MatcherConfig(auto_confirm_score=100),
    )

    实例.run()

    候选 = 实例.selected_candidates[0]
    assert 候选.scores.total < 100
    assert 候选.evidence["score_limited_by_business_evidence"] is True
    assert 候选.processing_status.value == "疑点事项"


def test_两边二乘二同额且明确同一对方仍可整组一致():
    配置 = MatcherConfig()
    实例 = Matcher(
        _数据([500, 500], "bank", "项目结算", "甲公司"),
        _数据([500, 500], "journal", "项目结算", "甲公司"),
        配置,
    )

    实例.run()

    assert any(
        len(候选.bank_idxs) == len(候选.journal_idxs) == 2
        and 候选.match_type == "closed_candidate_group"
        and 候选.processing_status.value == "整组勾稽一致"
        for 候选 in 实例.selected_candidates
    )
    结论 = dict(
        zip(
            Reporter(实例).build_report_tables(实例.config)["核对结论"]["项目"],
            Reporter(实例).build_report_tables(实例.config)["核对结论"]["数值"],
        )
    )
    assert "未穷尽" in 结论["组合搜索说明"]


def test_余额连续性Excel逐原行披露同日中间余额异常(tmp_path):
    银行 = _数据([100, 100, 100], "bank")
    银行["balance"] = [Decimal("1100"), Decimal("1250"), Decimal("1300")]
    银行["original_file_row"] = [12, 13, 14]
    账簿 = _数据([100], "journal")
    实例 = Matcher(银行, 账簿, MatcherConfig())
    输出 = tmp_path / "同日中间余额异常.xlsx"

    Reporter(实例).generate_report(str(输出), config=实例.config)

    工作表 = load_workbook(输出, data_only=True)["余额连续性异常"]
    表头 = {单元格.value: 单元格.column for 单元格 in 工作表[1]}
    数据行 = list(工作表.iter_rows(min_row=2, values_only=True))
    assert len(数据行) == 2
    assert [行[表头["差额"] - 1] for 行 in 数据行] == [50, 50]
    assert [行[表头["原文件行号"] - 1] for 行 in 数据行] == [13, 14]


def test_精确高密度候选先限池再评分且披露截断(monkeypatch):
    数量, 上限 = 80, 5
    实例 = Matcher(
        _数据([100] * 数量, "bank", "项目结算"),
        _数据([100] * 数量, "journal", "项目结算"),
        MatcherConfig(max_candidates=上限),
    )
    原评分, 次数 = 匹配模块.score_text_fields, 0

    def 计数评分(*参数, **关键字参数):
        nonlocal 次数
        次数 += 1
        return 原评分(*参数, **关键字参数)

    monkeypatch.setattr(匹配模块, "score_text_fields", 计数评分)
    实例.match_exact_1to1()

    assert 次数 <= 数量 * 上限
    统计 = 实例.run_parameters["candidate_search"]["exact"]
    assert 统计 == {
        "examined": 数量 * 数量,
        "retained": 数量 * 上限,
        "truncated_source_rows": 数量,
    }


def test_容差高密度候选先限池再评分且披露截断(monkeypatch):
    数量, 上限 = 80, 5
    银行 = _数据([100] * 数量, "bank", "项目结算")
    账簿 = _数据([101] * 数量, "journal", "项目结算")
    实例 = Matcher(
        银行,
        账簿,
        MatcherConfig(
            max_candidates=上限,
            clearly_trivial_threshold=Decimal("2"),
        ),
    )
    原评分, 次数 = 匹配模块.score_text_fields, 0

    def 计数评分(*参数, **关键字参数):
        nonlocal 次数
        次数 += 1
        return 原评分(*参数, **关键字参数)

    monkeypatch.setattr(匹配模块, "score_text_fields", 计数评分)
    实例.match_tolerance()

    assert 次数 <= 数量 * 上限
    统计 = 实例.run_parameters["candidate_search"]["tolerance"]
    assert 统计 == {
        "examined": 数量 * 数量,
        "retained": 数量 * 上限,
        "truncated_source_rows": 数量,
    }


def test_白名单高密度候选先限池再评分且披露截断(monkeypatch):
    数量, 上限 = 40, 4
    实例 = Matcher(
        _数据([100] * 数量, "bank", "银行手续费"),
        _数据([100] * 数量, "journal", "银行手续费"),
        MatcherConfig(max_candidates=上限),
    )
    原评分, 次数 = 匹配模块.score_text_fields, 0

    def 计数评分(*参数, **关键字参数):
        nonlocal 次数
        次数 += 1
        return 原评分(*参数, **关键字参数)

    monkeypatch.setattr(匹配模块, "score_text_fields", 计数评分)
    实例.match_whitelist_rules()

    assert 次数 <= 数量 * 上限
    统计 = 实例.run_parameters["candidate_search"]["whitelist"]
    assert 统计["examined"] == 数量 * 数量
    assert 统计["retained"] == 数量 * 上限
    assert 统计["truncated_source_rows"] == 数量


def test_完整凭证候选只扫描日期方向金额窗口():
    银行 = _数据([100] * 1000, "bank", "项目结算", "甲公司")
    银行["date"] = pd.date_range("2025-01-01", periods=len(银行), freq="D")
    凭证日 = 银行.loc[500, "date"]
    账簿 = _数据([40, 60], "journal", "项目结算", "甲公司")
    账簿["date"] = [凭证日, 凭证日]
    账簿["voucher_no"] = ["记001", "记001"]
    实例 = Matcher(
        银行,
        账簿,
        MatcherConfig(
            dfs_date_window=3,
            max_candidates=5,
            performance_materiality=Decimal("1"),
        ),
    )

    class 计数字典(dict):
        def __init__(self, 数据):
            super().__init__(数据)
            self.读取次数 = 0

        def __getitem__(self, 键):
            self.读取次数 += 1
            return super().__getitem__(键)

    银行业务行 = 计数字典(实例._business_rows["bank"])
    实例._business_rows["bank"] = 银行业务行
    期间 = pd.Timestamp(凭证日).strftime("%Y-%m")

    实例._add_atomic_journal_voucher_candidates(
        [(('凭证', 期间, '', '记001', 1), (0, 1))]
    )

    assert 银行业务行.读取次数 <= 7
    统计 = 实例.run_parameters["candidate_search"]["atomic_voucher"]
    assert 统计["examined"] <= 7
