"""业务依据、竞争对应和报告统计的对外行为。"""

from decimal import Decimal
from dataclasses import replace

import pandas as pd
import pytest
from openpyxl import load_workbook

from data_structures import MatchCandidate, MatcherConfig, ProcessingStatus
from gui import auto_select_auxiliary_columns
from matcher import Matcher
from matching_policy import build_group_metrics
from precision_engine import PrecisionEngine
from reporter import Reporter


def _数据(amounts, source):
    return pd.DataFrame([
        {
            "date": pd.Timestamp("2026-08-31"),
            "amount": Decimal(str(amount)),
            "amount_decimal": PrecisionEngine.to_integer_li(amount),
            "summary": "8月工资",
            "aux_text_fields": {"批次号": "000123456789012345"},
            "balance": None,
            "voucher_no": "",
            "source": source,
            "original_idx": index,
            "original_file_row": index + 12,
        }
        for index, amount in enumerate(amounts)
    ])


def _报告():
    matcher = Matcher(_数据([100, 900], "bank"), _数据([100, 400, 500], "journal"), MatcherConfig())

    def candidate(name, bank_idxs, journal_idxs, status):
        return MatchCandidate(
            candidate_id=name,
            bank_idxs=bank_idxs,
            journal_idxs=journal_idxs,
            match_type="business_group",
            match_stage="业务完整组",
            metrics=build_group_metrics(
                [int(matcher.bank.loc[i, "amount_decimal"]) for i in bank_idxs],
                [int(matcher.journal.loc[i, "amount_decimal"]) for i in journal_idxs],
            ),
            processing_status=status,
            final_match_id=name,
            evidence={"business_basis": "双方批次号000123456789012345一致"},
        )

    confirmed = candidate("C1", (0,), (0,), ProcessingStatus.AUTO_CONFIRMED)
    flagged = candidate("C2", (1,), (1, 2), ProcessingStatus.FLAGGED)
    alternative = candidate("C3", (1,), (0, 2), ProcessingStatus.FLAGGED)
    flagged.evidence["alternative_candidate_ids"] = ["C3"]
    matcher.candidates = [confirmed, flagged, alternative]
    matcher.selected_candidates = [confirmed, flagged]
    return Reporter(matcher)


def test_自动保留业务编号和批次列():
    columns = ["摘要", "业务编号", "付款批次号", "批次编号", "交易流水号", "回单号", "对方账号", "卡号", "凭证号"]
    assert auto_select_auxiliary_columns(columns) == columns[:6]


def test_报告公开业务依据和竞争候选的原文件组成(tmp_path):
    output = tmp_path / "业务依据报告.xlsx"
    _报告().generate_report(str(output), config=MatcherConfig())
    workbook = load_workbook(output)
    sheet = workbook["整组勾稽"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert "业务分组依据" in headers
    assert "其他可能对应" in headers
    basis = sheet.cell(2, headers["业务分组依据"]).value
    alternatives = sheet.cell(2, headers["其他可能对应"]).value
    assert "000123456789012345" in basis
    assert "银行原文件行：13" in alternatives
    assert "序时账原文件行：12、14" in alternatives
    assert not sheet.column_dimensions[sheet.cell(1, headers["业务分组依据"]).column_letter].hidden
    assert not sheet.column_dimensions[sheet.cell(1, headers["其他可能对应"]).column_letter].hidden


def test_零差额疑点不得计入核对一致覆盖率():
    table = _报告().build_report_tables(MatcherConfig())["核对结论"]
    summary = dict(zip(table["项目"], table["数值"]))
    assert summary["组级勾稽率"] == 0
    assert summary["银行已核对笔数覆盖率"] == pytest.approx(1 / 2)
    assert summary["序时账已核对笔数覆盖率"] == pytest.approx(1 / 3)
    assert summary["银行已核对金额覆盖率"] == pytest.approx(100 / 1000)
    assert summary["序时账已核对金额覆盖率"] == pytest.approx(100 / 1000)
    assert "疑点" in summary["自动完成率说明"]
    assert "绝对值" in summary["金额覆盖率口径"]


def test_旧结果没有业务依据字段仍可导出():
    reporter = _报告()
    for candidate in reporter.matcher.selected_candidates:
        candidate.evidence = {}
    tables = reporter.build_report_tables(MatcherConfig())
    assert tables["整组勾稽"].iloc[0]["业务分组依据"] == ""
    assert tables["整组勾稽"].iloc[0]["其他可能对应"] == ""


def test_大量竞争关系逐行保留不受单元格长度截断(tmp_path):
    reporter = _报告()
    reporter.matcher.journal = Matcher(_数据([100, 900], "bank"), _数据([100] * 1801, "journal"), MatcherConfig()).journal
    selected = reporter.matcher.selected_candidates[0]
    selected.processing_status = ProcessingStatus.FLAGGED
    alternatives = [replace(selected, candidate_id=f"A{i}", journal_idxs=(i,), evidence={}) for i in range(1, 1801)]
    selected.evidence["alternative_candidate_ids"] = [c.candidate_id for c in alternatives]
    reporter.matcher.candidates = [selected, *alternatives]
    reporter.matcher.selected_candidates = [selected]
    output = tmp_path / "大量竞争关系.xlsx"
    reporter.generate_report(str(output), config=MatcherConfig())
    details = pd.read_excel(output, sheet_name="其他可能对应明细")
    assert len(details) == 3600
    assert details["对应方案"].nunique() == 1800
    ledger_rows = details.loc[details["来源"] == "银行存款序时账", "原文件行号"]
    assert set(ledger_rows) == set(range(13, 1813))


def test_旧匹配未保存核对状态时覆盖率明确不可计算():
    reporter = _报告()
    reporter.matcher.selected_candidates = []
    for frame in (reporter.matcher.bank, reporter.matcher.journal):
        frame.loc[0, "matched"] = True
        frame.loc[0, "match_id"] = "旧匹配1"
        frame.loc[0, "match_type"] = "exact_1to1"
    summary = reporter.build_report_tables(MatcherConfig())["核对结论"]
    values = dict(zip(summary["项目"], summary["数值"]))
    assert "无法计算" in str(values["银行已核对笔数覆盖率"])
    assert "无法计算" in str(values["银行已核对金额覆盖率"])


def test_待查原始记录同时保留真实Excel行号():
    reporter = _报告()
    reporter.raw_bank = pd.DataFrame({"交易金额": [100, 900], "备注": ["第一条", "第二条"]})
    reporter.matcher.bank["original_idx"] = [1, 2]
    reporter.matcher.bank.loc[0, "matched"] = True
    unmatched = reporter.build_report_tables(MatcherConfig())["银行侧待查"]
    assert unmatched["原文件行号"].tolist() == [13]
    assert unmatched["交易金额"].tolist() == [900]
    assert unmatched["备注"].tolist() == ["第二条"]


def test_报告披露本次组合搜索覆盖与截断():
    reporter = _报告()
    reporter.matcher.run_parameters["combination_search"] = {
        "business_group_bank_rows": 35, "business_group_journal_rows": 1,
        "generic_source_rows": 9, "truncated_source_rows": 3, "candidate_limit": 30,
    }
    tables = reporter.build_report_tables(MatcherConfig())
    parameters = dict(zip(tables["运行参数"]["参数名称"], tables["运行参数"]["参数值"]))
    assert parameters["组合候选发生截断的来源笔数"] == 3
    summary = dict(zip(tables["核对结论"]["项目"], tables["核对结论"]["数值"]))
    assert "3" in summary["组合搜索说明"]
    assert "穷尽" in summary["组合搜索说明"]
