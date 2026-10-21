# -*- coding: utf-8 -*-
import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill


def _小计划():
    from 排版计划 import 簿计划, 表计划, 列计划, 行计划, 格样式

    表头样式 = 格样式(字体名="微软雅黑", 字号=10, 加粗=True, 填充色="FF1F4E79",
                    填充类型="solid", 水平="center", 垂直="center")
    df = pd.DataFrame({"名称": ["甲", "乙"], "金额": [1.5, "=B2*2"], "备注": ["批注格", None]})
    表 = 表计划(
        名称="测试表",
        表=df,
        表头样式=表头样式,
        列=[列计划(宽=20), 列计划(宽=12), 列计划(隐藏=True)],
        行={2: 行计划(高=30, outline=1), 3: 行计划(隐藏=True)},
        覆盖={(2, 2): 格样式(填充色="FFFFF2CC", 填充类型="solid", 数字格式="#,##0.00"),
              (2, 3): 格样式(批注=("提示", "核对"), 超链接="#测试表!A1")},
        验证=[{"type": "list", "formula1": '"是,否"', "sqref": "C2", "allow_blank": True}],
        条件格式=[("B2:B3", FormulaRule(formula=["B2>1"], fill=PatternFill("solid", fgColor="FFFF0000")))],
        冻结="B2",
        筛选="A1:C3",
        tab色="FF0000",
        网格线=False,
        缩放=90,
        打印={"orientation": "landscape", "paperSize": 9, "fitToWidth": 1,
              "fitToHeight": 0, "fitToPage": True, "title_rows": "1:1",
              "左边距": 0.5, "右边距": 0.5},
    )
    隐藏表 = 表计划(名称="技术表", 表=pd.DataFrame({"x": [1]}), 状态="hidden")
    return 簿计划(表=[隐藏表, 表], active=1, 自动重算=True)


def test_渲染簿计划(tmp_path):
    from 报告写入 import 渲染簿计划

    输出 = tmp_path / "出.xlsx"
    渲染簿计划(_小计划(), str(输出))
    wb = load_workbook(输出)
    assert wb.sheetnames == ["技术表", "测试表"]
    assert wb.active == wb["测试表"]
    assert wb["技术表"].sheet_state == "hidden"
    assert wb.calculation.fullCalcOnLoad

    ws = wb["测试表"]
    assert ws.sheet_properties.tabColor.rgb == "00FF0000" or ws.sheet_properties.tabColor.rgb == "FF0000"
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref == "A1:C3"
    assert ws.sheet_view.showGridLines is False
    assert ws.sheet_view.zoomScale == 90

    # 表头样式
    h = ws["A1"]
    assert h.value == "名称"
    assert h.font.name == "微软雅黑" and h.font.bold
    assert h.fill.fill_type == "solid" and h.fill.fgColor.rgb == "FF1F4E79"
    assert h.alignment.horizontal == "center"

    # 值与公式
    assert ws["A2"].value == "甲"
    assert ws["B2"].value == 1.5
    assert ws["B3"].value == "=B2*2"

    # 列
    assert ws.column_dimensions["A"].width == 20
    assert ws.column_dimensions["C"].hidden

    # 行
    assert ws.row_dimensions[2].height == 30
    assert ws.row_dimensions[2].outlineLevel == 1
    assert ws.row_dimensions[3].hidden

    # 格覆盖
    b2 = ws["B2"]
    assert b2.fill.fgColor.rgb == "FFFFF2CC"
    assert b2.number_format == "#,##0.00"
    c2 = ws["C2"]
    assert c2.comment.text == "提示"
    assert c2.hyperlink.target == "#测试表!A1"

    # 数据验证与条件格式
    dvs = list(ws.data_validations.dataValidation)
    assert len(dvs) == 1 and dvs[0].type == "list" and "C2" in str(dvs[0].sqref)
    assert len(ws.conditional_formatting._cf_rules) == 1

    # 打印
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1
    assert ws.print_title_rows == "$1:$1"
    assert ws.page_margins.left == 0.5


def test_样式覆盖优先级_盖优于列(tmp_path):
    from 排版计划 import 簿计划, 表计划, 列计划, 格样式
    from 报告写入 import 渲染簿计划

    表 = 表计划(
        名称="t",
        表=pd.DataFrame({"a": [1]}),
        列=[列计划(基础样式=格样式(填充色="FF0000FF", 填充类型="solid", 加粗=True))],
        覆盖={(2, 1): 格样式(填充色="FF00FF00", 填充类型="solid")},
    )
    输出 = tmp_path / "出.xlsx"
    渲染簿计划(簿计划(表=[表]), str(输出))
    c = load_workbook(输出)["t"]["A2"]
    assert c.fill.fgColor.rgb == "FF00FF00"   # 覆盖的填充胜出
    assert c.font.bold                          # 覆盖未设加粗，列基础样式保留


def test_新建簿计划_默认裸表(tmp_path):
    from 排版计划 import 新建簿计划
    from 报告写入 import 渲染簿计划

    计划 = 新建簿计划({"甲表": pd.DataFrame({"x": [1, 2]}), "乙表": pd.DataFrame({"y": [3]})})
    输出 = tmp_path / "出.xlsx"
    渲染簿计划(计划, str(输出))
    wb = load_workbook(输出)
    assert wb.sheetnames == ["甲表", "乙表"]
    assert wb["甲表"]["A2"].value == 1
