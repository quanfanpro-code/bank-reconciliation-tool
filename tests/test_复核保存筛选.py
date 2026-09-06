"""人工填写后按完整事项筛选，保留原行与公式的计算上下文。"""

from contextlib import closing
from hashlib import sha256
import os

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
import pytest

from 底稿筛选 import FilterCriteria, export_filtered_workpaper
from tests.test_复核工作簿 import _报告


def _行表(簿, 名称, 行):
    表 = 簿.create_sheet(名称)
    for 值 in 行:
        表.append(值)
    return 表


def _表头(表):
    return {格.value: 格.column for 格 in 表[1] if 格.value}


def _可见行(表):
    return [tuple(格.value for 格 in 表[行]) for 行 in range(2, 表.max_row + 1) if not 表.row_dimensions[行].hidden]


def _说明(簿):
    return dict((行[0], 行[1]) for 行 in 簿['筛选说明'].iter_rows(min_row=2, values_only=True) if 行[0])


def _构造报告(tmp_path):
    路径 = tmp_path / '全量复核.xlsx'
    簿 = Workbook()
    簿.active.title = '核对结论'
    簿.active.append(['项目', '数值'])
    簿.active.append(['全量事项数', '=COUNTA(复核事项索引!A2:A4)'])
    _行表(簿, '复核事项索引', [
        ['事项编号', '匹配ID', '程序结论', '核对方式', '类型', '组金额', '最早日期', '判断依据', '人工核对结果', '当前核对结论'],
        ['事项甲', '关系甲', '疑点', '人工全查', '退款重付', 400000, '2026-01-01', '完整退款重付组', '=人工全查!I2', '=人工全查!J2'],
        ['事项乙', '关系乙', '疑点', '人工抽样', '逐笔', 100000, '2026-01-02', '抽样', '=人工抽样!I2', '=人工抽样!J2'],
        ['事项丙', '', '未对应', '人工全查', '单边记录', 50000, '2026-02-01', '单边超过重要性水平', '=人工全查!I5', '=人工全查!J5'],
    ])
    人工表头 = ['事项编号', '行别', '核对原因', '银行记录', '银行金额', '序时账记录', '序时账金额', '差额', '人工核对结果', '当前核对结论', '备注']
    全查 = _行表(簿, '人工全查', [人工表头,
        ['事项甲', '核对事项', '退款重付', '甲公司', 0, '甲公司', 0, 0, '否定对应', '=复核事项索引!J2', '已经查看退款组成'],
        ['', '银行组成', '', '付款', -200000],
        ['', '银行组成', '', '退回', 200000],
        ['事项丙', '核对事项', '单边记录', '丙公司', -50000, '', None, 50000, '保留未对应', '=复核事项索引!J4', '两份数据仍未找到对应'],
    ])
    _行表(簿, '人工抽样', [人工表头,
        ['事项乙', '核对事项', '抽样', '乙公司', 100000, '乙公司', 100000, 0, '确认对应', '=复核事项索引!J3', '已核对业务编号'],
    ])
    验证 = DataValidation(type='list', formula1='"确认对应,否定对应"')
    全查.add_data_validation(验证)
    验证.add('I2')
    _行表(簿, '核对明细', [
        ['原始记录号', '事项编号', '来源', '日期', '对方及凭证', '摘要', '金额', '当前核对结果'],
        ['银行2', '事项甲', '银行流水', '2026-01-01', '甲公司', '支付设备款', -200000, '=复核事项索引!J2'],
        ['银行3', '事项甲', '银行流水', '2026-01-01', '甲公司', '设备款退回', 200000, '=复核事项索引!J2'],
        ['银行4', '事项乙', '银行流水', '2026-01-02', '乙公司', '服务费', 100000, '=复核事项索引!J3'],
        ['账面2', '事项乙', '序时账', '2026-01-02', '乙公司', '服务费', 100000, '=复核事项索引!J3'],
        ['银行5', '事项丙', '银行流水', '2026-02-01', '丙公司', '未对应货款', -50000, '=复核事项索引!J4'],
    ])
    _行表(簿, '复核候选选项', [
        ['选项编号', '事项编号', '选择文字', '自动确认'],
        ['事项甲|0', '事项甲', '确认对应', 0],
        ['事项甲|1', '事项甲', '采用其他对应1', 0],
        ['事项乙|0', '事项乙', '确认对应', 0],
    ])
    _行表(簿, '复核候选组成', [
        ['选项编号', '事项编号', '原始记录号'],
        ['事项甲|0', '事项甲', '银行2'], ['事项甲|0', '事项甲', '银行3'],
        ['事项甲|1', '事项甲', '银行4'], ['事项乙|0', '事项乙', '银行4'], ['事项乙|0', '事项乙', '账面2'],
    ])
    _行表(簿, '其他对应供选择', [['事项编号', '可选结果', '原始出处'], ['事项甲', '采用其他对应1', '银行4'], ['事项乙', '采用其他对应1', '银行2']])
    _行表(簿, '月度核对', [['月份', '银行收入', '银行支出'], ['2026-01', 300000, 200000], ['2026-02', 0, 50000]])
    _行表(簿, '每日统计', [['日期', '收入'], ['2026-01-01', 200000]])
    _行表(簿, '月度差异组成', [['月份', '事项编号', '收入差额'], ['2026-01', '事项甲', 200000], ['2026-02', '事项丙', 0]])
    _行表(簿, '银行侧待查', [['原始出处', '摘要'], ['银行5', '丙公司未对应']])
    for 名称 in ('复核事项索引', '复核候选选项', '复核候选组成'):
        簿[名称].sheet_state = 'hidden'
    簿.save(路径)
    簿.close()
    return 路径


def test_真实报告人工填选保存后筛选单边仍保留结果和公式(tmp_path):
    来源, _ = _报告(tmp_path)
    with closing(load_workbook(来源)) as 簿:
        表 = 簿['人工全查']
        头 = _表头(表)
        表.cell(2, 头['人工核对结果'], '保留未对应')
        表.cell(2, 头['备注'], '已核对两份记录')
        簿.save(来源)
        原公式 = {表.title: {格.coordinate: 格.value for 行 in 表 for 格 in 行 if 格.data_type == 'f'} for 表 in 簿}
    原指纹 = sha256(来源.read_bytes()).hexdigest()
    输出 = tmp_path / '单边筛选.xlsx'

    export_filtered_workpaper(来源, 输出, FilterCriteria(include_text=('甲公司',), statuses=('保留未对应',)))

    assert sha256(来源.read_bytes()).hexdigest() == 原指纹
    with closing(load_workbook(输出)) as 簿:
        assert _说明(簿)['筛选事项数'] == 1
        assert len(_可见行(簿['核对明细'])) == 1
        assert 簿['人工全查'].cell(2, 头['人工核对结果']).value == '保留未对应'
        assert 簿['人工全查'].cell(2, 头['备注']).value == '已核对两份记录'
        for 名称, 公式 in 原公式.items():
            assert {格.coordinate: 格.value for 行 in 簿[名称] for 格 in 行 if 格.data_type == 'f'} == 公式


def test_筛选保留完整组候选上下文和下拉且非选中行仅隐藏(tmp_path):
    来源 = _构造报告(tmp_path)
    输出 = tmp_path / '甲公司筛选.xlsx'
    export_filtered_workpaper(来源, 输出, FilterCriteria(include_text=('甲公司',)))

    with closing(load_workbook(输出)) as 簿:
        assert len(_可见行(簿['核对明细'])) == 2
        assert len(_可见行(簿['人工全查'])) == 3
        assert len(_可见行(簿['人工抽样'])) == 0
        assert len(_可见行(簿['其他对应供选择'])) == 1
        assert 簿['复核候选组成'].max_row == 6
        assert 簿['核对明细'].max_row == 6
        assert 簿['人工全查'].data_validations.dataValidation[0].sqref == 'I2'
        assert 簿['银行侧待查'].sheet_state == 'hidden'
        assert '全量' in _说明(簿)['口径说明'] and '公式' in _说明(簿)['口径说明']
        assert 簿.calculation.fullCalcOnLoad and 簿.calculation.forceFullCalc


@pytest.mark.parametrize('状态,预期', [('否定对应', '事项甲'), ('保留未对应', '事项丙'), ('确认对应', '事项乙')])
def test_不依赖未重算公式缓存按人工实际选择筛选(状态, 预期, tmp_path):
    来源 = _构造报告(tmp_path)
    输出 = tmp_path / (状态 + '.xlsx')
    export_filtered_workpaper(来源, 输出, FilterCriteria(statuses=(状态,)))

    with closing(load_workbook(输出)) as 簿:
        assert {行[1] for 行 in _可见行(簿['核对明细'])} == {预期}
        assert _说明(簿)['筛选事项数'] == 1


def test_覆盖比例按gross金额选择退款组而非净额(tmp_path):
    来源 = _构造报告(tmp_path)
    输出 = tmp_path / '覆盖筛选.xlsx'
    export_filtered_workpaper(来源, 输出, FilterCriteria(coverage_ratio=0.5))

    with closing(load_workbook(输出)) as 簿:
        说明 = _说明(簿)
        assert 说明['筛选事项数'] == 1
        assert 说明['筛选事项金额'] == 400000
        assert 说明['全量事项金额'] == 550000
        assert 说明['实际金额覆盖比例'] == pytest.approx(400000 / 550000)
        assert {行[1] for 行 in _可见行(簿['核对明细'])} == {'事项甲'}


def test_空筛选和再次筛选都保留完整原计算数据(tmp_path):
    来源 = _构造报告(tmp_path)
    空输出 = tmp_path / '空筛选.xlsx'
    再输出 = tmp_path / '再次筛选.xlsx'
    export_filtered_workpaper(来源, 空输出, FilterCriteria(include_text=('不存在的单位',)))
    with closing(load_workbook(空输出)) as 簿:
        assert _可见行(簿['核对明细']) == []
        assert 簿['核对明细'].max_row == 6
        assert _说明(簿)['筛选事项数'] == 0
    export_filtered_workpaper(空输出, 再输出, FilterCriteria(include_text=('乙公司',)))
    with closing(load_workbook(再输出)) as 簿:
        assert {行[1] for 行 in _可见行(簿['核对明细'])} == {'事项乙'}
        assert _说明(簿)['筛选事项数'] == 1


def test_拒绝硬链接别名覆盖来源(tmp_path):
    来源 = _构造报告(tmp_path)
    别名 = tmp_path / '全量报告别名.xlsx'
    os.link(来源, 别名)
    原指纹 = sha256(来源.read_bytes()).hexdigest()

    with pytest.raises(ValueError, match='覆盖'):
        export_filtered_workpaper(来源, 别名, FilterCriteria())
    assert sha256(来源.read_bytes()).hexdigest() == 原指纹



def test_实际报告主行和组成排序保存后仍按唯一人工主行筛选(tmp_path):
    from openpyxl.formula.translate import Translator

    甲 = {'摘要': '设备采购', '对方': '甲公司', '业务编号': 'P01'}
    丙 = {'摘要': '原料采购', '对方': '丙公司', '业务编号': 'P03'}
    来源, _ = _报告(
        tmp_path,
        [('2026-01-10', -600000, 甲), ('2026-01-10', -300000, 甲), ('2026-02-01', -200000, 丙)],
        [('2026-01-10', -400000, 甲), ('2026-01-10', -200000, 甲)],
    )
    with closing(load_workbook(来源)) as 簿:
        表 = 簿['人工全查']
        头 = _表头(表)
        主行 = [行 for 行 in range(2, 表.max_row + 1) if 表.cell(行, 头['填报编号']).value]
        assert len(主行) == 2
        选中编号 = 表.cell(主行[0], 头['事项编号']).value
        表.cell(主行[0], 头['人工核对结果'], '否定对应')
        表.cell(主行[0], 头['备注'], '先逐项看组成再否定这组对应')
        表.cell(主行[1], 头['人工核对结果'], '保留未对应')
        数据 = [(行, [格.value for 格 in 表[行]]) for 行 in range(2, 表.max_row + 1)]
        排序 = sorted(数据, key=lambda 行: (行[1][头['行别'] - 1] != '核对事项', str(行[1][头['事项编号'] - 1]), str(行[1][头['行别'] - 1])))
        assert [原行 for 原行, _ in 排序] != [原行 for 原行, _ in 数据]
        # 实际改写完整行，包括隐藏身份列；公式按排序后所在行平移，随后保存再筛选。
        for 新行, (原行, 值) in enumerate(排序, 2):
            for 列, 内容 in enumerate(值, 1):
                if isinstance(内容, str) and 内容.startswith('='):
                    内容 = Translator(内容, origin=表.cell(原行, 列).coordinate).translate_formula(表.cell(新行, 列).coordinate)
                表.cell(新行, 列).value = 内容
        簿.save(来源)
        预期组成数 = sum(表.cell(行, 头['事项编号']).value == 选中编号 for 行 in range(2, 表.max_row + 1))
        assert 预期组成数 == 5
    原指纹 = sha256(来源.read_bytes()).hexdigest()
    输出 = tmp_path / '排序后否定结果筛选.xlsx'

    export_filtered_workpaper(来源, 输出, FilterCriteria(statuses=('否定对应',)))

    assert sha256(来源.read_bytes()).hexdigest() == 原指纹
    with closing(load_workbook(输出)) as 簿:
        表 = 簿['人工全查']
        assert len(_可见行(表)) == 预期组成数
        assert {行[头['事项编号'] - 1] for 行 in _可见行(表)} == {选中编号}
        主行 = [行 for 行 in _可见行(表) if 行[头['填报编号'] - 1]]
        assert len(主行) == 1
        assert 主行[0][头['人工核对结果'] - 1] == '否定对应'
        assert 主行[0][头['备注'] - 1] == '先逐项看组成再否定这组对应'
        assert _说明(簿)['筛选事项数'] == 1



def test_人工尚未选择不把空单元格当nan状态(tmp_path):
    来源 = _构造报告(tmp_path)
    with closing(load_workbook(来源)) as 簿:
        簿['人工全查']['I2'] = None
        簿.save(来源)
    输出 = tmp_path / '未人工选择.xlsx'

    export_filtered_workpaper(来源, 输出, FilterCriteria(statuses=('疑点',)))

    with closing(load_workbook(输出)) as 簿:
        assert {行[1] for 行 in _可见行(簿['核对明细'])} == {'事项甲'}
