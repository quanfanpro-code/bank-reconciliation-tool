"""列位置只控制显示，报表读取必须按列名识别并保护人工内容。"""
from contextlib import closing
from hashlib import sha256

import pytest
from openpyxl import Workbook, load_workbook

from 底稿筛选 import FilterCriteria, _worksheet_frame, _filter_groups, preview_filter, export_filtered_workpaper
from tests.test_人工结果计算 import _报告, excel, _打开, _选择, _表


def test_表头空格和插入的其他列不影响识别():
    book = Workbook()
    sheet = book.active
    sheet.append(['其他说明', ' 金额 ', None, '事项编号'])
    sheet.append(['保留', -100, None, 'M1'])
    frame = _worksheet_frame(sheet)
    assert frame.loc[0, '金额'] == -100
    assert frame.loc[0, '事项编号'] == 'M1'


def test_重复列不能静默覆盖():
    book = Workbook()
    sheet = book.active
    sheet.title = '核对明细'
    sheet.append(['事项编号', '金额', ' 金额 '])
    sheet.append(['M1', 100, 200])
    with pytest.raises(ValueError, match='核对明细.*重复.*金额'):
        _worksheet_frame(sheet)


def test_缺少人工结果列必须提示不能误认为未填写(tmp_path):
    source = _报告(tmp_path, [100], [100], [((0,), (0,))])
    with closing(load_workbook(source)) as book:
        sheet = book['人工全查']
        column = next(c.column for c in sheet[1] if c.value == '人工核对结果')
        sheet.cell(1, column).value = '被误改的列名'
        book.save(source)
    with pytest.raises(ValueError, match='人工全查.*缺少.*人工核对结果'):
        preview_filter(source, FilterCriteria())


def test_未设金额上下限不因选择来源而漏掉单边事项():
    import pandas as pd
    groups = pd.DataFrame([{'事项编号':'账独有', '组金额':100}])
    assert len(_filter_groups(groups, FilterCriteria(amount_basis='银行单笔'))) == 1


def test_Excel挪列插列后人工结果备注公式及筛选仍然正确(excel, tmp_path):
    source = _报告(tmp_path, [100, 200], [100, 200], [((0,), (0,)), ((1,), (1,))])
    with _打开(excel, source) as book:
        # 用真正的 Excel 移动整列，使跨表引用随移动正常更新。
        for name, label in [('人工全查','人工核对结果'), ('核对明细','金额'), ('复核事项索引','事项编号')]:
            sheet = book.Worksheets(name)
            headers = list(sheet.UsedRange.Value[0])
            sheet.Columns(headers.index(label) + 1).Cut()
            sheet.Columns(2).Insert()
        sheet = book.Worksheets('人工全查')
        sheet.Columns(3).Insert()
        sheet.Cells(1, 3).Value = '项目自定义说明'
        sheet.Cells(2, 3).Value = '额外列保留'
        _选择(excel, book, 'M1', '否定对应')
        headers = list(sheet.UsedRange.Value[0])
        sheet.Cells(2, headers.index('备注') + 1).Value = '已查凭证，不能对应'
        excel.CalculateFullRebuild()
        before = _表(book, '逐笔核对')
        assert '否定' in before[0]['当前核对结论']
        book.Save()
    fingerprint = sha256(source.read_bytes()).hexdigest()
    criteria = FilterCriteria(statuses=('否定对应',), amount_basis='银行单笔', min_amount=100, max_amount=100)
    assert preview_filter(source, criteria)['selected'] == 1
    output = tmp_path / '挪列后筛选.xlsx'
    export_filtered_workpaper(source, output, criteria)
    assert sha256(source.read_bytes()).hexdigest() == fingerprint
    with closing(load_workbook(source)) as original, closing(load_workbook(output)) as filtered:
        for name in ('人工全查', '核对明细', '复核事项索引'):
            assert list(original[name].values) == list(filtered[name].values)
    with _打开(excel, output) as book:
        assert _表(book, '逐笔核对') == before
        row = _表(book, '人工全查')[0]
        assert row['人工核对结果'] == '否定对应'
        assert row['备注'] == '已查凭证，不能对应'
        assert row['项目自定义说明'] == '额外列保留'
