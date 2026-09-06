"""第五阶段独立复核：原始输入保护和历史保存失败后的报告交付。"""

import hashlib

import pandas as pd
import pytest
from openpyxl import load_workbook

from application import run_reconciliation
from data_structures import MatcherConfig
from 项目记录 import LocalProjectStore


映射 = {
    "date": "日期", "amount": "金额", "summary": "摘要", "mode": "signed_amount",
    "account": "账户", "currency": "币种", "amount_basis": "本位币",
}


def _输入(tmp_path):
    表 = pd.DataFrame({"日期": ["2026-08-10"], "金额": [100], "摘要": ["甲公司回款"],
                       "账户": ["A001"], "币种": ["CNY"]})
    银行, 账 = tmp_path / "银行.xlsx", tmp_path / "账.xlsx"
    for 路径 in (银行, 账):
        表.to_excel(路径, index=False)
    return 银行, 账


def _运行(银行, 账, 输出, 存储, **kwargs):
    return run_reconciliation(str(银行), str(账), 映射, 映射, MatcherConfig(),
                              output_path=输出, project_store=存储, **kwargs)


@pytest.mark.parametrize("来源", [0, 1])
@pytest.mark.parametrize("路径形式", ["原路径", "硬链接"])
def test_报告路径指向任一原始输入时拒绝且原文件不变(tmp_path, 来源, 路径形式):
    输入 = _输入(tmp_path)
    输出 = 输入[来源]
    if 路径形式 == "硬链接":
        输出 = tmp_path / "报告别名.xlsx"
        输出.hardlink_to(输入[来源])
    原哈希 = [hashlib.sha256(路径.read_bytes()).hexdigest() for 路径 in 输入]
    with pytest.raises(ValueError, match="原始输入"):
        _运行(*输入, 输出, LocalProjectStore(tmp_path / "历史"))
    assert [hashlib.sha256(路径.read_bytes()).hexdigest() for 路径 in 输入] == 原哈希


@pytest.mark.parametrize("故障", ["历史目录不可写", "映射文件损坏"])
def test_历史保存失败仍交付已生成报告并清楚提示(tmp_path, 故障):
    输入 = _输入(tmp_path)
    存储 = LocalProjectStore(tmp_path / "历史")
    if 故障 == "历史目录不可写":
        存储.root.write_text("这是文件，不能作为目录", encoding="utf-8-sig")
    else:
        存储.root.mkdir()
        存储.mapping_path.write_text("{损坏的JSON", encoding="utf-8-sig")
    日志, 进度 = [], []
    输出 = tmp_path / "报告.xlsx"
    结果 = _运行(*输入, 输出, 存储, logger=日志.append, progress_callback=进度.append)
    assert 结果 == 输出
    簿 = load_workbook(结果, read_only=True)
    try:
        assert "核对结论" in 簿.sheetnames
    finally:
        簿.close()
    assert any("历史" in 行 and "未保存" in 行 and str(输出) in 行 for 行 in 日志)
    assert "全部完成" in 日志[-1]
    assert 进度[-1] == 1.0
