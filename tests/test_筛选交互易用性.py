"""从真实窗口操作验证条件清单、报告自动读取和列拖动。"""
from contextlib import closing
from hashlib import sha256
from threading import Event
import time
import tkinter as tk

import pytest
from openpyxl import load_workbook

from 筛选窗口 import FilterDialog
from tests.test_人工结果计算 import _报告


@pytest.fixture
def root():
    window = tk.Tk()
    window.geometry('240x80+20+20')
    yield window
    window.destroy()


def _wait(root, predicate):
    deadline = time.monotonic() + 15
    while not predicate() and time.monotonic() < deadline:
        root.update()
        time.sleep(.01)
    assert predicate(), '后台读取未按期完成'


def _add(dialog, kind, text):
    dialog.rule_kind.set(kind)
    dialog.rule_value.set(text)
    dialog.add_rule()


def test_选好报告自动加载选项并可直接筛选(root, tmp_path):
    source = _报告(tmp_path, [100, 200], [100, 200], [((0,), (0,)), ((1,), (1,))])
    before = sha256(source.read_bytes()).hexdigest()
    dialog = FilterDialog(root, source, tmp_path/'偏好.json')
    _wait(root, lambda: dialog.report_data is not None)
    assert dialog.report_options['business_types']
    assert dialog.report_options['statuses']
    _add(dialog, '业务类型', dialog.report_options['business_types'][0])
    dialog.preview()
    _wait(root, lambda: dialog.preview_started is None)
    assert '2 / 2' in dialog.preview_text.get()
    assert sha256(source.read_bytes()).hexdigest() == before
    dialog.cancel()


def test_文字逐条录入修改删除且标点不充当分隔符(root, tmp_path):
    dialog = FilterDialog(root, preferences_path=tmp_path/'偏好.json')
    _add(dialog, '包含文字', '工资;奖金')
    _add(dialog, '包含文字', '甲公司')
    _add(dialog, '排除文字', '手续费')
    assert dialog.build_criteria().include_text == ('工资;奖金', '甲公司')
    first, second, third = dialog.rule_list.get_children()
    dialog.rule_list.selection_set(first)
    dialog.rule_kind.set('包含文字'); dialog.rule_value.set('工资')
    dialog.edit_rule()
    dialog.rule_list.selection_set(second)
    dialog.remove_rule()
    criteria = dialog.build_criteria()
    assert criteria.include_text == ('工资',)
    assert criteria.exclude_text == ('手续费',)
    dialog.clear()
    assert not dialog.build_criteria().include_text
    assert not dialog.rule_list.get_children()
    dialog.cancel()


def test_加载失败可以重新选报告恢复(root, tmp_path):
    source = _报告(tmp_path, [100], [100], [((0,), (0,))])
    dialog = FilterDialog(root, tmp_path/'不存在.xlsx', tmp_path/'偏好.json')
    _wait(root, lambda: bool(dialog.source_error))
    assert str(dialog.export_button['state']) == 'disabled'
    dialog.set_source(source)
    _wait(root, lambda: dialog.report_data is not None)
    assert not dialog.source_error
    assert str(dialog.export_button['state']) == 'normal'
    dialog.cancel()


def test_快速换报告不会使用较早任务的结果(root, tmp_path, monkeypatch):
    import 筛选窗口 as module
    one = tmp_path/'一'; two = tmp_path/'二'; one.mkdir(); two.mkdir()
    a = _报告(one, [100], [100], [((0,), (0,))])
    b = _报告(two, [100, 200], [100, 200], [((0,), (0,)), ((1,), (1,))])
    started, release = Event(), Event()
    original = module.read_filter_report
    def slow(path):
        if str(path) == str(a):
            started.set(); assert release.wait(10)
        return original(path)
    monkeypatch.setattr(module, 'read_filter_report', slow)
    dialog = FilterDialog(root, a, tmp_path/'偏好.json')
    try:
        _wait(root, started.is_set)
        dialog.set_source(b)
        dialog.clear()
        assert dialog.report_data is None
        release.set()
        _wait(root, lambda: dialog.report_data is not None)
        assert str(dialog.source_path) == str(b)
        assert '2 / 2' in dialog.preview_text.get()
    finally:
        release.set(); dialog.cancel()


def test_报告人工填写更新后重新读取而不是沿用缓存(root, tmp_path):
    source = _报告(tmp_path, [100], [100], [((0,), (0,))])
    dialog = FilterDialog(root, source, tmp_path/'偏好.json')
    _wait(root, lambda: dialog.report_data is not None)
    with closing(load_workbook(source)) as book:
        sheet = book['人工全查']
        col = next(c.column for c in sheet[1] if c.value == '人工核对结果')
        sheet.cell(2, col).value = '否定对应'
        book.save(source)
    dialog.preview()
    _wait(root, lambda: dialog.report_data is not None and dialog.preview_started is None)
    assert '否定对应' in dialog.report_options['statuses']
    dialog.cancel()


def test_拖动列后保存重开及键盘移动不丢列(root, tmp_path):
    path = tmp_path/'偏好.json'
    dialog = FilterDialog(root, preferences_path=path)
    dialog.tabs.select(dialog.column_tab)
    root.update()
    listing = dialog.column_list
    original = list(listing.get(0, 'end'))
    y0 = listing.bbox(0)[1] + 5
    y3 = listing.bbox(3)[1] + 5
    listing.event_generate('<ButtonPress-1>', x=12, y=y0)
    listing.event_generate('<B1-Motion>', x=12, y=y3)
    listing.event_generate('<ButtonRelease-1>', x=12, y=y3)
    root.update()
    expected = original[1:4] + original[:1] + original[4:]
    assert list(listing.get(0, 'end')) == expected
    dialog.save_columns(); dialog.cancel()
    reopened = FilterDialog(root, preferences_path=path)
    reopened.tabs.select(reopened.column_tab); root.update()
    assert list(reopened.column_list.get(0, 'end')) == expected
    reopened.column_list.selection_set(3)
    reopened.column_list.focus_force()
    reopened.column_list.event_generate('<Alt-Up>')
    root.update()
    assert reopened.column_list.get(2) == original[0]
    assert sorted(reopened.column_list.get(0, 'end')) == sorted(original)
    reopened.cancel()


def test_多条件预览与真正导出的整项一致(root, tmp_path):
    from 底稿筛选 import export_filtered_workpaper
    source = _报告(tmp_path, [100,200], [100,200], [((0,), (0,)), ((1,), (1,))])
    dialog = FilterDialog(root, source, tmp_path/'偏好.json')
    _wait(root, lambda: dialog.report_data is not None)
    _add(dialog, '包含文字', '设备款'); _add(dialog, '包含文字', '甲公司')
    _add(dialog, '排除文字', '工资')
    dialog.variables['amount_basis'].set('银行单笔'); dialog.variables['min_amount'].set('200')
    dialog.preview(); _wait(root, lambda: dialog.preview_started is None)
    assert '1 / 2' in dialog.preview_text.get()
    dialog.confirm()
    output = tmp_path/'实际导出.xlsx'
    export_filtered_workpaper(source, output, dialog.result)
    with closing(load_workbook(output)) as book:
        sheet = book['逐笔核对']
        headers = {cell.value:cell.column for cell in sheet[1]}
        visible = {sheet.cell(row, headers['事项编号']).value for row in range(2,sheet.max_row+1) if not sheet.row_dimensions[row].hidden}
        assert visible == {'M2'}


def test_报告改变后旧状态条件必须提示修正(root, tmp_path):
    source = _报告(tmp_path, [100], [100], [((0,), (0,))])
    dialog = FilterDialog(root, source, tmp_path/'偏好.json')
    _wait(root, lambda: dialog.report_data is not None)
    _add(dialog, '核对结果', dialog.report_options['statuses'][0])
    with closing(load_workbook(source)) as book:
        sheet = book['人工全查']
        col = next(c.column for c in sheet[1] if c.value == '人工核对结果')
        sheet.cell(2, col).value = '否定对应'; book.save(source)
    dialog.preview(); _wait(root, lambda: dialog.report_data is not None)
    dialog.confirm()
    assert dialog.result is None
    assert '当前报告' in dialog.preview_text.get()
    dialog.cancel()


def test_很多长文字条件也不挤掉底部操作(root, tmp_path):
    dialog = FilterDialog(root, preferences_path=tmp_path/'偏好.json')
    dialog.geometry('820x760')
    for number in range(25):
        _add(dialog, '包含文字', f'第{number}项供应商采购付款合同与发票核对说明')
    root.update()
    assert dialog.export_button.winfo_ismapped()
    assert dialog.export_button.winfo_rooty()+dialog.export_button.winfo_height() <= dialog.winfo_rooty()+dialog.winfo_height()
    dialog.cancel()
