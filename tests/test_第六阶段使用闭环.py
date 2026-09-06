import hashlib
import json
import subprocess
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from 底稿筛选 import FilterCriteria, export_filtered_workpaper
from 输入模板 import generate_input_templates
from 项目记录 import LocalProjectStore


def _sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def test_列映射模板按表头指纹保存和复用(tmp_path):
    store = LocalProjectStore(tmp_path / "records")
    mapping = {"date": "交易日期", "amount": "金额", "summary": "摘要"}

    fingerprint = store.save_mapping_template("bank", ["金额", "摘要", "交易日期"], mapping, "建行流水")
    loaded = store.load_mapping_template("bank", ["交易日期", "金额", "摘要"])

    assert loaded["fingerprint"] == fingerprint
    assert loaded["mapping"] == mapping
    assert loaded["name"] == "建行流水"


def test_项目记录保存输入与报告指纹并可重新打开(tmp_path):
    bank = tmp_path / "bank.xlsx"
    journal = tmp_path / "journal.xlsx"
    report = tmp_path / "report.xlsx"
    bank.write_bytes(b"bank")
    journal.write_bytes(b"journal")
    report.write_bytes(b"report")
    store = LocalProjectStore(tmp_path / "records")

    record = store.save_project(
        bank_path=bank,
        journal_path=journal,
        report_path=report,
        bank_mapping={"date": "日期"},
        journal_mapping={"date": "凭证日期"},
        parameters={"tolerance_days": 3},
        result_counts={"selected_groups": 5},
    )
    latest = store.list_projects()[0]

    assert latest["project_id"] == record["project_id"]
    assert latest["files"]["report"]["sha256"] == _sha256(report)
    assert store.resolve_report(record["project_id"]) == report


def test_筛选任一关系时完整带出组成且不覆盖全量报告(tmp_path):
    source = tmp_path / "full.xlsx"
    output = tmp_path / "filtered.xlsx"
    with pd.ExcelWriter(source, engine="openpyxl") as writer:
        pd.DataFrame({"项目": ["总银行流水金额"], "数值": [1000]}).to_excel(writer, sheet_name="核对结论", index=False)
        pd.DataFrame(
            {
                "匹配ID": ["M1", "M2"],
                "类型": ["手续费净额", "普通匹配"],
                "最终状态": ["整组勾稽一致", "自动确认"],
                "最早日期": ["2026-01-01", "2026-01-02"],
                "组金额": [800, 200],
                "判断依据": ["手续费明确", "普通业务"],
            }
        ).to_excel(writer, sheet_name="整组勾稽", index=False)
        pd.DataFrame(
            {
                "匹配ID": ["M1", "M1", "M1", "M2", "M2"],
                "来源": ["银行流水", "银行存款序时账", "银行存款序时账", "银行流水", "银行存款序时账"],
                "金额": [790, 800, 10, 200, 200],
                "摘要": ["回款", "应收", "手续费", "普通", "普通"],
            }
        ).to_excel(writer, sheet_name="匹配组成", index=False)
    before = _sha256(source)

    export_filtered_workpaper(source, output, FilterCriteria(business_types=("手续费净额",)))

    assert _sha256(source) == before
    result = pd.read_excel(output, sheet_name="匹配组成")
    explanation = pd.read_excel(output, sheet_name="筛选说明")
    assert list(result["匹配ID"].unique()) == ["M1"]
    assert len(result) == 3
    assert explanation.loc[explanation["项目"] == "筛选关系数", "数值"].iloc[0] == 1
    assert explanation.loc[explanation["项目"] == "覆盖比例", "数值"].iloc[0] == 0.8


def test_摘要或对方文字筛选会读取关系组成而不只看关系首页(tmp_path):
    source = tmp_path / "full.xlsx"
    output = tmp_path / "filtered.xlsx"
    with pd.ExcelWriter(source, engine="openpyxl") as writer:
        pd.DataFrame({"匹配ID": ["M1", "M2"], "类型": ["普通匹配", "普通匹配"], "组金额": [600, 400]}).to_excel(writer, sheet_name="逐笔匹配", index=False)
        pd.DataFrame({"匹配ID": ["M1", "M1", "M2", "M2"], "来源": ["银行流水", "银行存款序时账"] * 2, "金额": [600, 600, 400, 400], "摘要": ["设备款", "购置设备", "房租", "支付租金"], "辅助文字": ["对方：甲公司", "对方：甲公司", "对方：乙公司", "对方：乙公司"]}).to_excel(writer, sheet_name="匹配组成", index=False)

    export_filtered_workpaper(source, output, FilterCriteria(include_text=("甲公司",)))

    groups = pd.read_excel(output, sheet_name="逐笔匹配")
    components = pd.read_excel(output, sheet_name="匹配组成")
    assert list(groups["匹配ID"]) == ["M1"]
    assert set(components["匹配ID"]) == {"M1"}


def test_筛选导出长任务持续报告阶段进度(tmp_path):
    source = tmp_path / "full.xlsx"
    output = tmp_path / "filtered.xlsx"
    with pd.ExcelWriter(source, engine="openpyxl") as writer:
        pd.DataFrame({"匹配ID": ["M1"], "类型": ["普通匹配"], "组金额": [100]}).to_excel(
            writer, sheet_name="逐笔匹配", index=False
        )
        pd.DataFrame({"匹配ID": ["M1", "M1"], "来源": ["银行流水", "银行存款序时账"]}).to_excel(
            writer, sheet_name="匹配组成", index=False
        )
    progress = []
    logs = []

    export_filtered_workpaper(
        source,
        output,
        FilterCriteria(),
        progress_callback=progress.append,
        log_callback=logs.append,
    )

    assert progress[0] == 0.0
    assert progress[-1] == 1.0
    assert progress == sorted(progress)
    assert any("读取全量报告" in message for message in logs)
    assert any("写入筛选底稿" in message for message in logs)
    assert any("筛选导出完成" in message for message in logs)


def test_输入模板生成银行流水和序时账两个可直接填写文件(tmp_path):
    paths = generate_input_templates(tmp_path)

    assert {path.name for path in paths} == {"银行流水导入模板.xlsx", "银行存款序时账导入模板.xlsx"}
    bank_book = load_workbook(tmp_path / "银行流水导入模板.xlsx", read_only=True)
    journal_book = load_workbook(tmp_path / "银行存款序时账导入模板.xlsx", read_only=True)
    assert bank_book.active.cell(1, 1).value == "交易日期"
    assert journal_book.active.cell(1, 1).value == "凭证日期"
    assert {"交易日期", "收入金额", "支出金额", "账户", "币种", "流水号", "摘要"}.issubset(
        {cell.value for cell in bank_book.active[1]}
    )
    assert {"凭证日期", "借方金额", "贷方金额", "账户", "币种", "凭证号", "摘要"}.issubset(
        {cell.value for cell in journal_book.active[1]}
    )


def test_windows双击入口提供自检且无需手工输入命令():
    root = Path(__file__).resolve().parents[1]
    launcher = root / "启动银行流水核对工具.cmd"

    result = subprocess.run(["cmd.exe", "/c", str(launcher), "--check"], cwd=root, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=20)

    assert result.returncode == 0
    assert "READY" in result.stdout


def test_主程序提供打包后可用的自检入口():
    root = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        [
            r"C:\Users\27651\AppData\Local\Programs\Python\Python314\python.exe",
            str(root / "main.py"),
            "--check",
        ],
        cwd=root,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=20,
    )

    assert result.returncode == 0
    assert "READY" in result.stdout
