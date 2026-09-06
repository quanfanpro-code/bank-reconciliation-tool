"""输出阅读布局验收：只检查既有结果的展示，不改变计算规则。"""
from contextlib import closing
from hashlib import sha256

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from 底稿筛选 import FilterCriteria, export_filtered_workpaper
from tests.test_人工结果计算 import _报告


可见表 = {'核对结论','月度核对','逐笔核对','整组核对','人工全查','人工抽样','未对应记录'}
原始金额列 = ('银行收入','银行支出','序时账收入','序时账支出')


def _头(sheet):
    return {cell.value:cell.column for cell in sheet[1] if cell.value is not None}


def _行(sheet):
    headers=_头(sheet)
    return [(number,{key:sheet.cell(number,column).value for key,column in headers.items()})
            for number in range(2,sheet.max_row+1)
            if any(sheet.cell(number,column).value is not None for column in headers.values())]


def _隐藏身份(sheet):
    headers=_头(sheet)
    for key in ('事项编号','原始记录号'):
        if key in headers:
            assert sheet.column_dimensions[get_column_letter(headers[key])].hidden


def test_七个固定阅读入口及一对一并排单边完整(tmp_path):
    path=_报告(tmp_path,[100,300],[100,300,700],[((0,),(0,)),((1,),(1,))])
    with closing(load_workbook(path)) as book:
        assert {sheet.title for sheet in book if sheet.sheet_state=='visible'}==可见表
        pairs=book['逐笔核对']
        rows=_行(pairs)
        assert len(rows)==2
        assert [(row['银行金额'],row['序时账金额']) for _,row in rows]==[(100,100),(300,300)]
        assert all(row['银行记录'] and row['序时账记录'] for _,row in rows)
        unmatched=book['未对应记录']
        assert len(_行(unmatched))==1
        row=_行(unmatched)[0][1]
        assert row['原始记录号']=='账面4'
        assert row['金额']==700
        assert row['来源']=='序时账'
        assert row['日期'] and row['摘要'] and row['原始出处']
        _隐藏身份(pairs)
        _隐藏身份(unmatched)


def test_两对三先汇总后分块组成可折叠且不重复加总(tmp_path):
    path=_报告(tmp_path,[150,450],[100,200,300],[((0,1),(0,1,2))])
    with closing(load_workbook(path)) as book:
        sheet=book['整组核对']
        rows=_行(sheet)
        assert [row['行别'] for _,row in rows]==['核对事项','银行组成','银行组成','序时账组成','序时账组成','序时账组成']
        main_number,main=rows[0]
        assert '2笔' in str(main['银行记录']) and '3笔' in str(main['序时账记录'])
        assert all(main[key] is None for key in 原始金额列)
        assert not sheet.row_dimensions[main_number].hidden
        assert all(sum(float(row.get(key) or 0) for _,row in rows)==expected for key,expected in [('银行收入',600),('银行支出',0),('序时账收入',600),('序时账支出',0)])
        header=_头(sheet)
        assert sheet.cell(main_number,header['当前核对结论']).data_type=='f'
        assert '复核事项索引' in main['当前核对结论']
        for number,row in rows[1:]:
            assert sheet.row_dimensions[number].outlineLevel==1
            assert sheet.row_dimensions[number].hidden
            assert row['事项编号']=='M1' and row['原始记录号']
            assert sheet.cell(number,header['当前记录结果']).data_type=='f'
            assert '核对明细' in row['当前记录结果']
            opposite=('序时账收入','序时账支出') if row['行别']=='银行组成' else ('银行收入','银行支出')
            assert all(row[key] is None for key in opposite)
        _隐藏身份(sheet)


def test_混合收付整组原始列分别保留收入支出(tmp_path):
    path=_报告(tmp_path,[200,-50],[250,-100],[((0,1),(0,1))])
    with closing(load_workbook(path)) as book:
        rows=_行(book['整组核对'])
        main=rows[0][1]
        assert all('收入' in str(main[key]) and '支出' in str(main[key]) for key in ('银行记录','序时账记录'))
        assert all(main[key] is None for key in 原始金额列)
        totals={key:sum(float(row.get(key) or 0) for _,row in rows) for key in 原始金额列}
        assert totals=={'银行收入':200,'银行支出':50,'序时账收入':250,'序时账支出':100}


def test_人工组只在主行填一次且原组成展开替代金额不混加(tmp_path):
    path=_报告(tmp_path,[150,450],[100,200,300,600],[((0,1),(0,1,2))],[(0,(0,1),(3,))])
    with closing(load_workbook(path)) as book:
        sheet=book['人工全查']
        header=_头(sheet)
        rows=[(number,row) for number,row in _行(sheet) if row['事项编号']=='M1']
        mains=[(number,row) for number,row in rows if row['行别']=='核对事项']
        assert len(mains)==1 and mains[0][1]['填报编号']=='M1'
        assert all(not row.get('填报编号') for _,row in rows if row['行别']!='核对事项')
        inputs=[number for number,_ in rows if any(sheet.cell(number,header['人工核对结果']).coordinate in validation.sqref for validation in sheet.data_validations.dataValidation)]
        assert inputs==[mains[0][0]]
        original=[(number,row) for number,row in rows if row['行别'] in ('银行组成','序时账组成')]
        assert len(original)==5
        assert all(not sheet.row_dimensions[number].hidden for number,_ in original)
        assert all(sum(float(row.get(key) or 0) for _,row in rows)==expected for key,expected in [('银行收入',600),('银行支出',0),('序时账收入',600),('序时账支出',0)])
        alternatives=[row for _,row in rows if '其他对应' in '|'.join(str(value or '') for value in row.values())]
        assert alternatives
        assert all(all(row.get(key) is None for key in 原始金额列) for row in alternatives)


def test_月度发生额和缺失余额同表可见且余额不补零(tmp_path):
    path=_报告(tmp_path,[100],[90],[((0,),(0,))])
    with closing(load_workbook(path)) as book:
        sheet=book['月度核对']
        header=_头(sheet)
        required=('银行-收入金额','序时账-收入金额','银行-支出金额','序时账-支出金额','银行-月末余额','序时账-月末余额','余额差额')
        assert set(required)<=header.keys()
        assert all(not sheet.column_dimensions[get_column_letter(header[key])].hidden for key in required)
        row=next(row for _,row in _行(sheet) if str(row.get('月份'))=='2026-01')
        assert row['银行-收入金额']==100 and row['序时账-收入金额']==90
        assert row['收入金额差额']==10
        assert row['银行-月末余额']=='缺失' and row['序时账-月末余额']=='缺失'
        assert row['余额差额'] in (None,'缺失')
        for _,total in _行(sheet):
            if '合计' in str(total.get('月份','')):
                assert not isinstance(total['银行-月末余额'],(int,float))
                assert not isinstance(total['序时账-月末余额'],(int,float))


def test_筛选新阅读表整组保留所有组成且不改计算公式(tmp_path):
    source=_报告(tmp_path,[150,450,500],[100,200,300,500],[((0,1),(0,1,2)),((2,),(3,))])
    before=sha256(source.read_bytes()).hexdigest()
    with closing(load_workbook(source)) as book:
        formulas={name:{cell.coordinate:cell.value for row in book[name] for cell in row if cell.data_type=='f'} for name in ('核对明细','复核事项索引','复核候选选项','复核候选组成')}
    output=tmp_path/'整组筛选.xlsx'
    export_filtered_workpaper(source,output,FilterCriteria(coverage_ratio=0.5))
    assert sha256(source.read_bytes()).hexdigest()==before
    with closing(load_workbook(output)) as book:
        group=book['整组核对']
        assert group.sheet_state=='visible'
        rows=[(number,row) for number,row in _行(group) if row['事项编号']=='M1']
        assert len(rows)==6
        assert {row['原始记录号'] for _,row in rows if row.get('原始记录号')}=={'银行2','银行3','账面2','账面3','账面4'}
        assert not group.row_dimensions[rows[0][0]].hidden
        assert all(book['逐笔核对'].row_dimensions[number].hidden for number,row in _行(book['逐笔核对']) if row['事项编号']=='M2')
        for name,expected in formulas.items():
            assert {cell.coordinate:cell.value for row in book[name] for cell in row if cell.data_type=='f'}==expected


def test_折叠组主行可直接阅读双方日期范围和业务摘要(tmp_path,monkeypatch):
    import pandas as pd
    import tests.test_人工结果计算 as calculation
    original_df=calculation._df
    def dated_df(rows):
        frame=original_df(rows)
        start='2026-01-05' if len(rows)==2 else '2026-01-07'
        frame['date']=pd.date_range(start,periods=len(rows))
        return frame
    monkeypatch.setattr(calculation,'_df',dated_df)
    path=_报告(tmp_path,[150,450],[100,200,300],[((0,1),(0,1,2))])
    with closing(load_workbook(path)) as book:
        rows=_行(book['整组核对'])
        main=rows[0][1]
        assert all(book['整组核对'].row_dimensions[number].hidden for number,_ in rows[1:])
        for column,first,last in [('银行记录','2026-01-05','2026-01-06'),('序时账记录','2026-01-07','2026-01-09')]:
            text=str(main[column])
            assert first in text and last in text and '设备款' in text


from tests.test_人工结果计算 import excel, _打开, _选择, _表


def test_逐笔可见结果随否定和替代反映原始双方当前去向(excel,tmp_path):
    path=_报告(tmp_path,[100],[100,100],[((0,),(0,))],[(0,(0,),(1,))])
    with _打开(excel,path) as book:
        _选择(excel,book,'M1','否定对应')
        row=next(row for row in _表(book,'逐笔核对') if row['事项编号']=='M1')
        assert '否定' in row['当前核对结论'] and '已确认' not in row['当前核对结论']
        _选择(excel,book,'M1','采用其他对应1')
        row=next(row for row in _表(book,'逐笔核对') if row['事项编号']=='M1')
        assert '银行：' in row['当前核对结论'] and '序时账：' in row['当前核对结论']
        assert '已确认' in row['当前核对结论'] and '未对应' in row['当前核对结论']


def test_原逐笔双方同被新组采用显示新组核对编号(excel,tmp_path,monkeypatch):
    from data_structures import ProcessingStatus,RiskLevel
    from reporter import Reporter
    generate=Reporter.generate_report
    def with_auto_first(self,*args,**kwargs):
        first=self.matcher.selected_candidates[0]
        first.processing_status=ProcessingStatus.AUTO_CONFIRMED
        first.risk_level=RiskLevel.NORMAL
        first.is_ambiguous=False
        first.processing_reason='程序已确认对应关系'
        return generate(self,*args,**kwargs)
    monkeypatch.setattr(Reporter,'generate_report',with_auto_first)
    path=_报告(tmp_path,[100,200],[100,200],[((0,),(0,)),((1,),(1,))],[(1,(0,1),(0,1))])
    with _打开(excel,path) as book:
        _选择(excel,book,'M2','采用其他对应1')
        source=next(row for row in _表(book,'逐笔核对') if row['事项编号']=='M1')
        target=next(row for row in _表(book,'复核事项索引') if row['事项编号']=='M2')
        text=source['当前核对结论']
        assert str(target['核对编号']) in text and '对应' in text
        assert '未对应' not in text and '待人工' not in text


def test_新阅读表原始危险文本继续作为文本转义(tmp_path,monkeypatch):
    import tests.test_人工结果计算 as calculation
    original_df=calculation._df
    dangerous=['=1+1','+CMD','-2+3','@SUM(A1)']
    def dangerous_df(rows):
        frame=original_df(rows)
        for index in frame.index:
            value=dangerous[index%len(dangerous)]
            frame.at[index,'summary']=value
            fields=dict(frame.at[index,'aux_text_fields'])
            fields['摘要']=value
            frame.at[index,'aux_text_fields']=fields
        return frame
    monkeypatch.setattr(calculation,'_df',dangerous_df)
    path=_报告(tmp_path,[101,102,103,104],[101,102,103,104,700],[((i,),(i,)) for i in range(4)])
    with closing(load_workbook(path,data_only=False)) as book:
        cells=[cell for name in ('逐笔核对','人工全查','未对应记录') for row in book[name] for cell in row]
        for value in dangerous:
            found=[cell for cell in cells if value in str(cell.value or '')]
            assert found, value
            assert all(cell.data_type!='f' for cell in found), value

def test_人工表排序后Excel计算跳转目标仍指向本事项填写格(excel,tmp_path):
    path=_报告(tmp_path,[100,200],[100,200],[((0,),(0,)),((1,),(1,))])
    with _打开(excel,path) as book:
        manual=book.Worksheets('人工全查')
        headers=list(manual.UsedRange.Value[0])
        before=[row['事项编号'] for row in _表(book,'人工全查')]
        manual.UsedRange.Sort(Key1=manual.Cells(2,headers.index('事项编号')+1),Order1=2,Header=1)
        after=[row['事项编号'] for row in _表(book,'人工全查')]
        assert before!=after and after==sorted(before,reverse=True)
        excel.CalculateFullRebuild()
        pairs=book.Worksheets('逐笔核对')
        pair_headers=list(pairs.UsedRange.Value[0])
        target_row=next(number for number,row in enumerate(_表(book,'逐笔核对'),2) if row['事项编号']=='M1')
        cell=pairs.Cells(target_row,pair_headers.index('核对编号')+1)
        pairs.Activate()
        if str(cell.Formula).upper().startswith('=HYPERLINK('):
            expression=cell.Formula[len('=HYPERLINK('):].rsplit(',',1)[0]
            actual_target=pairs.Evaluate(expression)
            assert actual_target.startswith('#')
            sheet_name,address=actual_target[1:].rsplit('!',1)
            sheet_name=sheet_name.strip("'").replace("''", "'")
            excel.Goto(Reference=book.Worksheets(sheet_name).Range(address))
        else:
            cell.Hyperlinks(1).Follow(NewWindow=False,AddHistory=False)
        assert excel.ActiveSheet.Name=='人工全查'
        assert manual.Cells(excel.ActiveCell.Row,headers.index('填报编号')+1).Value=='M1'
        assert excel.ActiveCell.Column==headers.index('人工核对结果')+1

def test_空数据真实报告月度只有表头且不造零余额或年度合计(tmp_path):
    from data_structures import MatcherConfig
    from matcher import Matcher
    from reporter import Reporter
    from tests.test_第三阶段特殊业务 import _df
    config=MatcherConfig()
    matcher=Matcher(_df([]),_df([]),config,logger=lambda _:None)
    matcher.run()
    path=tmp_path/'空数据核对报告.xlsx'
    Reporter(matcher).generate_report(str(path),config=config)
    with closing(load_workbook(path)) as book:
        monthly=book['月度核对']
        assert monthly.max_row==1
        assert _行(monthly)==[]
        assert {'银行-月末余额','序时账-月末余额','余额差额'}<=_头(monthly).keys()
        assert not monthly.conditional_formatting
        assert {sheet.title for sheet in book if sheet.sheet_state=='visible'}==可见表
        assert not any('年合计' in str(cell.value or '') for row in monthly for cell in row)


def test_原整组被新组采用后主行不误称全未对应且组成显示新编号(excel,tmp_path,monkeypatch):
    from data_structures import ProcessingStatus,RiskLevel
    from reporter import Reporter
    generate=Reporter.generate_report
    def with_auto_first(self,*args,**kwargs):
        first=self.matcher.selected_candidates[0]
        first.processing_status=ProcessingStatus.AUTO_CONFIRMED
        first.risk_level=RiskLevel.NORMAL
        first.is_ambiguous=False
        first.processing_reason='程序已确认对应关系'
        return generate(self,*args,**kwargs)
    monkeypatch.setattr(Reporter,'generate_report',with_auto_first)
    path=_报告(tmp_path,[100,200,300],[100,200,300],[((0,1),(0,1)),((2,),(2,))],[(1,(0,1,2),(0,1,2))])
    with _打开(excel,path) as book:
        _选择(excel,book,'M2','采用其他对应1')
        rows=[row for row in _表(book,'整组核对') if row['事项编号']=='M1']
        main=next(row for row in rows if row['行别']=='核对事项')
        assert '未对应' not in main['当前核对结论']
        assert '替代' in main['当前核对结论'] and '展开' in main['当前核对结论']
        target=next(row for row in _表(book,'复核事项索引') if row['事项编号']=='M2')
        children=[row for row in rows if row['行别'] in ('银行组成','序时账组成')]
        assert len(children)==4
        assert all(str(target['核对编号']) in row['当前核对结论'] for row in children)


def test_采用候选一后候选二说明行不冒用事项已确认结论(excel,tmp_path):
    path=_报告(tmp_path,[100],[100,100,100],[((0,),(0,))],[(0,(0,),(1,)),(0,(0,),(2,))])
    with _打开(excel,path) as book:
        _选择(excel,book,'M1','采用其他对应1')
        rows=[row for row in _表(book,'人工全查') if row['事项编号']=='M1']
        candidate_two=[row for row in rows if row['行别']=='其他对应2']
        assert candidate_two
        assert all('已确认' not in str(row['当前核对结论'] or '') for row in candidate_two)


def test_原人工整组先否定后成员由新组采用不误称全未对应(excel,tmp_path):
    path=_报告(tmp_path,[100,200,300],[100,200,300],[((0,1),(0,1)),((2,),(2,))],[(1,(0,1,2),(0,1,2))])
    with _打开(excel,path) as book:
        _选择(excel,book,'M1','否定对应')
        _选择(excel,book,'M2','采用其他对应1')
        rows=[row for row in _表(book,'整组核对') if row['事项编号']=='M1']
        main=next(row for row in rows if row['行别']=='核对事项')
        assert '否定' in main['当前核对结论'] and '展开' in main['当前核对结论']
        assert '记录恢复未对应' not in main['当前核对结论']
        assert all('002' in row['当前核对结论'] for row in rows if row['行别']!='核对事项')
        manual=next(row for row in _表(book,'人工全查') if row['事项编号']=='M1' and row['行别']=='核对事项')
        assert '否定' in manual['当前核对结论'] and '展开' in manual['当前核对结论']
        assert '记录恢复未对应' not in manual['当前核对结论']


def test_原人工逐笔先否定后成员由新组采用提示查看逐笔当前去向(excel,tmp_path):
    path=_报告(tmp_path,[100,200],[100,200],[((0,),(0,)),((1,),(1,))],[(1,(0,1),(0,1))])
    with _打开(excel,path) as book:
        _选择(excel,book,'M1','否定对应')
        _选择(excel,book,'M2','采用其他对应1')
        main=next(row for row in _表(book,'人工全查') if row['事项编号']=='M1' and row['行别']=='核对事项')
        assert '否定' in main['当前核对结论'] and '逐笔核对' in main['当前核对结论']
        assert '记录恢复未对应' not in main['当前核对结论'] and '展开' not in main['当前核对结论']
        pair=next(row for row in _表(book,'逐笔核对') if row['事项编号']=='M1')
        assert '002' in pair['当前核对结论'] and '已确认' in pair['当前核对结论']
