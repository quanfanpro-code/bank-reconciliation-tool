"""从真实报告出口验证阅读顺序、记录完整性和筛选结果。"""

import hashlib
from contextlib import closing

import pandas as pd
import pytest
from openpyxl import load_workbook

from reporter import Reporter
from tests.test_单账户业务匹配 import 记录, 核对
from 底稿筛选 import FilterCriteria, export_filtered_workpaper


def test_跨年小额在首页提示且一项业务只列一次():
    m = 核对([记录(3000, "服务费", 日期="2025-12-31")],
             [记录(3000, "服务费", 日期="2026-01-01")])
    tables = Reporter(m).build_report_tables(m.config)
    assert "核对概览" in tables
    overview = dict(zip(tables["核对概览"]["项目"], tables["核对概览"]["结果"]))
    assert overview["待复核事项数"] == 1
    assert overview["涉及跨期事项数"] == 1
    assert len(tables["待复核事项"]) == 1
    assert "期间" in tables["待复核事项"].to_string(index=False)
    assert "2025-12-31" in tables["待复核事项"].to_string(index=False)
    assert "2026-01-01" in tables["待复核事项"].to_string(index=False)


def test_整组组成完整且不暗示逐条对应():
    m = 核对([记录(-150, "结算", "甲公司", "P03"), 记录(-450, "结算", "甲公司", "P03")],
             [记录(-100, "结算", "甲公司", "P03"), 记录(-200, "结算", "甲公司", "P03"), 记录(-300, "结算", "甲公司", "P03")])
    tables = Reporter(m).build_report_tables(m.config)
    assert "全部核对明细" in tables
    details = tables["全部核对明细"]
    assert len(details.loc[details["行别"] == "业务合计"]) == 1
    rows = details.loc[details["行别"] == "组成记录"]
    assert len(rows) == 5
    assert pd.to_numeric(rows["银行金额"], errors="coerce").sum() == -600
    assert pd.to_numeric(rows["日记账金额"], errors="coerce").sum() == -600
    assert not ((rows["银行金额"] != "") & (rows["日记账金额"] != "")).any()
    assert all("第" in value and "行" in value for value in rows["原始出处"])


def test_跨期且有金额疑点合并为同一事项():
    m = 核对([记录(-90000, "结算", "甲公司", "P03", "2025-12-31")],
             [记录(-60000, "结算", "甲公司", "P03", "2026-01-01")], performance_materiality=10000)
    tables = Reporter(m).build_report_tables(m.config)
    assert "待复核事项" in tables
    assert len(tables["待复核事项"]) == 1
    text = tables["待复核事项"].to_string(index=False)
    assert "30,000" in text and "跨期" in text


def test_工作簿只有三主表可见且编号可跳转(tmp_path):
    m = 核对([记录(3000, 日期="2025-12-31")], [记录(3000, 日期="2026-01-01")])
    path = tmp_path / "报告.xlsx"
    Reporter(m).generate_report(str(path))
    with closing(load_workbook(path)) as wb:
        assert [ws.title for ws in wb if ws.sheet_state == "visible"] == ["核对概览", "待复核事项", "全部核对明细"]
        assert wb.active.title == "核对概览"
        ws = wb["待复核事项"]
        headers = {c.value: c.column for c in ws[1]}
        assert ws.cell(2, headers["事项编号"]).hyperlink is not None
        assert "全部核对明细" in ws.cell(2, headers["事项编号"]).hyperlink.target
        assert len(ws.data_validations.dataValidation) == 1
        for name in ("待复核事项", "全部核对明细"):
            assert not {"候选ID", "匹配ID", "综合可信度", "组成键"}.intersection(c.value for c in wb[name][1])


def test_新报告筛选整组后重算概览并保留全部组成(tmp_path):
    m = 核对([记录(-150, "结算", "甲公司", "P03"), 记录(-450, "结算", "甲公司", "P03"), 记录(-700, "房租", "乙公司", "P02")],
             [记录(-600, "结算", "甲公司", "P03"), 记录(-700, "房租", "乙公司", "P02")])
    source, output = tmp_path / "全量.xlsx", tmp_path / "筛选.xlsx"
    Reporter(m).generate_report(str(source))
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    export_filtered_workpaper(source, output, FilterCriteria(include_text=("甲公司",)))
    with closing(load_workbook(output)) as wb:
        assert [ws.title for ws in wb if ws.sheet_state == "visible"] == ["核对概览", "待复核事项", "全部核对明细"]
    overview = pd.read_excel(output, sheet_name="核对概览").fillna("")
    values = dict(zip(overview["项目"], overview["结果"]))
    assert values["银行记录笔数"] == 2
    assert values["日记账记录笔数"] == 1
    assert "甲公司" in str(values["筛选条件"])
    assert "FilterCriteria" not in str(values["筛选条件"])
    details = pd.read_excel(output, sheet_name="全部核对明细").fillna("")
    assert len(details.loc[details["行别"] == "组成记录"]) == 3
    assert "乙公司" not in details.to_string(index=False)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before


def test_空筛选不沿用全量数量(tmp_path):
    m = 核对([记录(300)], [记录(300)])
    source, output = tmp_path / "全量.xlsx", tmp_path / "空筛选.xlsx"
    Reporter(m).generate_report(str(source))
    export_filtered_workpaper(source, output, FilterCriteria(include_text=("不存在的内容",)))
    tables = pd.read_excel(output, sheet_name=None)
    assert "核对概览" in tables
    values = dict(zip(tables["核对概览"]["项目"], tables["核对概览"]["结果"]))
    assert values["银行记录笔数"] == 0
    assert values["待复核事项数"] == 0


def test_单边记录不能从全部明细中丢失():
    m = 核对([记录(-100, "支付设备款", "甲公司")], [记录(7700, "收到租金", "乙公司")])
    tables = Reporter(m).build_report_tables(m.config)
    assert "全部核对明细" in tables
    assert len(tables["待复核事项"]) == 2
    assert len(tables["全部核对明细"]) == 2
    assert set(tables["全部核对明细"]["核对结果"]) == {"尚未找到对应"}


def test_主表文字不能成为外部公式(tmp_path):
    m = 核对([记录(100, '=HYPERLINK("https://example.com")')], [记录(100, '=HYPERLINK("https://example.com")')])
    path = tmp_path / "文本.xlsx"
    Reporter(m).generate_report(str(path))
    with closing(load_workbook(path)) as wb:
        assert "全部核对明细" in wb.sheetnames
        for row in wb["全部核对明细"]:
            assert all(c.data_type != "f" for c in row)


def test_主表能直接看到对方名称和业务编号():
    m = 核对([记录(-1000, "货款", "甲公司", "P03")], [记录(-1000, "货款", "甲公司", "P03")])
    tables = Reporter(m).build_report_tables(m.config)
    text = tables["全部核对明细"].to_string(index=False)
    assert "甲公司" in text and "P03" in text


def test_低风险和未抽中中风险不进入主要复核清单():
    from tests.test_疑点分层处置 import _report_records
    from matcher import Matcher
    from data_structures import MatcherConfig
    rows = []
    for i, amount in enumerate((6000, 7000, 8000, 9000)):
        for side in ("bank", "journal", "journal"):
            rows.append((f"2026-03-{i + 1:02d}", amount, "服务费", side))
    for side in ("bank", "journal", "journal"):
        rows.append(("2026-04-01", 200000, "设备款", side))
    m = Matcher(_report_records([r for r in rows if r[3] == "bank"]),
                _report_records([r for r in rows if r[3] == "journal"]), MatcherConfig())
    m.run()
    tables = Reporter(m).build_report_tables(m.config)
    index = tables["阅读事项索引"]
    medium = index.loc[index["风险等级"] == "中风险"]
    assert len(medium) == 4
    assert len(medium.loc[medium["需要复核"] == "是"]) == 1
    assert len(medium.loc[medium["核对结果"] == "留存备查"]) == 3
    m = 核对([记录(-1000, "货款", "甲公司", "P03")], [记录(-990, "货款", "甲公司", "P03")])
    assert Reporter(m).build_report_tables(m.config)["待复核事项"].empty


def test_筛选不能经硬链接覆盖原报告(tmp_path):
    import os
    m = 核对([记录(300)], [记录(300)])
    source, alias = tmp_path / "全量.xlsx", tmp_path / "另一个名称.xlsx"
    Reporter(m).generate_report(str(source))
    os.link(source, alias)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    with pytest.raises(ValueError, match="不能覆盖"):
        export_filtered_workpaper(source, alias, FilterCriteria())
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before


def test_全局范围和余额问题可以从首页追到具体事项():
    from 报告阅读 import build_readable_tables
    m = 核对([记录(300)], [记录(300)])
    tables = Reporter(m).build_report_tables(m.config)
    summary = tables["核对结论"].copy()
    summary.loc[summary["项目"] == "核对范围", "数值"] = "范围受限"
    summary.loc[summary["项目"] == "范围说明", "数值"] = "核对币种尚未确认"
    tables["核对结论"] = summary
    tables["余额连续性异常"] = pd.DataFrame([{"来源": "银行流水", "日期": "2026-08-10", "差额": 200}])
    rebuilt = build_readable_tables(tables)
    assert len(rebuilt["待复核事项"]) == 2
    text = rebuilt["待复核事项"].to_string(index=False)
    assert "币种" in text and "200" in text


def test_重新构造阅读表不丢失单边记录():
    from 报告阅读 import build_readable_tables
    m = 核对([记录(-100, "支付设备款")], [记录(7700, "收到租金")])
    tables = Reporter(m).build_report_tables(m.config)
    rebuilt = build_readable_tables(tables)
    assert len(rebuilt["全部核对明细"]) == 2


def test_高风险金额差异不以自动归集作为主表结论():
    m = 核对([记录(-120000, "设备款", "甲公司", "P03")],
             [记录(-90000, "设备款", "甲公司", "P03")], performance_materiality=10000)
    tables = Reporter(m).build_report_tables(m.config)
    assert tables["全部核对明细"].iloc[0]["核对结果"] == "金额差异，需核查"
    assert tables["逐笔匹配"].iloc[0]["最终状态"] == "自动归集事项"


def test_筛选概览说明来源及实际金额覆盖比例(tmp_path):
    m = 核对([记录(800, "货款", "甲公司", "P01"), 记录(200, "租金", "乙公司", "P02")],
             [记录(800, "货款", "甲公司", "P01"), 记录(200, "租金", "乙公司", "P02")])
    source, output = tmp_path / "全量来源.xlsx", tmp_path / "筛选.xlsx"
    Reporter(m).generate_report(str(source))
    export_filtered_workpaper(source, output, FilterCriteria(coverage_ratio=0.5))
    overview = pd.read_excel(output, sheet_name="核对概览").fillna("")
    values = dict(zip(overview["项目"], overview["结果"]))
    assert "全量来源.xlsx" in values["全量报告"]
    assert "80.00%" in values["选中金额与比例"]
    assert "800.00" in values["选中金额与比例"]
