# -*- coding: utf-8 -*-
"""排版新路径接入：reporter._generate_report 默认走 排版计划+报告写入 单次写入，
设置 银行流水核对_旧排版=1 时回退旧的"写入后重开重排"路径。"""
import os
from contextlib import closing

from openpyxl import load_workbook

from tests.test_复核工作簿 import _报告


def test_默认走单次写入新路径(tmp_path, monkeypatch):
    import 报告写入
    原始 = 报告写入.渲染簿计划
    调用 = []

    def 间谍(计划, 路径):
        调用.append(计划)
        return 原始(计划, 路径)

    monkeypatch.setattr(报告写入, "渲染簿计划", 间谍)
    monkeypatch.delenv("银行流水核对_旧排版", raising=False)
    输出, _ = _报告(tmp_path)
    assert 调用, "默认路径未调用 报告写入.渲染簿计划"
    with closing(load_workbook(输出)) as wb:
        assert "核对结论" in wb.sheetnames
        assert "待处理事项" in wb.sheetnames
        assert "核对总览" in wb.sheetnames


def test_环境变量切回旧路径(tmp_path, monkeypatch):
    import 报告写入

    def 禁用(*args, **kwargs):
        raise AssertionError("旧排版开关下不应调用新路径")

    monkeypatch.setattr(报告写入, "渲染簿计划", 禁用)
    monkeypatch.setenv("银行流水核对_旧排版", "1")
    输出, _ = _报告(tmp_path)
    assert os.path.exists(输出)
    with closing(load_workbook(输出)) as wb:
        assert "核对结论" in wb.sheetnames
