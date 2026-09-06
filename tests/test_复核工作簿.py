"""验证复核表保留原始业务、可直接选择结果及月度数字。"""
from contextlib import closing
from decimal import Decimal

from openpyxl import load_workbook

from data_structures import MatcherConfig
from matcher import Matcher
from reporter import Reporter
from tests.test_第三阶段特殊业务 import _df


def _报告(tmp_path, 银行行=None, 账行=None):
    银行行 = 银行行 or [("2026-01-10", -200000, {"摘要": "设备款", "对方": "甲公司", "业务编号": "P01"})]
    账行 = 账行 or [("2026-01-10", 7000, {"摘要": "服务费", "对方": "乙公司", "业务编号": "S01"})]
    配置 = MatcherConfig()
    匹配 = Matcher(_df(银行行), _df(账行), 配置, logger=lambda _: None)
    匹配.run()
    报告 = Reporter(匹配)
    输出 = tmp_path / '核对.xlsx'
    报告.generate_report(str(输出), config=配置)
    return 输出, 报告


def test_月度明细全查抽样及原每日核对均可直接访问(tmp_path):
    输出, _ = _报告(tmp_path)
    with closing(load_workbook(输出)) as 簿:
        可见 = {表.title for 表 in 簿 if 表.sheet_state == 'visible'}
        assert {'核对结论', '月度核对', '核对明细', '人工全查', '人工抽样', '每日统计'} <= 可见
        assert 簿.active.title == '核对结论'


def test_全查在完整记录旁选择结果且不要求资料或调整(tmp_path):
    输出, _ = _报告(tmp_path)
    with closing(load_workbook(输出)) as 簿:
        表 = 簿['人工全查']
        表头 = [格.value for 格 in 表[1]]
        assert '人工核对结果' in 表头
        assert '银行金额' in 表头 and '序时账金额' in 表头
        assert any(格.value == -200000 for 行 in 表 for 格 in 行)
        assert not {'后续状态', '调整凭证号', '责任人', '处理日期'} & set(表头)
        assert list(表.data_validations.dataValidation)


def test_核对明细每个原始行一次且金额可加总(tmp_path):
    输出, _ = _报告(tmp_path)
    with closing(load_workbook(输出)) as 簿:
        表 = 簿['核对明细']
        表头 = {格.value: 格.column for 格 in 表[1]}
        来源 = 表头['来源']
        金额 = 表头['金额']
        银行 = [表.cell(行, 金额).value for 行 in range(2, 表.max_row+1) if 表.cell(行, 来源).value == '银行流水']
        账 = [表.cell(行, 金额).value for 行 in range(2, 表.max_row+1) if 表.cell(行, 来源).value == '序时账']
        assert 银行 == [-200000] and 账 == [7000]


def test_跨期对应仍分真实年月进入月度差异组成(tmp_path):
    字段 = {'摘要': '服务费', '对方': '甲公司', '业务编号': 'S01'}
    输出, 报告 = _报告(tmp_path, [('2025-12-31', 3000, 字段)], [('2026-01-01', 3000, 字段)])
    表 = 报告.build_report_tables(config=报告.matcher.config)['月度差异组成']
    assert set(表['月份']) == {'2025-12', '2026-01'}
    assert 表.groupby('月份')['收入差额'].sum().to_dict() == {'2025-12': Decimal('3000'), '2026-01': Decimal('-3000')}
    with closing(load_workbook(输出)) as 簿:
        assert 簿['截止性差异'].sheet_state == 'visible'


def test_两笔对三笔完整组成可在明细中查看(tmp_path):
    字段 = {'摘要': '设备款', '对方': '甲公司', '业务编号': 'P02'}
    输出, _ = _报告(tmp_path, [('2026-01-10', 值, 字段) for 值 in [-150, -450]], [('2026-01-10', 值, 字段) for 值 in [-100, -200, -300]])
    with closing(load_workbook(输出)) as 簿:
        表 = 簿['核对明细']
        表头 = {格.value: 格.column for 格 in 表[1]}
        金额 = [表.cell(行, 表头['金额']).value for 行 in range(2, 表.max_row+1)]
        assert sorted(金额) == [-450, -300, -200, -150, -100]
