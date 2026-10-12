"""只检查新的阅读布局、人工输入衔接与筛选输出。"""
from contextlib import closing
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from 底稿筛选 import FilterCriteria, export_filtered_workpaper, read_filter_report
from 报告列识别 import identify_columns
from 紧凑报告 import 阅读表, 阅读颜色, 留痕颜色
from tests.test_人工结果计算 import _报告, excel, _打开, _表


def test_五张蓝色阅读表与绿色完整留痕均可见(tmp_path):
    path = _报告(tmp_path, [150, 450, 500], [100, 200, 300, 500, 700],
                 [((0, 1), (0, 1, 2)), ((2,), (3,))])
    with closing(load_workbook(path)) as book:
        assert tuple(book.sheetnames[:5]) == 阅读表
        assert all(sheet.sheet_state == 'visible' for sheet in book)
        for sheet in book:
            assert sheet.sheet_properties.tabColor.rgb == '00' + (阅读颜色 if sheet.title in 阅读表 else 留痕颜色)
            assert sheet.freeze_panes in (None, 'A2')
            if sheet.max_row > 1:
                assert sheet.row_dimensions[2].height < 60
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        assert cell.font.name == 'MiSans' and cell.font.sz == 9
            first = min(identify_columns(sheet).values())
            assert all(sheet.column_dimensions[get_column_letter(column)].hidden for column in range(1, first))
        assert book['核对明细'].max_row == 9
        assert len(list(book['待处理事项'].values)[1:]) == 3
        assert '收入差额' in identify_columns(book['对应明细'])
        assert all(sheet.row_dimensions[2].height == 15 for sheet in book if sheet.title not in 阅读表 and sheet.max_row > 1)
        assert book['核对总览'].freeze_panes is None


def test_新人工入口保留下拉选项并且支持保存筛选(tmp_path):
    path = _报告(tmp_path, [100, 200], [100, 200], [((0,), (0,)), ((1,), (1,))])
    with closing(load_workbook(path)) as book:
        sheet = book['待处理事项']
        headers = identify_columns(sheet)
        sheet.cell(2, headers['人工核对结果'], '否定对应')
        sheet.cell(2, headers['备注'], '已检查原始记录')
        assert any(sheet.cell(2, headers['人工核对结果']).coordinate in validation.sqref
                   for validation in sheet.data_validations.dataValidation)
        assert all(not book[name].data_validations.dataValidation for name in ('人工全查', '人工抽样'))
        book.save(path)
    groups, _ = read_filter_report(path)
    assert groups.loc[groups['事项编号'] == 'M1', '最终状态'].iloc[0] == '否定对应'
    output = tmp_path / '筛选.xlsx'
    export_filtered_workpaper(path, output, FilterCriteria(statuses=('否定对应',)))
    with closing(load_workbook(output)) as book:
        sheet = book['待处理事项']
        headers = identify_columns(sheet)
        visible = [sheet.cell(row, headers['事项编号']).value for row in range(2, sheet.max_row + 1)
                   if not sheet.row_dimensions[row].hidden]
        assert visible == ['M1']
        assert sheet.cell(2, headers['人工核对结果']).value == '否定对应'
        assert sheet.cell(2, headers['备注']).value == '已检查原始记录'
        assert book['筛选说明'].sheet_properties.tabColor.rgb == '00' + 阅读颜色
        assert book['每日统计'].sheet_state == 'visible'


def test_缺少余额明确显示未提供而不补零(tmp_path):
    path = _报告(tmp_path, [100], [90], [((0,), (0,))])
    with closing(load_workbook(path)) as book:
        for name in ('月度核对', '每日统计'):
            sheet = book[name]
            headers = identify_columns(sheet)
            assert '需查看' in headers
            balances = [label for label in headers if label.endswith(('期末余额', '月末余额'))]
            assert len(balances) == 2
            assert all(not sheet.column_dimensions[get_column_letter(headers[label])].hidden for label in balances)
            assert all(sheet.cell(2, headers[label]).value in ('未提供', '缺失') for label in balances)
        pending = book['待处理事项']
        headers = identify_columns(pending)
        assert pending.cell(2, headers['收入差额']).value == 10
        assert pending.cell(2, headers['支出差额']).value == 0


def test_全部自动确认时仍有五张阅读表但不增加人工待办(tmp_path, monkeypatch):
    from data_structures import ProcessingStatus, RiskLevel
    from reporter import Reporter
    generate = Reporter.generate_report

    def confirmed(self, *args, **kwargs):
        for candidate in self.matcher.selected_candidates:
            candidate.processing_status = ProcessingStatus.AUTO_CONFIRMED
            candidate.risk_level = RiskLevel.NORMAL
            candidate.is_ambiguous = False
        return generate(self, *args, **kwargs)

    monkeypatch.setattr(Reporter, 'generate_report', confirmed)
    path = _报告(tmp_path, [100], [100], [((0,), (0,))])
    with closing(load_workbook(path)) as book:
        assert tuple(book.sheetnames[:5]) == 阅读表
        assert book['待处理事项'].max_row == 1
        assert not book['待处理事项'].data_validations.dataValidation
        assert book['对应明细'].max_row == 2


def test_长说明留在批注且大额金额列足够宽(tmp_path, monkeypatch):
    import tests.test_人工结果计算 as calculation
    frame = calculation._df

    def long_text(rows):
        result = frame(rows)
        result['summary'] = '设备采购付款说明' * 30
        return result

    monkeypatch.setattr(calculation, '_df', long_text)
    value = 1234567890123.45
    path = _报告(tmp_path, [value], [value], [((0,), (0,))])
    with closing(load_workbook(path)) as book:
        sheet = book['对应明细']
        headers = identify_columns(sheet)
        cell = sheet.cell(2, headers['银行记录'])
        assert len(str(cell.value)) <= 32
        assert cell.comment and len(cell.comment.text) > 100
        money = get_column_letter(headers['银行收入'])
        assert sheet.column_dimensions[money].width >= len(f'{value:,.2f}') + 1
        assert sheet.cell(2, headers['银行收入']).value == value


def test_输出XML的前缀引用均有声明(tmp_path):
    path = _报告(tmp_path, [100], [100], [((0,), (0,))])
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith('.xml'):
                continue
            declared = set()
            with archive.open(name) as stream:
                for event, value in ElementTree.iterparse(stream, events=('start-ns', 'start')):
                    if event == 'start-ns':
                        declared.add(value[0])
                    else:
                        ignorable = value.attrib.get('{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable', '')
                        assert set(ignorable.split()) <= declared, name


def _输入(excel, book, identity, result):
    sheet = book.Worksheets('待处理事项')
    values = sheet.UsedRange.Value
    headers = list(values[0])
    for row, values in enumerate(values[1:], 2):
        if values[headers.index('事项编号')] == identity:
            sheet.Cells(row, headers.index('人工核对结果') + 1).Value = result
            excel.CalculateFullRebuild()
            return
    raise AssertionError(identity)


def test_Excel中新的唯一输入仍驱动原有公式(excel, tmp_path):
    path = _报告(tmp_path, [100], [90], [((0,), (0,))])
    with _打开(excel, path) as book:
        assert _表(book, '待处理事项')[0]['当前结果'] == '待核对'
        _输入(excel, book, 'M1', '确认对应并保留差额')
        summary = {row['项目']: row['结果'] for row in _表(book, '核对总览')}
        assert summary['还需人工核对'] == 0
        assert summary['已完成人工核对'] == 1
        assert summary['银行金额已对平比例'] == 0
        assert summary['非工作日银行收付笔数'] == 1
        assert summary['非工作日银行收款金额'] == 100
        assert _表(book, '待处理事项')[0]['收入差额'] == 10
        assert _表(book, '待处理事项')[0]['当前结果'] == '已确认，差额保留'
        _输入(excel, book, 'M1', '否定对应')
        assert _表(book, '待处理事项')[0]['当前结果'] == '已取消对应'
        assert all(not row['当前对应'] for row in _表(book, '核对明细'))


def test_待处理表排序后仍按事项编号引用人工输入(excel, tmp_path):
    path = _报告(tmp_path, [100, 200], [100, 200], [((0,), (0,)), ((1,), (1,))])
    with closing(load_workbook(path)) as book:
        sheet = book['待处理事项']
        # 模拟按整行排序，公式相对位置由Excel处理；静态输入和编号随整行移动。
        headers = identify_columns(sheet)
        for label in ('事项编号', '人工核对结果'):
            a, b = sheet.cell(2, headers[label]), sheet.cell(3, headers[label])
            a.value, b.value = b.value, a.value
        book.save(path)
    with _打开(excel, path) as book:
        _输入(excel, book, 'M1', '否定对应')
        index = {row['事项编号']: row for row in _表(book, '复核事项索引')}
        assert index['M1']['人工核对结果'] == '否定对应'
        assert not index['M2']['人工核对结果']
