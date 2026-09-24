"""2026-09-24 全面独立复核新增边界用例。

独立于既有测试重新构造，从真实文件入口 run_reconciliation 端到端验证：
- 非工作日示警、首页指标、八张可见表、原子交付与源文件只读；
- 中风险等距抽样与高风险全查分流（跨月分散，避免月度累计升级干扰）；
- 同额竞争组不重叠选择，原记录不重复占用。
"""
import hashlib
from contextlib import closing
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from application import run_reconciliation
from data_structures import MatcherConfig


def _mapping():
    return {
        "date": "日期",
        "amount": "金额",
        "direction": "方向",
        "summary": "摘要",
        "account": "账号",
        "currency": "币种",
        "amount_basis": "原币",
        "auxiliary_text_columns": ["摘要"],
        "mode": "single_amount_with_direction",
    }


def _attach_account(rows):
    return [{**row, "账号": "1234567890", "币种": "人民币"} for row in rows]


def _write_pair(tmp_path, bank_rows, journal_rows):
    bank_rows, journal_rows = _attach_account(bank_rows), _attach_account(journal_rows)
    bank_path = tmp_path / "银行流水.xlsx"
    journal_path = tmp_path / "序时账.xlsx"
    pd.DataFrame(bank_rows).to_excel(bank_path, index=False)
    pd.DataFrame(journal_rows).to_excel(journal_path, index=False)
    return bank_path, journal_path


def _sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _sheet_rows(sheet):
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
    return headers, [
        {name: sheet.cell(row, col).value for name, col in headers.items()}
        for row in range(2, sheet.max_row + 1)
    ]


def test_端到端非工作日示警与原子交付(tmp_path):
    bank_rows = [
        {"日期": "2026-02-15", "金额": 100, "方向": "贷", "摘要": "春节收货款"},   # 官方放假日
        {"日期": "2026-03-07", "金额": 200, "方向": "借", "摘要": "周末付款"},     # 普通周末
        {"日期": "2026-02-14", "金额": 300, "方向": "借", "摘要": "补班日付款"},   # 调休上班，不示警
        {"日期": "2026-03-09", "金额": 1000, "方向": "贷", "摘要": "工作日收款"},  # 工作日
        {"日期": "2026-03-09", "金额": 0, "方向": "贷", "摘要": "零金额"},         # 零金额不示警
        {"日期": "2027-01-01", "金额": 400, "方向": "贷", "摘要": "日历外收款"},   # 日历未覆盖
    ]
    journal_rows = [
        {"日期": "2026-03-09", "金额": 100, "方向": "借", "摘要": "春节货款入账"},
        {"日期": "2026-03-07", "金额": 200, "方向": "贷", "摘要": "周末付款入账"},
        {"日期": "2026-02-14", "金额": 300, "方向": "贷", "摘要": "补班付款入账"},
        {"日期": "2026-03-09", "金额": 1000, "方向": "借", "摘要": "工作日收款入账"},
        {"日期": "2027-01-05", "金额": 400, "方向": "借", "摘要": "日历外入账"},
    ]
    bank_path, journal_path = _write_pair(tmp_path, bank_rows, journal_rows)
    bank_hash, journal_hash = _sha256(bank_path), _sha256(journal_path)
    output = tmp_path / "核对报告.xlsx"

    result = run_reconciliation(
        str(bank_path), str(journal_path), _mapping(), _mapping(), MatcherConfig(),
        output_path=output,
    )

    assert result == output and output.exists()
    with closing(load_workbook(output)) as book:
        visible = {name for name in book.sheetnames if book[name].sheet_state == "visible"}
        assert visible == {"核对结论", "非工作日交易", "月度核对", "逐笔核对", "整组核对",
                           "人工全查", "人工抽样", "未对应记录"}
        _, alerts = _sheet_rows(book["非工作日交易"])
        assert len(alerts) == 2
        assert {row["休息日类型"] for row in alerts} == {"官方放假日", "普通周末"}
        assert {row["收支方向"] for row in alerts} == {"收入", "支出"}
        assert sorted(row["金额"] for row in alerts) == [100, 200]
        assert sorted(row["原文件行号"] for row in alerts) == [2, 3]

        _, summary_rows = _sheet_rows(book["核对结论"])
        summary = {row["项目"]: row["数值"] for row in summary_rows}
        assert summary["非工作日银行收付笔数"] == 2
        assert summary["非工作日银行收款笔数"] == 1
        assert summary["非工作日银行收款金额"] == 100
        assert summary["非工作日银行付款笔数"] == 1
        assert summary["非工作日银行付款金额"] == 200
        assert summary["银行交易日历未覆盖笔数"] == 1
        # 人工事项统计由公式承接人工选择，报告内必须是活公式而非静态值
        assert str(summary["人工全查事项数"]).startswith("=COUNTIF")
        assert str(summary["人工抽样事项数"]).startswith("=COUNTIF")

        # 隐藏证据表保留完整原始明细，银行侧绝对金额合计与输入一致
        assert book["核对明细"].sheet_state == "hidden"
        _, details = _sheet_rows(book["核对明细"])
        bank_total = sum(abs(row["金额"]) for row in details if row["来源"] == "银行流水")
        assert bank_total == 2000

    # 源文件只读；同名报告拒绝覆盖；报告路径不得指向原始输入
    assert _sha256(bank_path) == bank_hash
    assert _sha256(journal_path) == journal_hash
    with pytest.raises(FileExistsError):
        run_reconciliation(str(bank_path), str(journal_path), _mapping(), _mapping(),
                           MatcherConfig(), output_path=output)
    with pytest.raises(ValueError, match="不能覆盖原始输入"):
        run_reconciliation(str(bank_path), str(journal_path), _mapping(), _mapping(),
                           MatcherConfig(), output_path=bank_path)


def test_中风险等距抽样与高风险全查(tmp_path):
    # 五个中风险单边（各6,000，超过明显微小临界值5,000、不超过重要性水平100,000）
    # 分布在不同月份，避免同一月度差异池累计升级干扰样本总体
    bank_rows = [
        {"日期": day, "金额": 6000, "方向": "贷", "摘要": f"中风险单边{index}"}
        for index, day in enumerate(
            ["2026-01-12", "2026-02-10", "2026-03-10", "2026-04-10", "2026-05-11"], 1)
    ]
    # 两个高风险单边（各120,000，超过重要性水平），同样分月
    bank_rows += [
        {"日期": "2026-06-10", "金额": 120000, "方向": "贷", "摘要": "高风险单边甲"},
        {"日期": "2026-07-10", "金额": 120000, "方向": "贷", "摘要": "高风险单边乙"},
    ]
    journal_rows = [
        {"日期": "2026-08-12", "金额": 777, "方向": "贷", "摘要": "无关序时账记录"},
    ]
    bank_path, journal_path = _write_pair(tmp_path, bank_rows, journal_rows)

    run_reconciliation(
        str(bank_path), str(journal_path), _mapping(), _mapping(), MatcherConfig(),
        output_path=tmp_path / "报告.xlsx",
    )

    with closing(load_workbook(tmp_path / "报告.xlsx")) as book:
        _, full = _sheet_rows(book["人工全查"])
        full_items = [row for row in full if row["行别"] == "核对事项"]
        assert len(full_items) == 2  # 两个高风险全部全查

        _, sampled = _sheet_rows(book["人工抽样"])
        sampled_items = [row for row in sampled if row["行别"] == "核对事项"]
        # 样本数 = min(中风险5项, 高风险2项) = 2
        assert len(sampled_items) == 2
        assert all("抽中" in row["核对原因"] for row in sampled_items)

        # 复核事项索引：5 个中风险中 2 个人工抽样、3 个留存备查
        _, index_rows = _sheet_rows(book["复核事项索引"])
        mediums = [row for row in index_rows if row["风险等级"] == "中风险"]
        assert len(mediums) == 5
        modes = [row["核对方式"] for row in mediums]
        assert modes.count("人工抽样") == 2
        assert modes.count("留存备查") == 3
        # 等距：5 项抽 2 项应抽中排序后第 1、3 项
        ordered = sorted(mediums, key=lambda row: (row["最早日期"], row["事项编号"]))
        assert [row["核对方式"] for row in ordered] == ["人工抽样", "留存备查", "人工抽样", "留存备查", "留存备查"]


def test_同额竞争组不重叠且不重复占用原记录(tmp_path):
    bank_rows = [
        {"日期": "2026-01-12", "金额": 500, "方向": "贷", "摘要": "收货款"},
        {"日期": "2026-01-13", "金额": 500, "方向": "贷", "摘要": "收货款"},
        {"日期": "2026-01-14", "金额": 500, "方向": "贷", "摘要": "收货款"},
    ]
    journal_rows = [
        {"日期": "2026-01-12", "金额": 500, "方向": "借", "摘要": "收货款"},
        {"日期": "2026-01-13", "金额": 500, "方向": "借", "摘要": "收货款"},
    ]
    bank_path, journal_path = _write_pair(tmp_path, bank_rows, journal_rows)

    run_reconciliation(
        str(bank_path), str(journal_path), _mapping(), _mapping(), MatcherConfig(),
        output_path=tmp_path / "报告.xlsx",
    )

    with closing(load_workbook(tmp_path / "报告.xlsx")) as book:
        _, unmatched = _sheet_rows(book["未对应记录"])
        assert len(unmatched) == 1  # 银行侧恰剩一笔未对应

        # 每条原始记录只进入一个复核事项：核对明细中原始记录号不得重复
        _, details = _sheet_rows(book["核对明细"])
        keys = [row["原始记录号"] for row in details]
        assert len(keys) == len(set(keys)) == 5


def test_方向列数据值不参与表头层级评分(tmp_path):
    """借贷混合的单层表头不得因方向值命中表头关键词而误判为多级表头。"""
    from data_loader import DataLoader
    rows = [
        {"日期": "2026-02-15", "金额": 100, "方向": "贷", "摘要": "春节收货款"},
        {"日期": "2026-03-07", "金额": 200, "方向": "借", "摘要": "周末付款"},
        {"日期": "2026-02-14", "金额": 300, "方向": "借", "摘要": "补班日付款"},
    ]
    path = tmp_path / "混合方向.xlsx"
    pd.DataFrame(rows).to_excel(path, index=False)
    structure = DataLoader().detect_table_structure(path)
    assert structure.header_rows == 1
    assert "日期" in structure.columns


def test_真实两级表头仍按两级识别(tmp_path):
    """修复只排除精确等于借/贷的单元格，真实复合表头识别不受影响。"""
    from data_loader import DataLoader
    frame = pd.DataFrame([
        ["日期", "金额", "金额", "摘要"],
        [None, "借方", "贷方", None],
        ["2026-01-12", 100, None, "收货款"],
        ["2026-01-13", None, 200, "付货款"],
    ])
    path = tmp_path / "两级表头.xlsx"
    frame.to_excel(path, index=False, header=False)
    structure = DataLoader().detect_table_structure(path)
    assert structure.header_rows == 2
