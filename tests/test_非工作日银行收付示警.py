"""银行流水非工作日示警的可观察行为。"""
from contextlib import closing
from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from data_structures import MatcherConfig
from matcher import Matcher
from reporter import Reporter
from tests.test_第三阶段特殊业务 import _df
from tests.test_人工结果计算 import _报告 as 生成报告
from 非工作日日历 import 判定日期


@pytest.mark.parametrize(
    "放假日,补班日,普通周末",
    [
        ("2021-02-12", "2021-02-07", "2021-03-06"),
        ("2022-02-01", "2022-01-29", "2022-03-12"),
        ("2023-01-24", "2023-01-28", "2023-03-04"),
        ("2024-02-10", "2024-02-18", "2024-03-09"),
        ("2025-01-28", "2025-01-26", "2025-03-08"),
        ("2026-02-15", "2026-02-14", "2026-03-07"),
    ],
)
def test_六年官方放假和调休优先于星期(放假日, 补班日, 普通周末):
    assert 判定日期(date.fromisoformat(放假日)) == ("官方放假日", "春节")
    assert 判定日期(date.fromisoformat(补班日)) == ("工作日", "")
    assert 判定日期(date.fromisoformat(普通周末)) == ("普通周末", "")


def test_年份越界不伪装成普通工作日():
    assert 判定日期(date(2027, 1, 1)) == ("未覆盖", "")
    assert 判定日期(date(2024, 2, 9)) == ("工作日", "")


def test_只扫描银行真实收付且保留已匹配记录():
    bank = _df([
        ("2026-02-15", 100, {"摘要": "收货款", "对方": "甲公司"}),
        ("2026-03-07", -200, {"摘要": "付款", "对方": "乙公司"}),
        ("2026-02-14", -300, {"摘要": "补班付款"}),
        ("2027-01-01", 400, {"摘要": "超范围收款"}),
        ("2026-03-08", 0, {"摘要": "零金额"}),
        ("2026-03-09", 50, {"摘要": "工作日收款"}),
    ])
    bank["matched"] = [True, False, False, False, False, True]
    bank["date_evidence"] = [
        {"date_column": "交易日期", "original_value": str(day.date())}
        for day in bank["date"]
    ]
    reporter = Reporter(SimpleNamespace(bank=bank, journal=_df([
        ("2026-03-07", 50, {"摘要": "周末补账"}),
    ])))

    table, uncovered = reporter._build_nonworkday_bank_table()

    assert uncovered == 1
    assert table["原文件行号"].tolist() == [2, 3]
    assert table["休息日类型"].tolist() == ["官方放假日", "普通周末"]
    assert table["收支方向"].tolist() == ["收入", "支出"]
    assert table["金额"].tolist() == [100, 200]
    assert table["原日期列名"].tolist() == ["交易日期", "交易日期"]
    assert "甲公司" in table.loc[0, "对方及辅助文字"]


def test_报告首页与银行示警清单按收支分别合计():
    bank = _df([
        ("2026-03-07", 100, {"摘要": "周末收款"}),
        ("2026-03-08", -200, {"摘要": "周末付款"}),
    ])
    journal = _df([
        ("2026-03-09", 100, {"摘要": "工作日补账"}),
        ("2026-03-09", -200, {"摘要": "工作日补账"}),
    ])
    config = MatcherConfig()
    matcher = Matcher(bank, journal, config, logger=lambda _: None)
    matcher.run()

    tables = Reporter(matcher).build_report_tables(config)

    alerts = tables["非工作日交易"]
    summary = dict(zip(tables["核对结论"]["项目"], tables["核对结论"]["数值"]))
    assert alerts["原文件行号"].tolist() == [2, 3]
    assert summary["非工作日银行收付笔数"] == 2
    assert summary["非工作日银行收款笔数"] == 1
    assert summary["非工作日银行收款金额"] == 100
    assert summary["非工作日银行付款笔数"] == 1
    assert summary["非工作日银行付款金额"] == 200
    assert summary["银行交易日历未覆盖笔数"] == 0


def test_没有示警时仍保留清单表头和零笔汇总():
    bank = _df([("2026-03-09", 100, {"摘要": "工作日收款"})])
    journal = _df([("2026-03-09", 100, {"摘要": "工作日入账"})])
    matcher = Matcher(bank, journal, MatcherConfig(), logger=lambda _: None)
    matcher.run()

    tables = Reporter(matcher).build_report_tables(matcher.config)
    assert tables["非工作日交易"].empty
    assert "原文件行号" in tables["非工作日交易"].columns
    summary = dict(zip(tables["核对结论"]["项目"], tables["核对结论"]["数值"]))
    assert summary["非工作日银行收付笔数"] == 0


def test_生成的工作簿直接显示银行节假日示警(tmp_path):
    path = 生成报告(tmp_path, [100], [100], [((0,), (0,))])

    with closing(load_workbook(path)) as book:
        sheet = book["非工作日交易"]
        assert sheet.sheet_state == "visible"
        headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
        assert sheet.max_row == 2
        assert sheet.cell(2, headers["原文件行号"]).value == 2
        assert sheet.cell(2, headers["交易日期"]).value.date() == date(2026, 1, 10)
        assert sheet.cell(2, headers["休息日类型"]).value == "普通周末"
        home = book["核对结论"]
        home_headers = {cell.value: cell.column for cell in home[1] if cell.value}
        summary = {
            row[home_headers["项目"] - 1].value: row[home_headers["数值"] - 1].value
            for row in home.iter_rows(min_row=2)
        }
        assert summary["非工作日银行收付笔数"] == 1
        visible_rows = {
            home.cell(row, home_headers["项目"]).value: not home.row_dimensions[row].hidden
            for row in range(2, home.max_row + 1)
        }
        assert visible_rows["非工作日银行收付笔数"]
        assert visible_rows["银行交易日历未覆盖笔数"]
