# -*- coding: utf-8 -*-
"""紧凑层回放：紧凑报告.apply_compact_report。

布局（前置空列=1，除新建三张表为 0）：
复核事项索引 B核对编号 C事项编号 D核对方式 E当前核对结论 F判断依据 G跨期 H风险等级 I入选原因
核对明细 B核对编号 C事项编号 D来源 E日期 F摘要 G对方及凭证 H金额
人工全查 B行别 C事项编号 D人工核对结果 E备注 F填报编号
人工抽样 同人工全查（空表）
逐笔核对 B核对编号 C事项编号
整组核对 B核对编号 C事项编号
未对应记录 B核对编号 C事项编号
每日统计 B日期 C收入金额差额 D支出金额差额
月度核对 B月份 C银行-月末余额 D序时账-月末余额 E余额差额 F收入金额差额 G支出金额差额 H余额衔接差
余额连续性异常 B日期
核对结论 B项目 C数值
"""
import pandas as pd
import pytest
from openpyxl import load_workbook

from 排版计划 import 新建簿计划

_替换表 = (
    ("所选对应重复使用原记录", "重复占用，请重选"),
    ("原对应已否定，记录恢复未对应", "已取消对应"),
    ("原对应被人工选择替代，记录未对应", "原对应已被替代"),
    ("已人工核对，保留未对应", "已查看，仍未对应"),
    ("已确认所选对应；原始差额及期间保留", "已确认，差额保留"),
    ("部分原记录已由其他事项建立对应，其余仍未对应", "部分记录已对应"),
    ("原记录已由其他事项建立对应", "已由其他事项对应"),
    ("程序已确认对应", "程序已对上"),
    ("待人工核对", "待核对"),
    ("留存备查，未作人工确认", "留存查看"),
    ("选择无效", "选择无效，请重选"),
)


def _缩短(expression):
    for old, new in _替换表:
        expression = f'SUBSTITUTE({expression},"{old}","{new}")'
    return "=" + expression


def _紧凑簿():
    tables = {
        "复核事项索引": pd.DataFrame([
            ["001", "I1", "人工全查", "", "", "", "", ""],
            ["002", "I2", "留存备查", "", "", "", "", ""],
        ], columns=["核对编号", "事项编号", "核对方式", "当前核对结论", "判断依据", "跨期", "风险等级", "入选原因"]),
        "核对明细": pd.DataFrame([
            ["001", "I1", "银行流水", "2026-01-05", "货款", "甲公司", 100.0],
            ["001", "I1", "序时账", "2026-01-05", "货款", "甲公司", 100.0],
            ["002", "I2", "银行流水", "2026-01-06", "服务费", "乙公司", 50.0],
        ], columns=["核对编号", "事项编号", "来源", "日期", "摘要", "对方及凭证", "金额"]),
        "人工全查": pd.DataFrame([["核对事项", "I1", "", "", "I1"]],
            columns=["行别", "事项编号", "人工核对结果", "备注", "填报编号"]),
        "人工抽样": pd.DataFrame(columns=["行别", "事项编号", "人工核对结果", "备注", "填报编号"]),
        "逐笔核对": pd.DataFrame([["001", "I1"]], columns=["核对编号", "事项编号"]),
        "整组核对": pd.DataFrame([["003", "X9"]], columns=["核对编号", "事项编号"]),
        "未对应记录": pd.DataFrame([["004", "X8"]], columns=["核对编号", "事项编号"]),
        "每日统计": pd.DataFrame([["2026-01-05", 1.0, 0.0]],
            columns=["日期", "收入金额差额", "支出金额差额"]),
        "月度核对": pd.DataFrame([
            ["2026-02", 100.0, 100.0, 0.0, 0.0, 0.0, 0.0],
            ["2026年合计", None, None, None, 0.0, 0.0, None],
        ], columns=["月份", "银行-月末余额", "序时账-月末余额", "余额差额",
                    "收入金额差额", "支出金额差额", "余额衔接差"]),
        "余额连续性异常": pd.DataFrame([["2026-01-05"]], columns=["日期"]),
        "核对结论": pd.DataFrame([
            ["核对范围", "范围可用"],
            ["范围说明", "未发现范围疑点"],
            ["银行有效交易笔数", 10],
            ["余额核对", "余额核对一致"],
        ], columns=["项目", "数值"]),
    }
    簿 = 新建簿计划(tables)
    for tp in 簿.表:
        tp.前置空列 = 1
    # 复核层留下的行状态与人工表验证，供紧凑层搬迁/清除
    簿.取表("人工全查").验证.append({
        "type": "list", "formula1": '"确认对应,否定对应"', "allow_blank": True,
        "errorTitle": "请选择本事项的结果", "error": "请选择下拉列表中的对应结果。",
        "showErrorMessage": True, "errorStyle": "stop", "sqref": "D2"})
    整组 = 簿.取表("整组核对")
    整组.行计划取(2).outline = 1
    整组.行计划取(2).隐藏 = True
    整组.行计划取(2).收起 = True
    return 簿


def _值(tp, 行, 列名):
    import math
    v = tp.表.iloc[行 - 2, list(tp.表.columns).index(列名)]
    if v is None or v is pd.NaT or (isinstance(v, float) and math.isnan(v)):
        return None
    return v


@pytest.fixture
def 簿():
    from 排版计划 import 应用紧凑层
    簿 = _紧凑簿()
    应用紧凑层(簿)
    return 簿


def test_待处理事项内容(簿):
    p = 簿.取表("待处理事项")
    assert p.前置空列 == 0
    assert list(p.表.columns)[:3] == ["核对编号", "处理要求", "哪里需要查看"]
    # 排序：人工全查在前，留存备查在后
    assert _值(p, 2, "事项编号") == "I1" and _值(p, 3, "事项编号") == "I2"
    assert _值(p, 2, "处理要求") == "逐项核对" and _值(p, 3, "处理要求") == "留存查看"
    assert _值(p, 2, "哪里需要查看") == "金额相同，对应关系仍需确认"
    assert _值(p, 2, "银行记录") == "2026-01-05 1笔 货款；甲公司"
    # 当前结果：缩短公式（11 层 SUBSTITUTE 包 INDEX/MATCH）
    lookup = "IFERROR(INDEX('复核事项索引'!$E$2:$E$3,MATCH($L2,'复核事项索引'!$C$2:$C$3,0)),\"\")"
    assert _值(p, 2, "当前结果") == _缩短(lookup)
    # 人工事项：结果格取原值（空→None），验证搬到 H2
    assert _值(p, 2, "人工核对结果") is None
    assert len(p.验证) == 1 and p.验证[0]["sqref"] == "H2"
    assert p.验证[0]["errorStyle"] == "stop"
    # 非人工事项：无需填写，无验证
    assert _值(p, 3, "人工核对结果") == "本次无需填写"
    # 完整明细超链接
    assert p.覆盖[(2, 11)].超链接 == "#'人工全查'!A2"
    assert p.覆盖[(3, 11)].超链接 == "#'核对明细'!A4"
    # 批注：原判断依据
    assert p.覆盖[(2, 3)].批注 == ("；", "原判断依据")


def test_人工表入口回收(簿):
    全查 = 簿.取表("人工全查")
    assert 全查.验证 == []
    v = "IFERROR(INDEX('待处理事项'!$H$2:$H$3,MATCH(F2,'待处理事项'!$L$2:$L$3,0)),\"\")"
    assert _值(全查, 2, "人工核对结果") == f'=IF({v}="","",{v})'
    v2 = "IFERROR(INDEX('待处理事项'!$J$2:$J$3,MATCH(F2,'待处理事项'!$L$2:$L$3,0)),\"\")"
    assert _值(全查, 2, "备注") == f'=IF({v2}="","",{v2})'


def test_对应明细与编号重指(簿):
    m = 簿.取表("对应明细")
    assert _值(m, 2, "对应方式") == "逐笔对应"
    assert _值(m, 2, "银行收入") == 100.0
    assert _值(m, 2, "核对编号") == (
        "=HYPERLINK(\"#'待处理事项'!H\"&(MATCH($N2,'待处理事项'!$L$2:$L$3,0)+1),\"001\")")
    assert _值(m, 2, "完整明细") == (
        "=HYPERLINK(\"#'逐笔核对'!A\"&(MATCH($N2,'逐笔核对'!$C$2:$C$2,0)+1),\"查看\")")
    assert _值(簿.取表("逐笔核对"), 2, "核对编号") == (
        "=HYPERLINK(\"#'待处理事项'!H\"&(MATCH(C2,'待处理事项'!$L$2:$L$3,0)+1),\"001\")")
    # 不在待处理的事项不重指
    assert _值(簿.取表("整组核对"), 2, "核对编号") == "003"


def test_期间标记(簿):
    每日 = 簿.取表("每日统计")
    cols = list(每日.表.columns)
    assert cols[3:6] == ["银行-期末余额", "日记账-期末余额", "余额差额"]
    assert cols[6] == "需查看"
    assert 每日.表["银行-期末余额"].tolist() == ["未提供"]
    # 异常日期直接写"余额不连续"
    assert 每日.表["需查看"].tolist() == ["余额不连续"]
    月度 = 簿.取表("月度核对")
    assert list(月度.表.columns)[-1] == "需查看"
    assert _值(月度, 2, "需查看") == (
        '=IF(OR(AND(ISNUMBER(F2),F2<>0),AND(ISNUMBER(G2),G2<>0),'
        'AND(ISNUMBER(E2),E2<>0),AND(ISNUMBER(H2),H2<>0)),"有差额",'
        'IF(OR(NOT(ISNUMBER(C2)),NOT(ISNUMBER(D2))),"余额未提供",""))')
    # 合计行跳过
    assert _值(月度, 3, "需查看") is None


def test_核对总览(簿):
    o = 簿.取表("核对总览")
    assert list(o.表.columns) == ["项目", "结果", "怎么看"]
    base = "IFERROR(INDEX('核对结论'!$C$2:$C$5,MATCH(\"核对范围\",'核对结论'!$B$2:$B$5,0)),\"未提供\")"
    for old, new in (("范围受限", "资料有缺口，部分结果需核实"),
                     ("范围可用", "本次资料可以用于核对"), ("范围未验证", "尚未确认资料范围")):
        base = f'SUBSTITUTE({base},"{old}","{new}")'
    assert _值(o, 2, "结果") == "=" + base
    lookup = "IFERROR(INDEX('核对结论'!$C$2:$C$5,MATCH(\"范围说明\",'核对结论'!$B$2:$B$5,0)),\"未提供\")"
    assert _值(o, 3, "结果") == (
        f'=IF({lookup}="未发现范围疑点","未发现额外资料问题","有需要确认的资料，见原始说明")')
    assert _值(o, 3, "怎么看") == "未发现范围疑点"
    # 百分比行数字格式
    比例行 = o.表["项目"].tolist().index("银行金额已对平比例") + 2
    assert o.覆盖[(比例行, 2)].数字格式 == "0.00%"
    # 两张期间表的 COUNTIF 行 + 超链接
    行号 = o.表["项目"].tolist().index("有差额的月份") + 2
    assert _值(o, 行号, "结果") == (
        "=COUNTIF('月度核对'!$I$2:$I$3,\"有差额\")+COUNTIF('月度核对'!$I$2:$I$3,\"余额不连续\")")
    assert o.覆盖[(行号, 3)].超链接 == "#'月度核对'!A1"
    行号2 = o.表["项目"].tolist().index("有差额或余额不连续的日期") + 2
    assert "每日统计'!$H$2:$H$2" in _值(o, 行号2, "结果")
    assert _值(o, len(o.表) + 1, "项目") == "使用方法"


def test_紧凑样式与渲染(tmp_path):
    from 排版计划 import 应用紧凑层
    from 报告写入 import 渲染簿计划

    簿 = _紧凑簿()
    应用紧凑层(簿)
    名称序 = [t.名称 for t in 簿.表]
    assert 名称序[:5] == ["核对总览", "待处理事项", "月度核对", "每日统计", "对应明细"]
    assert 簿.active == 0
    # 整组核对（非阅读表）：隐藏/收起被清除，outline 保留
    整组 = 簿.取表("整组核对")
    assert 整组.行[2].隐藏 is False and 整组.行[2].收起 is False and 整组.行[2].outline == 1
    # 条件格式
    p = 簿.取表("待处理事项")
    区域集 = sorted(区域 for 区域, _ in p.条件格式)
    assert 区域集 == ["F2:F3", "G2:G3"]

    输出 = tmp_path / "出.xlsx"
    渲染簿计划(簿, str(输出))
    wb = load_workbook(输出)
    assert wb.sheetnames[:5] == ["核对总览", "待处理事项", "月度核对", "每日统计", "对应明细"]
    ws = wb["待处理事项"]
    assert ws.sheet_properties.tabColor.rgb == "00BDD7EE" or ws.sheet_properties.tabColor.rgb == "FFBDD7EE"
    assert ws.freeze_panes == "A2"
    assert ws.sheet_format.defaultRowHeight == 15
    assert ws.row_dimensions[1].height == 28
    a2 = ws["A2"]
    assert a2.font.name == "MiSans" and a2.font.size == 9
    assert a2.fill.fill_type is None            # 填充被清空
    assert a2.alignment.horizontal == "left"    # 文本左对齐
    d2 = ws["D2"]
    assert d2.alignment.horizontal == "left"    # 银行记录为文本左对齐
    assert ws["F2"].alignment.horizontal == "right"   # 数字右对齐
    assert ws["F2"].number_format == "#,##0.00;[Red]-#,##0.00"
    # 黄色格在紧凑清填充之后仍保留
    assert ws["H2"].fill.fgColor.rgb == "00FFF2CC" or ws["H2"].fill.fgColor.rgb == "FFFFF2CC"
    # 超链接格蓝字下划线
    k2 = ws["K2"]
    assert k2.font.color.rgb[-6:] == "0563C1" and k2.font.underline == "single"
    assert k2.hyperlink is not None
    # 表头
    a1 = ws["A1"]
    assert a1.font.name == "MiSans" and a1.font.bold and a1.fill.fgColor.rgb[-6:] == "BDD7EE"
    # 核对总览不冻结
    assert wb["核对总览"].freeze_panes is None
    # 事项编号列隐藏
    assert ws.column_dimensions["L"].hidden


def test_紧凑新建表空格物化(tmp_path):
    """紧凑层新建表（待处理事项/对应明细/核对总览）旧路径在重载后直接建表：
    空串格也物化并带紧凑样式。"""
    from 排版计划 import 应用紧凑层, 表计划
    from 报告写入 import 渲染簿计划

    簿 = _紧凑簿()
    簿.表.append(表计划(名称="非工作日交易", 表=pd.DataFrame([["2026-01-04"]], columns=["交易日期"])))
    应用紧凑层(簿)
    assert 簿.取表("待处理事项").物化空格
    assert 簿.取表("对应明细").物化空格
    assert 簿.取表("核对总览").物化空格
    输出 = tmp_path / "出.xlsx"
    渲染簿计划(簿, str(输出))
    wb = load_workbook(输出)
    ws = wb["待处理事项"]
    # H2 人工核对结果、J2 备注均为 None：旧路径落格为带紧凑样式的空格
    for coord in ("H2", "J2"):
        cell = ws[coord]
        assert cell.value is None
        assert cell.font.name == "MiSans" and cell.font.size == 9
        assert cell.alignment.horizontal == "left" and cell.alignment.vertical == "top"
    assert ws["H2"].fill.fgColor.rgb[-6:] == "FFF2CC"   # 黄格后覆盖仍在
    # 核对总览：非工作日行的"怎么看"为空串，同样物化
    总览 = wb["核对总览"]
    项目 = [总览.cell(r, 1).value for r in range(1, 总览.max_row + 1)]
    行 = 项目.index("非工作日银行收款笔数") + 1
    cell = 总览.cell(行, 3)
    assert cell.value is None and cell.font.name == "MiSans"


def test_期间标记追加列边框中和(tmp_path):
    """期间标记追加列是紧凑层后建新列：不背复核层行样式/表头边框。"""
    from 排版计划 import 应用紧凑层, 格样式
    from 报告写入 import 渲染簿计划

    簿 = _紧凑簿()
    每日 = 簿.取表("每日统计")
    每日.表头样式 = 格样式(上框=("medium", "000000"), 下框=("dashed", "808080"))
    每日.行计划取(2).样式 = 格样式(上框=("dashed", "808080"), 下框=("dashed", "808080"))
    应用紧凑层(簿)
    输出 = tmp_path / "出.xlsx"
    渲染簿计划(簿, str(输出))
    wb = load_workbook(输出)
    ws = wb["每日统计"]
    # 追加列（E 银行-期末余额、H 需查看）表头与数据格无边框
    for coord in ("E1", "H1", "E2", "H2"):
        边 = ws[coord].border
        assert 边.top.style is None and 边.bottom.style is None, coord
    # 非追加列保持行样式边框
    assert ws["B2"].border.top.style == "dashed"
    # 月度核对合计行的需查看格旧路径不存在（空缺格不落格）
    月 = wb["月度核对"]
    assert 月["I3"].value is None and not 月["I3"].has_style
