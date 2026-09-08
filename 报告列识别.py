"""读取本工具报告的列名，列位置只用于定位已识别字段的单元格。"""
from openpyxl.utils import get_column_letter


def identify_columns(sheet, required=()):
    """报告第一行为表头；允许换序、额外列和两端空格，不猜测歧义列。"""
    columns = {}
    for cell in sheet[1]:
        label = str(cell.value).strip() if cell.value is not None else ''
        if not label:
            continue
        if label in columns:
            positions = f'{get_column_letter(columns[label])}、{get_column_letter(cell.column)}'
            raise ValueError(f'工作表“{sheet.title}”存在重复列名“{label}”（{positions}列），请区分列名后再操作')
        columns[label] = cell.column
    missing = [name for name in required if name not in columns]
    if missing:
        raise ValueError(f'工作表“{sheet.title}”缺少必要列：{"、".join(missing)}；请恢复列名或使用完整报告')
    return columns
