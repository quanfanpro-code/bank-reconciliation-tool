# -*- coding: utf-8 -*-
"""报告写入：把 排版计划.簿计划 单次渲染为 xlsx（openpyxl 普通模式，保存一次）。

样式对象按属性元组缓存复用，不逐格新建 Font/Fill。
解析样式优先级：列基础样式 ← 行样式 ← 格覆盖（后者优先）。
"""
import math

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from 排版计划 import 格样式, 合并样式, 簿计划

_字体缓存 = {}
_填充缓存 = {}
_对齐缓存 = {}
_边框缓存 = {}


def _字体(s: 格样式) -> Font | None:
    key = (s.字体名, s.字号, s.加粗, s.字色, s.下划线)
    if key == (None,) * 5:
        return None
    if key not in _字体缓存:
        # 空字符串表示"显式清空"（旧路径整体替换 Font 时 name/size 回落默认）
        _字体缓存[key] = Font(name=s.字体名 or None, size=s.字号 or None,
                              bold=s.加粗, color=s.字色, underline=s.下划线)
    return _字体缓存[key]


def _填充(s: 格样式) -> PatternFill | None:
    key = (s.填充类型, s.填充色)
    if key == (None, None):
        return None
    if key not in _填充缓存:
        if s.填充类型 == "none":
            _填充缓存[key] = PatternFill(fill_type=None)
        else:
            _填充缓存[key] = PatternFill(fill_type=s.填充类型 or "solid",
                                         fgColor=s.填充色 or "00000000")
    return _填充缓存[key]


def _对齐(s: 格样式) -> Alignment | None:
    key = (s.水平, s.垂直, s.自动换行)
    if key == (None,) * 3:
        return None
    if key not in _对齐缓存:
        # 空字符串表示"显式清空"（旧路径整体替换 Alignment 时会清掉 horizontal）
        _对齐缓存[key] = Alignment(horizontal=s.水平 or None, vertical=s.垂直 or None,
                                   wrap_text=s.自动换行)
    return _对齐缓存[key]


def _边(框: tuple | None) -> Side:
    if 框 is None:
        return Side(style=None)
    style, color = 框
    # style 为空字符串表示"显式清除该边"（旧路径整体替换 Border 的效果）
    return Side(style=style or None, color=color)


def _边框(s: 格样式) -> Border | None:
    key = (s.上框, s.下框, s.左框, s.右框)
    if key == (None,) * 4:
        return None
    if key not in _边框缓存:
        _边框缓存[key] = Border(top=_边(s.上框), bottom=_边(s.下框),
                                left=_边(s.左框), right=_边(s.右框))
    return _边框缓存[key]


def _套样式(cell, s: 格样式):
    if s is None:
        return
    f = _字体(s)
    if f is not None:
        cell.font = f
    p = _填充(s)
    if p is not None:
        cell.fill = p
    a = _对齐(s)
    if a is not None:
        cell.alignment = a
    b = _边框(s)
    if b is not None:
        cell.border = b
    if s.数字格式 is not None:
        cell.number_format = s.数字格式
    if s.批注 is not None:
        cell.comment = Comment(s.批注[0], s.批注[1])
    if s.超链接 is not None:
        cell.hyperlink = s.超链接


def _值转换(v):
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if v is pd.NaT:
        return None
    return v


def _渲染表(wb: Workbook, tp):
    ws = wb.create_sheet(tp.名称)
    if tp.tab色:
        ws.sheet_properties.tabColor = tp.tab色
    if tp.状态 != "visible":
        ws.sheet_state = tp.状态
    if tp.冻结:
        ws.freeze_panes = tp.冻结
    if tp.筛选:
        ws.auto_filter.ref = tp.筛选
    if not tp.网格线:
        ws.sheet_view.showGridLines = False
    if tp.缩放 is not None:
        ws.sheet_view.zoomScale = tp.缩放
    if tp.左上角:
        ws.sheet_view.topLeftCell = tp.左上角
    if tp.outline汇总在下 is not None:
        ws.sheet_properties.outlinePr.summaryBelow = tp.outline汇总在下
    if tp.显示outline符号 is not None:
        ws.sheet_view.showOutlineSymbols = tp.显示outline符号

    偏 = tp.前置空列
    # 前置空列（A 列等）的宽与隐藏
    for k, cp in enumerate(tp.前部列):
        dim = ws.column_dimensions[get_column_letter(k + 1)]
        if cp.宽 is not None:
            dim.width = cp.宽
        if cp.隐藏:
            dim.hidden = True
    # 列宽与隐藏
    for i, cp in enumerate(tp.列):
        letter = get_column_letter(偏 + i + 1)
        dim = ws.column_dimensions[letter]
        if cp.宽 is not None:
            dim.width = cp.宽
        if cp.隐藏:
            dim.hidden = True

    # 表头行（第 1 行）；覆盖对表头同样生效（如表头批注）
    if tp.表头含前列:
        for k in range(偏):
            cell = ws.cell(row=1, column=k + 1)
            _套样式(cell, 合并样式(tp.表头样式, tp.覆盖.get((1, k + 1))))
    for j, 列名 in enumerate(tp.表.columns):
        cell = ws.cell(row=1, column=偏 + j + 1, value=列名)
        _套样式(cell, 合并样式(tp.表头样式, tp.覆盖.get((1, 偏 + j + 1))))

    # 数据行
    行样式 = {行号: rp.样式 for 行号, rp in tp.行.items() if rp.样式 is not None}
    记录 = tp.表.itertuples(index=False, name=None)
    for i, row in enumerate(记录):
        行号 = i + 2
        rs = 行样式.get(行号)
        # 前置空列：仅在有样式要求时落格（复刻旧路径产生的带样式空格）
        for k in range(偏):
            ps = tp.前部列[k].基础样式 if k < len(tp.前部列) else None
            样式 = 合并样式(ps, rs)
            if 样式 is not None:
                _套样式(ws.cell(row=行号, column=k + 1), 样式)
        for j, v in enumerate(row):
            列号 = 偏 + j + 1
            cell = ws.cell(row=行号, column=列号, value=_值转换(v))
            s = tp.覆盖.get((行号, 列号))
            样式 = None
            if j < len(tp.列):
                样式 = tp.列[j].基础样式
            if rs is not None:
                样式 = 合并样式(样式, rs)
            if s is not None:
                样式 = 合并样式(样式, s)
            if 样式 is not None:
                _套样式(cell, 样式)

    # 行计划
    for 行号, rp in tp.行.items():
        dim = ws.row_dimensions[行号]
        if rp.高 is not None:
            dim.height = rp.高
        if rp.隐藏:
            dim.hidden = True
        if rp.outline:
            dim.outlineLevel = rp.outline
        if rp.收起:
            dim.collapsed = True

    # 数据验证
    for kwargs in tp.验证:
        kwargs = dict(kwargs)
        sqref = kwargs.pop("sqref")
        dv = DataValidation(**kwargs)
        ws.add_data_validation(dv)
        dv.add(sqref)

    # 条件格式
    for 区域, 规则 in tp.条件格式:
        ws.conditional_formatting.add(区域, 规则)

    # 打印
    p = tp.打印
    if p:
        if p.get("orientation"):
            ws.page_setup.orientation = p["orientation"]
        if p.get("paperSize") is not None:
            ws.page_setup.paperSize = p["paperSize"]
        if p.get("fitToWidth") is not None:
            ws.page_setup.fitToWidth = p["fitToWidth"]
        if p.get("fitToHeight") is not None:
            ws.page_setup.fitToHeight = p["fitToHeight"]
        if p.get("fitToPage"):
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        if p.get("title_rows"):
            ws.print_title_rows = p["title_rows"]
        if p.get("area"):
            ws.print_area = p["area"]
        if p.get("左边距") is not None:
            ws.page_margins.left = p["左边距"]
        if p.get("右边距") is not None:
            ws.page_margins.right = p["右边距"]


def 渲染簿计划(计划: 簿计划, 输出路径: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for tp in 计划.表:
        _渲染表(wb, tp)
    wb.active = 计划.active
    if 计划.自动重算:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    wb.save(输出路径)
