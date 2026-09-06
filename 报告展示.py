"""只组织已有核对结果的阅读展示；不参与匹配、风险分流或结果计算。"""
from collections import defaultdict
from decimal import Decimal

import pandas as pd
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


阅读表 = ('核对结论', '月度核对', '逐笔核对', '整组核对', '人工全查', '人工抽样', '未对应记录')
组成金额列 = ('银行收入', '银行支出', '序时账收入', '序时账支出')
组展示列 = ['核对编号', '行别', '银行记录', '银行收入', '银行支出', '序时账记录', '序时账收入', '序时账支出', '核对差额', '人工核对结果', '当前核对结论', '当前记录结果', '核对原因', '备注', '原始出处', '事项编号', '原始记录号', '填报编号', '银行金额', '序时账金额', '差额']


def _文字(record):
    return '\n'.join(str(x) for x in (pd.Timestamp(record['日期']).strftime('%Y-%m-%d'), record['摘要'], record['对方及凭证']) if x)


def _收支(records):
    amounts = [Decimal(str(record['金额'])) for record in records]
    return sum((max(x, Decimal(0)) for x in amounts), Decimal(0)), sum((max(-x, Decimal(0)) for x in amounts), Decimal(0))


def _汇总(records):
    if not records:
        return '未找到对应记录'
    income, expense = _收支(records)
    dates = sorted({pd.Timestamp(record['日期']).strftime('%Y-%m-%d') for record in records})
    period = dates[0] if len(dates) == 1 else dates[0] + '至' + dates[-1]
    summaries = list(dict.fromkeys(str(record['摘要']) for record in records if record['摘要']))
    description = '、'.join(summaries[:2]) + (f'等{len(summaries)}种摘要' if len(summaries) > 2 else '')
    return f'{period}\n{description}\n{len(records)}笔\n收入 {income:,.2f}\n支出 {expense:,.2f}'


def _金额(row, side, record):
    amount = float(record['金额'])
    row[side + '收入'] = max(amount, 0)
    row[side + '支出'] = max(-amount, 0)


def _组成(item, records):
    rows = []
    for record in records:
        side = '银行' if record['来源'] == '银行流水' else '序时账'
        row = {'核对编号': item['核对编号'], '事项编号': item['事项编号'], '原始记录号': record['原始记录号'], '行别': side + '组成', side + '记录': _文字(record), side + '金额': record['金额'], '原始出处': record['原始出处']}
        _金额(row, side, record)
        rows.append(row)
    return rows


def _主行(item, records, original=None):
    banks = [r for r in records if r['来源'] == '银行流水']
    journals = [r for r in records if r['来源'] == '序时账']
    grouped = len(banks) > 1 or len(journals) > 1
    bi, be = _收支(banks)
    ji, je = _收支(journals)
    row = dict(original or {})
    row.update({'核对编号': item['核对编号'], '事项编号': item['事项编号'], '行别': '核对事项', '核对原因': row.get('核对原因') or item['判断依据'], '核对差额': f'收入差 {bi-ji:,.2f}\n支出差 {be-je:,.2f}', '原始出处': '' if grouped else '；'.join(r['原始出处'] for r in records)})
    for side, source in (('银行', banks), ('序时账', journals)):
        row[side + '记录'] = _汇总(source) if grouped or not source else _文字(source[0])
        if not grouped and source:
            _金额(row, side, source[0])
    return row, grouped


def build_display_tables(tables):
    """从既有逐条原记录产生阅读视图；技术表和原计算数据继续保留。"""
    by_item = defaultdict(list)
    for record in tables['核对明细'].to_dict('records'):
        by_item[record['事项编号']].append(record)
    index = tables['复核事项索引'].to_dict('records')
    originals = {row['事项编号']: row for name in ('人工全查', '人工抽样') for row in tables[name].to_dict('records') if row['行别'] == '核对事项'}
    alternatives = defaultdict(list)
    for record in tables['其他对应供选择'].to_dict('records'):
        alternatives[record['事项编号']].append(record)
    pair_rows, group_rows, single_rows = [], [], []
    manual_rows = {'人工全查': [], '人工抽样': []}
    for item in index:
        records = by_item[item['事项编号']]
        banks = [r for r in records if r['来源'] == '银行流水']
        journals = [r for r in records if r['来源'] == '序时账']
        if len(banks) == len(journals) == 1:
            bank, journal = banks[0], journals[0]
            row = {'核对编号': item['核对编号'], '事项编号': item['事项编号'], '银行日期': bank['日期'], '银行记录': '\n'.join(str(x) for x in (bank['摘要'], bank['对方及凭证']) if x), '银行金额': bank['金额'], '序时账日期': journal['日期'], '序时账记录': '\n'.join(str(x) for x in (journal['摘要'], journal['对方及凭证']) if x), '序时账金额': journal['金额'], '差额': float(Decimal(str(bank['金额'])) - Decimal(str(journal['金额']))), '核对依据': item['判断依据'], '核对方式': item['核对方式'], '银行原始出处': bank['原始出处'], '序时账原始出处': journal['原始出处'], '银行记录号': bank['原始记录号'], '序时账记录号': journal['原始记录号']}
            pair_rows.append(row)
        elif banks and journals:
            main, _ = _主行(item, records)
            group_rows.extend([main, *_组成(item, records)])
        else:
            single_rows.extend(records)
        mode = item['核对方式']
        if mode in manual_rows:
            main, grouped = _主行(item, records, originals[item['事项编号']])
            manual_rows[mode].append(main)
            if grouped:
                manual_rows[mode].extend(_组成(item, records))
            # 其他候选只是供选择的说明，不把候选金额重复计入原始组成列。
            for other in alternatives[item['事项编号']]:
                side = '银行' if other['来源'] == '银行流水' else '序时账'
                manual_rows[mode].append({'核对编号': item['核对编号'], '事项编号': item['事项编号'], '行别': other['可选结果'].replace('采用', ''), side + '记录': str(other['摘要及对方']) + f"\n金额 {other['金额']:,.2f}", '核对原因': other['关系说明'], '原始出处': other['原始出处']})
    tables['逐笔核对'] = pd.DataFrame(pair_rows, columns=['核对编号', '银行日期', '银行记录', '银行金额', '序时账日期', '序时账记录', '序时账金额', '差额', '当前核对结论', '核对依据', '核对方式', '银行原始出处', '序时账原始出处', '事项编号', '银行记录号', '序时账记录号'])
    tables['整组核对'] = pd.DataFrame(group_rows, columns=[x for x in 组展示列 if x not in ('人工核对结果', '备注', '填报编号', '银行金额', '序时账金额', '差额')])
    tables['未对应记录'] = pd.DataFrame(single_rows, columns=['核对编号', '来源', '日期', '摘要', '对方及凭证', '金额', '当前核对结果', '核对方式', '原始出处', '事项编号', '原始记录号'])
    for name, rows in manual_rows.items():
        tables[name] = pd.DataFrame(rows, columns=组展示列)
    month = tables['月度核对'].copy()
    for name in ('银行-月末余额', '序时账-月末余额', '余额差额'):
        if name not in month:
            month[name] = None
    first = ['月份', '银行-收入金额', '序时账-收入金额', '收入金额差额', '银行-支出金额', '序时账-支出金额', '支出金额差额', '银行-月末余额', '序时账-月末余额', '余额差额', '余额衔接差', '余额说明']
    tables['月度核对'] = month[[x for x in first if x in month] + [x for x in month if x not in first]]
    return tables


def _headers(sheet):
    return {cell.value: cell.column for cell in sheet[1] if cell.value is not None}


def format_display_sheets(book):
    """引用现有公式的结果，再设置阅读顺序、折叠与唯一填写入口。"""
    references = {name: (_headers(book[name]), max(2, book[name].max_row)) for name in ('核对明细', '复核事项索引', '人工全查', '人工抽样')}
    def find(name, key_label, key, label):
        headers, end = references[name]
        value_col, key_col = (get_column_letter(headers[x]) for x in (label, key_label))
        return f"""IFERROR(INDEX('{name}'!${value_col}$2:${value_col}${end},MATCH({key},'{name}'!${key_col}$2:${key_col}${end},0)),"")"""
    manual_links = {}
    for name in ('人工全查', '人工抽样'):
        sheet = book[name]
        hs = _headers(sheet)
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, hs['填报编号']).value:
                manual_links[str(sheet.cell(row, hs['事项编号']).value)] = (name, row, hs['人工核对结果'])
    for name in ('逐笔核对', '整组核对', '未对应记录', '人工全查', '人工抽样'):
        sheet = book[name]
        hs = _headers(sheet)
        for row in range(2, sheet.max_row + 1):
            item = sheet.cell(row, hs['事项编号'])
            current = find('复核事项索引', '事项编号', item.coordinate, '当前核对结论')
            if name == '逐笔核对':
                keys = [sheet.cell(row, hs[label]).coordinate for label in ('银行记录号', '序时账记录号')]
                owners = [find('核对明细', '原始记录号', key, '当前对应') for key in keys]
                numbers = [find('复核事项索引', '事项编号', owner, '核对编号') for owner in owners]
                bank_location = f'IF({owners[0]}="","未对应","已确认对应（核对"&{numbers[0]}&"）")'
                journal_location = f'IF({owners[1]}="","未对应","已确认对应（核对"&{numbers[1]}&"）")'
                current = f'IF(AND({owners[0]}={owners[1]},{owners[0]}<>""),IF({owners[0]}={item.coordinate},{current},"已并入核对"&{numbers[0]}&"，对应已确认"),IF(AND({owners[0]}="",{owners[1]}=""),{current},"银行："&{bank_location}&"；序时账："&{journal_location}))'
            if name == '整组核对' or (name in ('人工全查', '人工抽样') and sheet.cell(row, hs['行别']).value == '核对事项' and all(sheet.cell(row, hs[label]).value is None for label in 组成金额列)):
                current = f'SUBSTITUTE({current},"原对应被人工选择替代，记录未对应","原对应已被替代；展开查看各笔当前去向")'
                current = f'SUBSTITUTE({current},"原对应已否定，记录恢复未对应","原对应已否定；展开查看各笔当前去向")'
            elif name in ('人工全查', '人工抽样') and sheet.cell(row, hs['行别']).value == '核对事项':
                current = f'SUBSTITUTE({current},"原对应被人工选择替代，记录未对应","原对应已被替代；当前去向见逐笔核对")'
                current = f'SUBSTITUTE({current},"原对应已否定，记录恢复未对应","原对应已否定；当前去向见逐笔核对")'
            if '当前核对结论' in hs and (name == '逐笔核对' or sheet.cell(row, hs['行别']).value == '核对事项'):
                sheet.cell(row, hs['当前核对结论'], '=' + current)
            if '原始记录号' in hs and sheet.cell(row, hs['原始记录号']).value:
                target = '当前核对结果' if name == '未对应记录' else '当前记录结果'
                key = sheet.cell(row, hs['原始记录号']).coordinate
                record_result = '=' + find('核对明细', '原始记录号', key, '当前核对结果')
                sheet.cell(row, hs[target], record_result)
                owner = find('核对明细', '原始记录号', key, '当前对应')
                number = find('核对明细', '原始记录号', key, '当前核对编号')
                display = f'=IF(AND({owner}<>"",{owner}<>{item.coordinate}),"已并入核对"&{number}&"，对应已确认",{record_result[1:]})'
                if '当前核对结论' in hs:
                    sheet.cell(row, hs['当前核对结论'], display)
                else:
                    sheet.cell(row, hs[target], display)
            if '行别' in hs and '其他对应' in str(sheet.cell(row, hs['行别']).value):
                sheet.cell(row, hs['当前核对结论']).value = None
            if name not in ('人工全查', '人工抽样') and str(item.value) in manual_links:
                dest, dest_row, dest_col = manual_links[str(item.value)]
                dest_headers, dest_end = references[dest]
                key_col = get_column_letter(dest_headers['填报编号'])
                label = sheet.cell(row, hs['核对编号']).value
                sheet.cell(row, hs['核对编号'], f'''=HYPERLINK("#'{dest}'!"&ADDRESS(MATCH({item.coordinate},'{dest}'!${key_col}$2:${key_col}${dest_end},0)+1,{dest_col},4),"{label}")''')
        _style(sheet)
    _monthly(book['月度核对'])
    summary = book['核对结论']
    hs = _headers(summary)
    messages = {
        '使用顺序': '先看月度收支与余额，再看逐笔或整组。整组左侧“＋”展开原始记录。人工全查、人工抽样直接在黄色格选择结果。',
        '金额差额口径': '收入与支出分别比较。整组汇总金额写在文字格；收入、支出数值列只记录原始组成，避免重复加总。逐笔金额正数为收入、负数为支出。',
        '月度说明': '月度收支和余额在同一张表。人工确认不改变原始金额、记账月份和客观差额；年度合计只加总发生额。',
        '专题说明': '每日、退款重付、手续费、跨期、重复线索及输入详情保留在隐藏表中，可取消隐藏查看。当前核对结果仍显示原有判断与限制。',
        '人工结果说明': '黄色格每项只选一次并保存。点击逐笔、整组或未对应表的核对编号，可跳到人工填写处。原始分组保留，当前结果随选择更新。',
    }
    for row in range(2, summary.max_row + 1):
        label = summary.cell(row, hs['项目']).value
        if label in messages:
            summary.cell(row, hs['数值'], messages[label])
    summary.column_dimensions[get_column_letter(hs['项目'])].width = 28
    summary.column_dimensions[get_column_letter(hs['数值'])].width = 100
    for row in range(2, summary.max_row + 1):
        value = summary.cell(row, hs['数值'])
        lines = max(1, (len(str(value.value or '')) * 2 + 99) // 100) if value.data_type != 'f' else 1
        summary.row_dimensions[row].height = max(25, lines * 15 + 6)
    order = [book[name] for name in 阅读表]
    book._sheets = order + [sheet for sheet in book if sheet not in order]
    for sheet in book:
        sheet.sheet_state = 'visible' if sheet.title in 阅读表 else 'hidden'
    book.active = 0


def _style(sheet):
    hs = _headers(sheet)
    manual = sheet.title in ('人工全查', '人工抽样')
    grouped = '行别' in hs
    hidden = {'事项编号', '原始记录号', '银行记录号', '序时账记录号', '填报编号', '当前记录结果'}
    if manual:
        hidden.update(('银行金额', '序时账金额', '差额'))
    widths = {'核对编号': 9, '行别': 12, '银行日期': 12, '序时账日期': 12, '日期': 12, '银行记录': 34, '序时账记录': 34, '银行金额': 14, '序时账金额': 14, '金额': 14, '差额': 13, '核对差额': 22, '人工核对结果': 24, '当前核对结论': 29, '当前记录结果': 27, '当前核对结果': 30, '核对原因': 36, '核对依据': 36, '核对方式': 12, '摘要': 30, '对方及凭证': 27, '原始出处': 28, '银行原始出处': 26, '序时账原始出处': 26, '备注': 25}
    for col in range(1, min(hs.values())):
        sheet.column_dimensions[get_column_letter(col)].hidden = True
    for label, col in hs.items():
        dimension = sheet.column_dimensions[get_column_letter(col)]
        dimension.hidden = label in hidden
        dimension.width = widths.get(label, 14)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.freeze_panes = 'C2' if grouped else 'B2'
    sheet.sheet_properties.outlinePr.summaryBelow = False
    sheet.sheet_view.showOutlineSymbols = True
    sheet.auto_filter.ref = None if grouped else f'A1:{get_column_letter(sheet.max_column)}{sheet.max_row}'
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='24445B')
        cell.font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    last_main = None
    for row in range(2, sheet.max_row + 1):
        kind = sheet.cell(row, hs['行别']).value if grouped else ''
        main = kind == '核对事项'
        other = '其他对应' in str(kind)
        if main:
            last_main = row
        child = grouped and not main
        sheet.row_dimensions[row].outlineLevel = int(child)
        sheet.row_dimensions[row].hidden = child and sheet.title == '整组核对'
        sheet.row_dimensions[row].collapsed = False
        if child and last_main and sheet.title == '整组核对':
            sheet.row_dimensions[last_main].collapsed = True
        height = 40
        for label, col in hs.items():
            cell = sheet.cell(row, col)
            cell.font = Font(name='微软雅黑', size=10, bold=main, color='243746')
            cell.fill = PatternFill('solid', fgColor='DEEAF2' if main else 'FFF8E8' if other else 'F3F7FA' if kind == '银行组成' else 'FFFFFF')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            if main:
                cell.border = Border(top=Side(style='thin', color='91A8B8'))
            if label in 组成金额列 or label in ('银行金额', '序时账金额', '金额', '差额'):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
            if label.endswith('日期'):
                cell.number_format = 'yyyy-mm-dd'
            if label == '人工核对结果' and manual and main:
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
            elif label == '人工核对结果' and child:
                # 组成只展示当前记录结果，人工操作格仅主行出现。
                cell.number_format = ';;;'
            if cell.hyperlink or (cell.data_type == 'f' and str(cell.value).startswith('=HYPERLINK(')):
                cell.font = Font(name='微软雅黑', size=10, color='0563C1', underline='single')
            if label not in hidden and cell.data_type != 'f':
                width = widths.get(label, 14)
                lines = sum(max(1, int((sum(2 if ord(x) > 127 else 1 for x in part) + width - 1) // width)) for part in str(cell.value or '').split('\n'))
                height = max(height, lines * 15 + 8)
        sheet.row_dimensions[row].height = min(height, 350)
    for label in ('银行金额', '序时账金额', '金额'):
        if label in hs and label not in hidden:
            sheet.cell(1, hs[label]).comment = Comment('正数为收入，负数为支出。', '核对说明')
    if grouped:
        sheet.cell(1, hs['银行记录']).comment = Comment('汇总行先看双方笔数、收入与支出；组成行先银行后序时账。两侧排在同一组不表示已经逐笔配对。', '核对说明')
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'
    sheet.print_area = f'A1:{get_column_letter(max(c for k, c in hs.items() if k not in hidden))}{sheet.max_row}'
    sheet.page_margins.left = sheet.page_margins.right = 0.2


def _monthly(sheet):
    hs = _headers(sheet)
    end = sheet.max_row
    balance_cols = [hs[x] for x in ('银行-月末余额', '序时账-月末余额', '余额差额')]
    for row in range(2, end + 1):
        for col in balance_cols:
            if sheet.cell(row, col).value is None:
                sheet.cell(row, col, '缺失')
                sheet.cell(row, col).font = Font(name='微软雅黑', size=10, color='9C6500')
    if end >= 2:
        for label in ('收入金额差额', '支出金额差额', '余额差额', '余额衔接差'):
            letter = get_column_letter(hs[label])
            sheet.conditional_formatting.add(f'{letter}2:{letter}{end}', FormulaRule(formula=[f'AND(ISNUMBER({letter}2),{letter}2<>0)'], fill=PatternFill('solid', fgColor='FFF2CC'), font=Font(color='9C0006')))
    years = sorted({str(sheet.cell(row, hs['月份']).value)[:4] for row in range(2, end + 1)})
    for year in years:
        rows = [r for r in range(2, end + 1) if str(sheet.cell(r, hs['月份']).value).startswith(year + '-')]
        if not rows:
            continue
        target = sheet.max_row + 1
        sheet.cell(target, hs['月份'], year + '年合计')
        for label, col in hs.items():
            if '金额' in label or '笔数' in label or '变动净额' in label or label.startswith(('已确认', '未确认')):
                letter = get_column_letter(col)
                sheet.cell(target, col, '=' + '+'.join(f'{letter}{r}' for r in rows))
                sheet.cell(target, col).number_format = '#,##0.00;[Red]-#,##0.00'
        sheet.cell(target, hs['余额说明'], '发生额合计；余额不加总')
        sheet.row_dimensions[target].height = 32
        for cell in sheet[target]:
            cell.font = Font(name='微软雅黑', size=10, bold=True, color='243746')
            cell.fill = PatternFill('solid', fgColor='DEEAF2')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.auto_filter.ref = f'A1:{get_column_letter(sheet.max_column)}{end}'
    sheet.print_area = f'A1:{get_column_letter(sheet.max_column)}{sheet.max_row}'
    sheet.freeze_panes = 'B2'
    for label, col in hs.items():
        sheet.column_dimensions[get_column_letter(col)].width = 13 if label != '余额说明' else 28
